"""The product-code and order-number hard rules (James Kuo, 30 Sep 2026), in Python and in the interface page."""
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import config  # noqa: E402
import product_code as pc  # noqa: E402


def test_code_rule():
    # "10 means 1mm and 90 means 9mm. A0 is 10mm and B0 is 11mm"
    assert pc.thickness_mm('RPA10WB1') == 1.0 and pc.thickness_mm('RPA90WB1') == 9.0
    assert pc.thickness_mm('RPAA0WB318') == 10.0 and pc.thickness_mm('RPAB0WB1') == 11.0 and pc.thickness_mm('RPP63RF1') == 6.3
    assert pc.problems('RPA40WB3051') == []
    assert pc.problems('RPAQWRAN72')                      # "AQ is incorrect from the start"
    assert pc.problems('RPA40WR3051')                     # "the color code is WB. WR dont exist"
    assert pc.problems('RPA40WB30S1')                     # "after that is just number for the product. No english characters"
    assert pc.problems('DPPAOKS27')                       # letter O in the thickness: must be A0
    assert pc.repair('DPPAOKS27') == ('DPPA0KS27', True)
    assert pc.repair('RPA4OWB3143')[0] == 'RPA40WB3143' and pc.repair('RBP30ER22')[0] is None


def test_product_master_keeps_the_rule():
    """No code breaking the rule may be in the published Product Master (29 Sep: DPPAOKS27, RBPAOKS14/37/40 were)."""
    from openpyxl import load_workbook
    p = config.published_path('Product Master.xlsx')
    if not (p and p.exists()):
        pytest.skip('Product Master not published on this PC')
    codes = [r[0] for r in load_workbook(p, read_only=True)['Product Master'].iter_rows(min_row=2, max_col=1, values_only=True) if r[0]]
    assert [c for c in codes if pc.CODE_RX.match(c) is None] == []


def test_packets_keep_the_rule():
    for f in sorted((ROOT / 'data' / 'packets').glob('packet_*.json')):
        pk = json.loads(f.read_text(encoding='utf-8'))
        bad = [r['prod_code'] for e in pk['ext'] for r in e['rows'] if pc.CODE_RX.match(r['prod_code']) is None]
        bad += [r['product_code'] for c in pk['cnv'] for r in c['rows'] if pc.CODE_RX.match(r['product_code']) is None]
        assert bad == [], (f.name, bad)


PAGE_JS = r'''
const fs = require('fs');
const html = fs.readFileSync(process.argv[2], 'utf8');
const D = JSON.parse(/<script type="application\/json" id="data">([\s\S]*?)<\/script>/.exec(html)[1]);
const PRODUCTS = new Set(D.products);
const src = /<script>([\s\S]*)<\/script>/.exec(html.replace(/<script type="application\/json"[\s\S]*?<\/script>/, ''))[1];
const a = src.indexOf('const PAIRS = new Set('), b = src.indexOf('async function ocrPage(');
eval(src.slice(a, b).replace(/^const /gm, 'var ').replace(/^let /gm, 'var '));
const pairs = [['H69A039-1', 'DPP30WB1023'], ['RP26826-3', 'RPA40WB3072'], ['H66A116-1', 'DPP50WB308']];
console.log(JSON.stringify({
  suffix17: parseOrderCell('H69A203-17'),                 // a 2-digit suffix is allowed only as read in its own cell
  order: parseOrderCell('H69A203-1'), badOrder: parseOrderCell('HSA0S-1'),
  aq: codeProblems('RPAQWRAN72').length > 0, wr: codeProblems('RPA40WR3051').length > 0, ok: codeProblems('RPA40WB3051'),
  repair: repairCode('RPA4OWB3143'), take: takeProd('RPAQWRAN72'), take2: takeProd('RPA4OWB3072'),
  judgeHist: judge('RP26826-3', 'RPAQWRAN72', pairs, '').status,       // product garbage: not taken from history
  judgeNear: judge('H06A116-1', 'DPP50WB308', pairs, '').status,       // 0/6 is not a look-alike: not taken
  judgeOk: judge('RP26826-3', 'RPA4OWB3072', pairs, '').status,
}));
'''


def test_page_reader_rules(tmp_path):
    """The interface page's reader keeps the same rules (ui/page.template.html, built by ui/build.py)."""
    node = shutil.which('node')
    page = ROOT / 'out' / 'profile-formulation.html'
    if not node or not page.exists():
        pytest.skip('node or the built page is not available')
    js = tmp_path / 't.js'
    js.write_text(PAGE_JS, encoding='utf-8')
    r = subprocess.run([node, str(js), str(page)], capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    out = json.loads(r.stdout)
    assert out['order'] == 'H69A203-1' and out['badOrder'] is None
    assert out['aq'] and out['wr'] and out['ok'] == []
    assert out['repair'] == 'RPA40WB3143' and out['take'] is None and out['take2'] == 'RPA40WB3072'
    assert out['judgeHist'] == 'check'                  # RPAQWRAN72 never reaches the printed copy
    assert out['judgeNear'] != 'read' and out['judgeNear'] != 'matched to schedule history'
    assert out['judgeOk'] == 'matched to schedule history'
