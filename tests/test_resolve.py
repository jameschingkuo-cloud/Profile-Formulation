"""Backtest: the FRM Draft for 25 Sep, built only from the 23-24 Sep issued FRMs, against what Tech issued on 25 Sep."""
import json
import os
import subprocess
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def test_draft_matches_issued_25_sep(tmp_path):
    env = dict(os.environ, PYTHONUTF8="1", PKT_DATE="2026-09-25", RUN_DATE="2026-09-28", HISTORY_BEFORE="2026-09-25",
               OUTPUT_DIR=str(tmp_path / "out"), WORK_DIR=str(tmp_path / "work"))
    r = subprocess.run([sys.executable, str(ROOT / "daily/resolve.py")], env=env, capture_output=True, text=True, cwd=ROOT)
    assert r.returncode == 0, r.stderr[-2000:]
    wb = load_workbook(tmp_path / "out" / "FRM Draft 2026-09-25.xlsx")
    h = [c.value for c in wb["Draft"][1]]
    prop = defaultdict(set)
    for row in wb["Draft"].iter_rows(min_row=2, values_only=True):
        d = dict(zip(h, row))
        prop[(d["Line Code"], d["Order"])].add((d["Formula Code"], d["Feeder"], d["Material (as issued)"], str(d["Set"])))
    issued = defaultdict(set)
    pk = json.loads((ROOT / "data/packets/packet_2026-09-25.json").read_text(encoding="utf-8"))
    for pg in pk["frm"]:
        for g in pg["groups"]:
            for o in g["orders"]:
                for f in g["formulas"]:
                    for col, v in f["feeders"].items():
                        if v["material"] or v["set"]:
                            feeder = col.replace("Extruder ", "")
                            feeder = feeder.split(" ", 1)[1] if feeder[:2] in ("A ", "B ", "C ", "D ") else feeder
                            issued[(pg["line_code"], o)].add((f["formula_code"], feeder, v["material"], v["set"]))
    differ = sorted(k for k in prop if prop[k] != issued.get(k))
    assert len(prop) == 69
    assert differ == [("SE25", "RP26923-1")]          # the one formula Tech changed that day
    exc = {row[1] for row in wb["Exceptions"].iter_rows(min_row=2, values_only=True)}
    assert exc == {"H67A120-1", "H68A153-1"}          # the two new orders go to an engineer
