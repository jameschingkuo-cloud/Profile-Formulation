"""Tech's past formulation pages (history/frm_scans.py -> work/history/frm_scans.json) into the Formulation Master (James Kuo,
30 Sep 2026: "go ahead and add everything after 2020"; "if the fomulation is obsolete one (replaced by something newer)
put it in the log instead").

    python db/import_frm_history.py --by "James Kuo" --why "..." [--dry-run]

Issues are taken newest first. For each formula on a page (line, code, variant):
  - the master has it on that line with the same settings  -> nothing to do;
  - the master has it on that line with other settings     -> OBSOLETE (the master's is newer): Change Log only,
                                                              Field '(superseded - history only)', nothing changed;
  - a row on the same page is marked '(New Formula)'        -> the row it replaces is OBSOLETE (log only);
  - not in the master                                      -> added as Draft (Formulas + Line Settings), logged;
    unless a printed material is not mapped: then nothing of that formula is written, it is listed for James.
For each order's product: a Product to Formula row already there -> nothing; the product already has a formula on that
line from a later issue -> OBSOLETE (log only); else added as Draft (Last Run = the issue date), logged.
Two material spellings found on these pages are added to Materials 'Other Spellings' (logged): FR-GPP30003 MP and
WB-NPC PE-W22151 (clear matches). 'Vistamaxx 3588FL Pre-mix' is not mapped (a pre-mix, not the pure 3588FL).
Works on a copy of the published master in out/; then publish.py, preflight check, preflight accept.
"""
import argparse
import datetime
import json
import re
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT)); sys.path.insert(0, str(ROOT / 'daily'))
import config  # noqa: E402
from db import preflight, schema  # noqa: E402
from db.import_frm import replaced  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from resolve import split_feeder, variant  # noqa: E402

NAME = schema.FM
NEW_SPELLINGS = {'50-7002-361': 'FR-GPP30003 MP', '50-7002-393': 'WB-NPC PE-W22151'}
OBSOLETE = '(superseded - history only)'


def main(argv):
    ap = argparse.ArgumentParser()
    ap.add_argument('--by', required=True); ap.add_argument('--why', required=True)
    ap.add_argument('--dry-run', action='store_true')
    ap.add_argument('--json', default=str(config.WORK_DIR / 'history' / 'frm_scans.json'))
    a = ap.parse_args(argv)
    config.record_read(a.json, 'Tech formulation pages read from older scans')
    issues = sorted(json.loads(Path(a.json).read_text(encoding='utf-8')), key=lambda p: p['packet_date'], reverse=True)
    ok, lines, (cur, _) = preflight.check(NAME)
    if not ok:
        print('\n'.join(lines))
        raise SystemExit('STOP: the published master has unaccepted changes (or no accepted version); settle that first.')

    spell = {}
    for mid, m in sorted(cur['Materials'].items(), key=lambda kv: kv[1].get('Status') != 'Retired'):
        for t in [m.get('FRM Text')] + str(m.get('Other Spellings') or '').split(' | '):
            if t and t.strip():
                spell[t.strip()] = mid
    spell.update(replaced())
    spell.update({v: k for k, v in NEW_SPELLINGS.items()})
    sets = {}                                            # (line, code, var) -> {(ext, feeder): (mid, set)}
    for s in cur['Line Settings'].values():
        sets.setdefault((s['Line Code'], s['Formula Code'], s['Variant']), {})[(s['Extruder'] or '', s['Feeder'])] = (s['Material ID'], str(s['Set']))
    have_f = {(f['Formula Code'], f['Variant']) for f in cur['Formulas'].values()}
    p2f = {}                                             # (product, line) -> [(code, var, last run)]
    for p in cur['Product to Formula'].values():
        p2f.setdefault((p['Product Code'], p['Line Code']), []).append((p['Formula Code'], p['Variant'], str(p.get('Last Run') or '')[:10]))

    add = {'Formulas': [], 'Line Settings': [], 'Product to Formula': []}
    log_only, not_written, same = [], [], 0
    for pk in issues:
        day = pk['packet_date']
        prod = {r['order']: r['prod_code'] for pg in pk['ext'] for r in pg['rows']}
        for pg in pk['frm']:
            line = pg['line_code']
            src = f"FRM {day} ({pk['frm_source_scan']} p{pg['scan_page']})"
            for g in pg['groups']:
                for i, f in enumerate(g['formulas']):
                    code = f['formula_code']
                    var = 'Primary' if f.get('replaces_row') else variant(f.get('note', ''), i == 0)
                    if var == 'Reclaim run-out' and 'reclaim' not in (f.get('note') or '').lower():
                        var = 'Primary' if i == 0 else 'Other'   # a colour run-out note (NPC PE-W22151) is not the reclaim variant
                    printed, unmapped = {}, []
                    for c, v in f['feeders'].items():
                        if not (v.get('material') or v.get('set')):
                            continue
                        mid = spell.get((v.get('material') or '').strip())
                        if not mid:
                            unmapped.append(f"{c} \"{v.get('material')}\" {v.get('set')}")
                        printed[split_feeder(c)] = (mid, str(v.get('set')))
                    text = '; '.join(f"{e}{' ' if e else ''}{fd} {m} {s}" for (e, fd), (m, s) in sorted(printed.items()))
                    k = (line, code, var)
                    if f.get('superseded_on_page'):
                        log_only.append(('Line Settings', '|'.join([line, code, var]), '', text, src,
                                         f'Issued {day}; replaced on the same page by the row marked (New Formula)'))
                        continue
                    if k in sets:
                        if sets[k] == printed:
                            same += 1
                        else:
                            old = '; '.join(f"{e}{' ' if e else ''}{fd} {m} {s}" for (e, fd), (m, s) in sorted(sets[k].items()))
                            log_only.append(('Line Settings', '|'.join(k), old, text, src,
                                             f'Issued {day}; the master holds a newer version of {code} ({var}) on {line}'))
                    elif unmapped:
                        not_written.append(f"{day} {line} {code} {var}: material not mapped: {', '.join(unmapped)} (orders {', '.join(g['orders'])})")
                        continue
                    else:
                        sets[k] = printed
                        if (code, var) not in have_f:
                            have_f.add((code, var))
                            add['Formulas'].append({'Formula Code': code, 'Variant': var, 'Family': code[:2], 'When to Use': f.get('note') or None,
                                                    'Status': 'Draft', 'Seeded From': src})
                        for (ext, fd), (mid, st) in sorted(printed.items()):
                            add['Line Settings'].append({'Line Code': line, 'Formula Code': code, 'Variant': var, 'Extruder': ext or None,
                                                         'Feeder': fd, 'Material ID': mid, 'Set': st, 'Source': 'FRM', 'Status': 'Draft'})
                    for o in g['orders']:
                        p = prod.get(o)
                        if not p:
                            not_written.append(f'{day} {line} {o}: product not found on any extrusion schedule (Product to Formula not written)')
                            continue
                        rows = p2f.get((p, line), [])
                        if any((c, v) == (code, var) for c, v, _ in rows):
                            continue
                        later = [x for x in rows if x[2] > day]
                        if later:
                            log_only.append(('Product to Formula', '|'.join([p, line, code, var]), ', '.join(f'{c} {v} ({d})' for c, v, d in later),
                                             f'{code} {var} ({day}, order {o})', src,
                                             f'Issued {day}; {p} has a later formula on {line}'))
                            continue
                        rows.append((code, var, day)); p2f[(p, line)] = rows
                        add['Product to Formula'].append({'Product Code': p, 'Line Code': line, 'Formula Code': code, 'Variant': var,
                                                          'Priority': 'Primary' if var == 'Primary' else 'Alternate',
                                                          'Last Run': datetime.date.fromisoformat(day), 'Status': 'Draft'})

    for sheet, rows in add.items():
        for r in rows:
            print(f'ADD {sheet}: ' + ' | '.join(f'{k}={v}' for k, v in r.items() if v not in (None, '')))
    for r in log_only:
        print(f'LOG ONLY (obsolete) {r[0]} {r[1]}: {r[5]}')
    for u in not_written:
        print('NOT WRITTEN: ' + u)
    print(f'{same} formula(s) already in the master as printed; {sum(map(len, add.values()))} row(s) to add; '
          f'{len(log_only)} obsolete logged; {len(not_written)} not written')
    if a.dry_run:
        print('dry run - nothing written')
        return add, log_only, not_written

    out = config.OUTPUT_DIR / NAME
    shutil.copy(config.published_path(NAME), out)
    wb = load_workbook(out)
    log = wb['Change Log']
    nid = max([r[0] for r in log.iter_rows(min_row=2, values_only=True) if isinstance(r[0], int)] or [0]) + 1
    first, today = nid, datetime.date.today()
    keys = {'Formulas': ['Formula Code', 'Variant'], 'Line Settings': ['Line Code', 'Formula Code', 'Variant', 'Extruder', 'Feeder'],
            'Product to Formula': ['Product Code', 'Line Code', 'Formula Code', 'Variant']}
    mat = wb['Materials']                                # the two new spellings
    mh = {c.value: j for j, c in enumerate(mat[1], 1)}
    for row in mat.iter_rows(min_row=2):
        mid = row[mh['Material ID'] - 1].value
        if mid in NEW_SPELLINGS:
            cell = row[mh['Other Spellings'] - 1]
            old = cell.value
            cell.value = ' | '.join(x for x in [old, NEW_SPELLINGS[mid]] if x)
            log.append([nid, today, 'Materials', mid, 'Other Spellings', old, cell.value,
                        f'Spelling printed on Tech\'s past FRM pages ({NEW_SPELLINGS[mid]})', a.by, a.by, 'history/frm_scans.py'])
            nid += 1
    for sheet, rows in add.items():
        ws = wb[sheet]
        h = {c.value: j for j, c in enumerate(ws[1], 1)}
        for r in rows:
            vals = [None] * len(h)
            for k, v in r.items():
                vals[h[k] - 1] = v
            ws.append(vals)
            log.append([nid, today, sheet, '|'.join(str(r.get(k) or '') for k in keys[sheet]), '(new row)', None,
                        ', '.join(f'{k} {r[k]}' for k in ('Material ID', 'Set', 'Priority') if r.get(k)) or r.get('Status'),
                        a.why, a.by, a.by, r.get('Seeded From') or 'history/frm_scans.py'])
            nid += 1
    for sheet, key, old, new, src, why in log_only:
        log.append([nid, today, sheet, key, OBSOLETE, old or None, new,
                    f"{why}. History only, not applied (James Kuo, 30 Sep 2026: 'if the fomulation is obsolete one (replaced by "
                    f"something newer) put it in the log instead')", a.by, a.by, src])
        nid += 1
    wb.save(out)
    print(out, f'- Changes {first}-{nid - 1} logged')
    return add, log_only, not_written


if __name__ == '__main__':
    main(sys.argv[1:])
