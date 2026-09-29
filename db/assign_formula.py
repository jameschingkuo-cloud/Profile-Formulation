"""Record an engineer's decision for an Exception: this product, on this line, runs these formulas (Formulation Master,
Product to Formula), approved and logged in the Change Log. The next `daily/resolve.py` + `daily/render_frm.py` run
fills the order in (James Kuo, 29 Sep 2026: "this allow the engineer to update the data base and trigger a re run").

    python db/assign_formula.py --line SE22 --product RPAA0WB318 --by "James Kuo" --why "..." \\
        --formula FUA060WBA Primary --formula FUA060WBA Other --formula FUA010WBA "Reclaim run-out"

The first --formula is Primary, the rest Alternate; list them in run order (reclaim first, run-out last). Each
(line, formula, variant) must already have Line Settings on that line: otherwise it stops and says so (settings are
added in Excel or with db/edit_master.py first). Existing Draft rows for the same key are approved, not duplicated.
Then: publish.py "Formulation Master.xlsx", preflight check, preflight accept, and re-run the day.
"""
import argparse
import datetime
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import config  # noqa: E402
from db import preflight, schema  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

NAME = schema.FM


def main(argv):
    ap = argparse.ArgumentParser()
    ap.add_argument('--line', required=True); ap.add_argument('--product', required=True)
    ap.add_argument('--by', required=True); ap.add_argument('--why', required=True)
    ap.add_argument('--formula', nargs=2, action='append', required=True, metavar=('CODE', 'VARIANT'))
    a = ap.parse_args(argv)
    variants = next(c.values for s in schema.BY_NAME[NAME].sheets if s.name == 'Formulas' for c in s.cols if c.name == 'Variant')
    for code, var in a.formula:
        if var not in variants:
            raise SystemExit(f'Variant {var!r} is not one of {", ".join(variants)}')
    ok, lines, (cur, _) = preflight.check(NAME)
    if not ok:
        print('\n'.join(lines))
        raise SystemExit('STOP: the published master has unaccepted changes (or no accepted version); settle that first.')
    missing = [f'{code} {var}' for code, var in a.formula
               if not any(s['Line Code'] == a.line and s['Formula Code'] == code and s['Variant'] == var for s in cur['Line Settings'].values())]
    if missing:
        raise SystemExit(f'STOP: no Line Settings on {a.line} for {", ".join(missing)}. Add the settings first. Nothing written.')

    out = config.OUTPUT_DIR / NAME
    shutil.copy(config.published_path(NAME), out)
    wb = load_workbook(out)
    ws, log = wb['Product to Formula'], wb['Change Log']
    h = {c.value: i for i, c in enumerate(ws[1], 1)}
    nid = max([r[0] for r in log.iter_rows(min_row=2, values_only=True) if isinstance(r[0], int)] or [0]) + 1
    today = datetime.date.today()

    def logrow(rk, field, old, new):
        nonlocal nid
        log.append([nid, today, 'Product to Formula', rk, field, old, new, a.why, a.by, a.by, 'Engineer decision for an Exception'])
        nid += 1

    first = nid
    for i, (code, var) in enumerate(a.formula):
        rk = f'{a.product}|{a.line}|{code}|{var}'
        prio = 'Primary' if i == 0 else 'Alternate'
        row = next((r for r in range(2, ws.max_row + 1)
                    if [ws.cell(r, h[k]).value for k in ('Product Code', 'Line Code', 'Formula Code', 'Variant')] == [a.product, a.line, code, var]), None)
        want = {'Priority': prio, 'Status': 'Approved', 'Approved By': a.by, 'Approved Date': today}
        if row is None:
            vals = [None] * len(h)
            for k, v in {'Product Code': a.product, 'Line Code': a.line, 'Formula Code': code, 'Variant': var, **want}.items():
                vals[h[k] - 1] = v
            ws.append(vals)
            logrow(rk, '(new row)', None, f'{prio}, Approved')
        else:
            for k, v in want.items():
                old = ws.cell(row, h[k]).value
                if (old.date() if isinstance(old, datetime.datetime) else old) != v:
                    ws.cell(row, h[k]).value = v
                    logrow(rk, k, old, v.isoformat() if isinstance(v, datetime.date) else v)
    wb.save(out)
    print(out, f'- Changes {first}-{nid - 1} logged' if nid > first else '- already approved; nothing changed')


if __name__ == '__main__':
    main(sys.argv[1:])
