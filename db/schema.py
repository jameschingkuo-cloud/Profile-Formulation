"""The formulation database: every workbook, sheet and column, in one place.

The database is a set of Excel workbooks (James Kuo, 28 Sep 2026: Excel only), one folder per kind
under PUBLISH_DIR (`General\\Engineering Pipeline\\Production Formulation`, James's folders of 24 Sep 2026).
This module is the single description of them. From it:

    python db/schema.py templates            blank workbooks (with reference data) -> OUTPUT_DIR/templates/
    python db/schema.py check <file.xlsx>    check a workbook against its description (headers by name,
                                             keys unique, required cells filled, allowed values)
    python db/schema.py doc                  writes docs/DATABASE_TABLES.md (the column reference)

Kinds of workbook:
    master    approved reference data, changed only through its Change Log (Tech/James approve)
    record    append-only history: one row per thing issued or produced; rows are never edited
    daily     one workbook per day, as issued; never rewritten (publish.py refuses; --reissue only if James asks)
    evidence  rebuilt from sources (Tech's calc workbooks); read to seed or check a master, never edited

Keys: a column marked key=True is part of the sheet's primary key (together unique). `ref` names the
sheet a value must exist in ('Workbook/Sheet/Column').
"""
from __future__ import annotations

import sys
from dataclasses import dataclass, field
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "calc"))


@dataclass
class Col:
    name: str
    type: str = "text"          # text | int | number | date | datetime | enum
    key: bool = False
    required: bool = False
    values: tuple = ()          # allowed values for enum
    ref: str = ""               # 'Workbook/Sheet/Column' the value must exist in
    note: str = ""


@dataclass
class Sheet:
    name: str
    purpose: str
    cols: list[Col]
    append_only: bool = False


@dataclass
class Book:
    name: str                   # file name; <date> = YYYY-MM-DD for daily workbooks
    folder: str                 # subfolder of PUBLISH_DIR
    kind: str                   # master | record | daily | evidence
    purpose: str
    written_by: str
    read_by: str
    sheets: list[Sheet]
    status: str = "new"         # existing | new


# ---- shared column sets -------------------------------------------------------------------------------------
SEVERITY = ("High", "Medium", "Low", "Info")
APPROVAL = ("Draft", "Approved", "Retired")
DOSING_TYPES = ("WEIGHT", "AUGER")
ROLES = ("VIRGIN", "HOMO", "RECLAIM", "TALC", "CACO3", "COLOUR", "MODIFIER", "ADDITIVE", "HDPE", "SKIN", "FOAM", "OTHER")
VARIANTS = ("Primary", "Reclaim run-out", "VOIDFORM", "Sign blank", "Corn box", "Roll", "Other")

ISSUES = Sheet("Issues", "Problems found by the build, for James/Tech to answer (house columns)", [
    Col("Severity", "enum", required=True, values=SEVERITY),
    Col("Document", required=True), Col("Line"), Col("Order"), Col("Check", required=True), Col("Detail"),
    Col("Source"), Col("James / Tech response")])

README = Sheet("Read Me", "Sources read (file, modified time, hash), pre-flight result, notes", [
    Col("Item"), Col("Value")])

CHANGE_LOG = Sheet("Change Log", "Every change to this master: who, why, old -> new. The only way a master changes", [
    Col("Change ID", "int", key=True, required=True),
    Col("Date", "date", required=True),
    Col("Sheet", required=True),
    Col("Row Key", required=True, note="the key of the changed row, e.g. 'FUA152WB4|Primary|A|PC416'"),
    Col("Field", required=True),
    Col("Old Value"), Col("New Value"),
    Col("Why", required=True),
    Col("Requested By", required=True),
    Col("Approved By", required=True, note="Tech or James; a change with no approver is not applied"),
    Col("Source", note="calc workbook + sheet, FRM date, catch test, e-mail ...")], append_only=True)

FM = "Formulation Master.xlsx"
AC = "Auger Calibration.xlsx"
PM = "Product Master.xlsx"

# ---- the workbooks ------------------------------------------------------------------------------------------
BOOKS: list[Book] = [
    Book(PM, "Product Master", "master",
         "One row per product (material master number): extrusion and converting data, formula codes seen",
         "daily/build_master.py (merges each packet); product_master/prepare.py (prepared copy + Issues)",
         "resolve step (product -> formula), Tech",
         [Sheet("Product Master", "One row per Product Code. Columns as built by daily/build_master.py (38 columns); "
                "only the ones the flow depends on are listed here", [
             Col("Product Code", key=True, required=True, note="material master number; never duplicated"),
             Col("Material", "enum", values=("PPP", "BBB")), Col("Grade", "enum", values=("P", "A")),
             Col("Spec"), Col("Colours (3 layers)"),
             Col("Thk (mm)", "number", note="the spec decides; the code's thickness is nominal (James, 28 Sep 2026)"),
             Col("GSM", "number"), Col("Formula Code(s)"), Col("Formula Last Run", "date"),
             Col("Source"), Col("Check"),
             Col("Status", "enum", required=True, values=("Draft", "Verified", "Needs Review", "Obsolete")),
             Col("Last Updated", "date", required=True)]),
          README], status="existing"),

    Book(FM, "Formulation Data Base", "master",
         "The approved formulations. The recipe is the weight % (per extruder on co-ex); settings per line are derived "
         "from it (HANDOFF §7.16). Seeded from the Calc Library, Draft until Tech approves each formula",
         "db/seed (from Formulation Calc Library, then Tech edits through the Change Log)",
         "resolve step, FRM renderer, Tech",
         [Sheet("Lines", "One row per extrusion line (reference data; DOSING is hardcoded, James 26 Sep 2026)", [
             Col("Line Code", key=True, required=True, note="R1: 2 letters + 2 digits"),
             Col("Line No", "int", required=True),
             Col("Dosing", "enum", required=True, values=DOSING_TYPES),
             Col("Feeder Layout", required=True), Col("Extruders", required=True),
             Col("AC"), Col("Form Effective Date"),
             Col("Active", "enum", required=True, values=("Yes", "No"))]),
          Sheet("Materials", "One row per material, keyed on the plant's material code list (IWPFT062)", [
             Col("Material ID", key=True, required=True, note="IWPFT062 item no., e.g. 50-1560-050"),
             Col("Material Code", note="e.g. PC416"), Col("Name", required=True, note="e.g. Formosa F6502A"),
             Col("Supplier"),
             Col("IWPFT062 Status", "enum", required=True, values=("Active", "In-active", "Withdrawn", "Not listed"),
                 note="as on the qualified list; 'Not listed' = plant-internal (INT-) or not on the list (NL-)"),
             Col("Role", "enum", required=True, values=ROLES),
             Col("FRM Text", note="how the FRM page writes it, e.g. 'PP Virgin-silo 3 (6502A)'"),
             Col("Other Spellings", note="every spelling seen in calcs/FRM, ' | ' separated"),
             Col("Approved Substitutes"), Col("Bulk Density (g/cm3)", "number"),
             Col("Status", "enum", required=True, values=APPROVAL)]),
          Sheet("Formulas", "One row per formula code + variant", [
             Col("Formula Code", key=True, required=True),
             Col("Variant", "enum", key=True, required=True, values=VARIANTS),
             Col("Family", note="FU / FS / RU / BF (HANDOFF §6.1, unconfirmed)"),
             Col("Description"), Col("When to Use", note="for variants: the FRM note, e.g. 'in case PP WB Reclaim runs out'"),
             Col("Status", "enum", required=True, values=APPROVAL),
             Col("Approved By"), Col("Approved Date", "date"), Col("Seeded From")]),
          Sheet("Recipe", "The recipe: weight % per material. Each extruder adds to 100", [
             Col("Formula Code", key=True, required=True, ref=f"{FM}/Formulas/Formula Code"),
             Col("Variant", "enum", key=True, required=True, values=VARIANTS),
             Col("Extruder", key=True, note="A-D on co-ex lines; blank = the only extruder"),
             Col("Material ID", key=True, required=True, ref=f"{FM}/Materials/Material ID"),
             Col("Weight %", "number", note="blank when Balance = Yes"),
             Col("Balance", "enum", values=("Yes", "No"), note="Yes = takes the rest to 100 ('Auto' on weight lines)"),
             Col("Note")]),
          Sheet("Line Settings", "What the floor sets, per line. Weight lines: Set = %. Auger lines: Set = speed 0-100, "
                "derived from the recipe and the current slope (setting = % x T / slope, whole numbers)", [
             Col("Line Code", key=True, required=True, ref=f"{FM}/Lines/Line Code"),
             Col("Formula Code", key=True, required=True, ref=f"{FM}/Formulas/Formula Code"),
             Col("Variant", "enum", key=True, required=True, values=VARIANTS),
             Col("Extruder", key=True),
             Col("Feeder", key=True, required=True, note="H1-H5, V1-V9, A1 ... as printed on that line's FRM page"),
             Col("Material ID", required=True, ref=f"{FM}/Materials/Material ID"),
             Col("Set", required=True, note="number, or 'Auto' on weight lines"),
             Col("Slope Used", "number", note="auger lines: from Auger Calibration (current)"),
             Col("Weight % (from Set)", "number", note="auger: slope x Set / total; label: target by calibration"),
             Col("Recipe Weight %", "number"),
             Col("Deviation (pts)", "number", note="Weight % (from Set) - Recipe Weight %"),
             Col("Source", "enum", required=True, values=("Calc block", "FRM", "Derived", "Tech")),
             Col("Status", "enum", required=True, values=APPROVAL)]),
          Sheet("Product to Formula", "Which formula a product gets on a line. Primary + alternates", [
             Col("Product Code", key=True, required=True, ref=f"{PM}/Product Master/Product Code"),
             Col("Line Code", key=True, required=True, ref=f"{FM}/Lines/Line Code"),
             Col("Formula Code", key=True, required=True, ref=f"{FM}/Formulas/Formula Code"),
             Col("Variant", "enum", key=True, required=True, values=VARIANTS),
             Col("Priority", "enum", required=True, values=("Primary", "Alternate")),
             Col("Last Run", "date"), Col("Status", "enum", required=True, values=APPROVAL),
             Col("Approved By"), Col("Approved Date", "date")]),
          Sheet("Standing Notes", "Footer notes printed on a line's FRM page", [
             Col("Line Code", key=True, required=True, ref=f"{FM}/Lines/Line Code"),
             Col("No", "int", key=True, required=True), Col("Note", required=True)]),
          CHANGE_LOG, ISSUES, README]),

    Book(AC, "Formulation Data Base", "master",
         "Auger hardware and calibration slopes for the auger lines. Weight lines need none",
         "db/seed (from Calc Library), then catch tests through the Change Log",
         "Line Settings derivation, checks A1-A7 (calc/auger_rules.py)",
         [Sheet("Hoppers", "Hardware per line and hopper, and the role it may carry (auger_rules.RULES, draft Q13)", [
             Col("Line Code", key=True, required=True, ref=f"{FM}/Lines/Line Code"),
             Col("Hopper", key=True, required=True),
             Col("Gear Ratio", required=True), Col("Screw", required=True),
             Col("Main Role", "enum", values=ROLES), Col("Allowed Roles"),
             Col("Since", "date", note="hardware change voids the slopes (A4)")]),
          Sheet("Calibration", "One current slope per line + hopper + material (or approved calibration family)", [
             Col("Line Code", key=True, required=True, ref=f"{FM}/Lines/Line Code"),
             Col("Hopper", key=True, required=True),
             Col("Material ID", key=True, required=True, ref=f"{FM}/Materials/Material ID"),
             Col("Slope (g/min per unit)", "number", required=True),
             Col("Family", note="set when a material borrows another's slope with approval (A3)"),
             Col("Calibrated Date", "date"), Col("Last Verified", "date"),
             Col("Verified By"), Col("Method", "enum", values=("Catch test", "From calc sheet", "IWPFM031", "Other")),
             Col("Overdue", "enum", values=("Yes", "No")),
             Col("Status", "enum", required=True, values=("Current", "Superseded", "Suspect"))]),
          Sheet("Verification Log", "Timed catch tests (the only measurement of an auger)", [
             Col("Test Date", "date", key=True, required=True),
             Col("Line Code", key=True, required=True), Col("Hopper", key=True, required=True),
             Col("Material ID", required=True), Col("Setting", "number", required=True),
             Col("Catch (g)", "number", required=True), Col("Time (s)", "number", required=True),
             Col("g/min", "number"), Col("Slope in Use", "number"), Col("Error %", "number"),
             Col("Tested By", required=True)], append_only=True),
          CHANGE_LOG, ISSUES, README]),

    Book("Formulation Calc Library.xlsx", "Formulation Data Base", "evidence",
         "Everything read from Tech's SExx Formulation.xls calc workbooks (today's 'Formulation Master.xlsx', renamed "
         "so the approved master can take that name). Rebuilt when the calcs change; never edited",
         "calc/parse_fcal.py -> calc/build_formulation_master.py", "seeding and checking the masters",
         [Sheet("Formula Library", "as built today", [Col("Line"), Col("Formula Code")]),
          Sheet("Current Recipes", "as built today", [Col("Line"), Col("Formula Code"), Col("Feeder")]),
          Sheet("Auger Calibration", "as built today", [Col("Line"), Col("Feeder")]),
          ISSUES, README], status="existing"),

    Book("FRM Draft <date>.xlsx", "Daily Formulation Report", "daily",
         "What the pipeline proposes for the day: every EXT order resolved to a formula and settings, plus Exceptions. "
         "Nothing guessed: an order without an exact match goes to Exceptions for an engineer",
         "daily/resolve.py", "Tech (review and sign), FRM renderer",
         [Sheet("Draft", "One row per order x formula x feeder, as last issued for that order on that line", [
             Col("Line Code", key=True, required=True), Col("Order", key=True, required=True),
             Col("Product Code", required=True), Col("Formula Code", key=True, required=True),
             Col("Variant", "enum", key=True, required=True, values=VARIANTS),
             Col("Extruder", key=True), Col("Feeder", key=True, required=True),
             Col("Material ID", note="blank until the Materials table exists"), Col("Material (as issued)"),
             Col("Set", required=True), Col("Weight %", "number"),
             Col("How Resolved", "enum", required=True,
                 values=("Last issued (same order, same line)", "Product to Formula", "Engineer")),
             Col("Last Issued", "date"), Col("Source"), Col("Formula Row", "int", key=True), Col("Note")]),
          Sheet("Exceptions", "Orders the pipeline would not resolve (HANDOFF §7.4)", [
             Col("Line Code", required=True), Col("Order", required=True), Col("Product Code"),
             Col("Reason", required=True), Col("Suggestion (not used)"), Col("Engineer Decision"), Col("Decided By")]),
          ISSUES, README]),

    Book("FRM Formulation Report <date>.xlsx", "Daily Formulation Report", "daily",
         "The day's formulation as issued (today: transcribed from Tech's page; later: rendered from the signed draft)",
         "daily/build_xlsx.py", "floor, Formulation Report Record",
         [Sheet("Formulations", "as built today", [Col("Line"), Col("Orders"), Col("Formula Code"), Col("Feeder"), Col("Set")]),
          ISSUES, README], status="existing"),

    Book("FRM Formulation <date>.docx", "Daily Formulation Report", "daily",
         "The print-ready formulation for the floor (Word): one page per line, every formulation of each order in run "
         "order, cover with exceptions and the IWPFO055 §5.3 issue block. DRAFT until an engineer completes the "
         "exceptions and Technical signs (James Kuo, 29 Sep 2026: operators use the paper copy)",
         "daily/render_frm.py (from the FRM Draft)", "extrusion operators (printed), Schedule binder (IWPFO055 §5.4)",
         [], status="new"),

    Book("Formulation Report Record.xlsx", "Daily Formulation Report", "record",
         "Every formulation issued, per day, order and feeder. Append-only: the plant's history of what ran",
         "daily/record.py (to build) after Tech signs the day", "traceability, complaints, trends",
         [Sheet("Issued", "One row per issue date x order x feeder", [
             Col("Issue Date", "date", key=True, required=True),
             Col("Line Code", key=True, required=True), Col("Order", key=True, required=True),
             Col("Product Code", required=True), Col("Formula Code", key=True, required=True),
             Col("Variant", key=True, required=True), Col("Extruder", key=True), Col("Feeder", key=True, required=True),
             Col("Formula Row", "int", key=True, required=True,
                 note="1 = first formula listed for the order; one code can be printed twice with different sets"),
             Col("Material ID"), Col("Material (as printed)"), Col("Set", required=True),
             Col("Weight %", "number", note="weight lines only: Set, or the balance to 100 for 'Auto'; blank on auger lines"),
             Col("Note", note="the FRM row's note, as printed"),
             Col("Source", "enum", required=True, values=("Tech FRM", "Pipeline draft", "Engineer")),
             Col("Source Scan", note="scan file + page"),
             Col("Signed By"), Col("Signed At", "datetime")], append_only=True),
          README]),

    Book("EXT Extrusion Schedule <date>.xlsx", "Extrusion Schedule", "daily",
         "The day's extrusion schedule (WPPPOPRC) as printed", "daily/build_xlsx.py",
         "resolve step, Extrusion Production Record",
         [Sheet("Orders", "as built today", [Col("Line"), Col("Order"), Col("Prod Code")]), ISSUES, README],
         status="existing"),

    Book("Extrusion Production Record.xlsx", "Extrusion Schedule", "record",
         "Each order's progress by day, from the EXT schedule (and converting status beside it)",
         "daily/record.py (to build)", "planning, reporting",
         [Sheet("Orders by Day", "One row per schedule date x order", [
             Col("Schedule Date", "date", key=True, required=True),
             Col("Line Code", key=True, required=True), Col("Order", key=True, required=True),
             Col("Product Code", required=True), Col("Total Sheets", "int"), Col("Weight (LBs)", "number"),
             Col("Plts Done (EXT)", "int", note="'NNN PLTS DONE' in the special instructions"),
             Col("Plts Ordered", "int", note="# Plt as printed; the field caps at 999"),
             Col("Special Instructions"),
             Col("Handwritten", note="handwriting on the print (e.g. a corrected pallet count); never read as data"),
             Col("Source Scan", note="scan file + page")], append_only=True),
          README]),

    Book("CNV Converting Schedule <date>.xlsx", "Converting Schedule", "daily",
         "The day's converting schedule as printed", "daily/build_xlsx.py", "Extrusion Production Record",
         [Sheet("Orders", "as built today", [Col("Order"), Col("Product Code")]), ISSUES, README],
         status="existing"),

    Book("Converting Production Record.xlsx", "Converting Schedule", "record",
         "Each converting order's progress by day, from the CNV sheets (James, 28 Sep 2026: a history file for "
         "converting as well)", "daily/record.py (to build)", "planning, reporting, the X OF Y check against EXT",
         [Sheet("Orders by Day", "One row per schedule date x converting line x order row, as printed", [
             Col("Schedule Date", "date", key=True, required=True),
             Col("Converting Line", key=True, required=True, note="e.g. SD31, SD11/SD12, SC31"),
             Col("Order", key=True, required=True),
             Col("Row", "int", key=True, required=True, note="an order can be listed twice on one sheet"),
             Col("Product Code", required=True),
             Col("Extrusion Status (printed)", note="X OF Y as typed on the sheet"),
             Col("Plts Extruded", "int"), Col("Plts Ordered", "int"),
             Col("Semi Size"), Col("Die #"), Col("Die Status"), Col("Plate Status"), Col("Ink Color"),
             Col("Total Sheets", "int"), Col("Pack Code"), Col("# of Plts", "int"), Col("Pc/Plt", "int"),
             Col("Req. Date"), Col("Done Note"),
             Col("Handwritten", note="handwriting on the sheet; never read as data"),
             Col("Source Scan", note="scan file + page")], append_only=True),
          README]),
]

BY_NAME = {b.name: b for b in BOOKS}

# Reference data the templates carry (settled rules, not guesses).
LINES = [  # (code, no, feeder layout, extruders, AC) — HANDOFF §1 table; dosing from auger_rules.DOSING
    ("SE11", 1, "Hopper 1-5", "1", "1"), ("SE12", 2, "Hopper 1-5", "1", "90"), ("SE13", 3, "Hopper 1-5", "1", "90"),
    ("SE21", 4, "Hopper 1-5", "1", ""), ("SE22", 5, "Hopper 1-5", "1", ""), ("SE23", 6, "Hopper 1-5", "1", "1"),
    ("SE24", 7, "A V1-V5, B V1-V4, C V1-V4", "A, B, C", ""), ("SE31", 8, "A V1-V5, B V1-V4, C V1-V4", "A, B, C", ""),
    ("SE32", 9, "A V1-V5, B V1-V4, C V1-V4", "A, B, C", ""), ("SE25", 10, "Hopper 1-5", "1", "1"),
    ("SE42", 12, "V1-V9", "1", ""), ("SE43", 13, "V1-V9", "1", ""),
    ("SE61", 16, "A 1-6, B 1, C 1-6, D 1 (co-ex)", "A, B, C, D", ""),
]


def book_for(filename: str) -> Book | None:
    """The description of a file, daily names matched on their prefix."""
    if filename in BY_NAME:
        return BY_NAME[filename]
    for b in BOOKS:
        if "<date>" in b.name and filename.startswith(b.name.split("<date>")[0]):
            return b
    return None


def destination(filename: str) -> str | None:
    """Subfolder of PUBLISH_DIR a database file belongs in; None = not a database file (goes to DOCS_DIR)."""
    b = book_for(filename)
    return b.folder if b else None


# ---- templates ----------------------------------------------------------------------------------------------
def templates(out_dir: Path) -> list[Path]:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from auger_rules import DOSING, RULES

    hfont = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="1F3864")
    body = Font(name="Arial", size=10)
    out_dir.mkdir(parents=True, exist_ok=True)
    made = []
    for b in BOOKS:
        if b.status != "new" or "<date>" in b.name:
            continue
        wb = Workbook()
        wb.remove(wb.active)
        for s in b.sheets:
            ws = wb.create_sheet(s.name)
            ws.append([c.name for c in s.cols])
            for i, c in enumerate(s.cols, 1):
                ws.cell(1, i).font, ws.cell(1, i).fill = hfont, hfill
                ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(c.name) + 4))
            ws.freeze_panes = "A2"
            if s.name == "Lines":
                for code, no, layout, ext, acv in LINES:
                    ws.append([code, no, DOSING[code], layout, ext, acv, None, "Yes"])
            if s.name == "Hoppers":
                for line, hoppers in RULES.items():
                    for hop, (hw, main, allowed) in hoppers.items():
                        ratio, _, screw = hw.partition(" ")
                        ws.append([line, hop, ratio, screw, main, ", ".join(sorted(allowed)), None])
            if s.name == "Read Me":
                ws.append(["Workbook", b.name])
                ws.append(["Purpose", b.purpose])
                ws.append(["Kind", b.kind])
                ws.append(["Written by", b.written_by])
                ws.append(["Template", "Blank structure from db/schema.py. Reference rows (Lines, Hoppers) are settled "
                                       "rules; the hopper roles are James's draft (HANDOFF §7.14, Q13)."])
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.font = body
        path = out_dir / b.name
        wb.save(path)
        made.append(path)
    return made


# ---- checking a workbook against its description ------------------------------------------------------------
def check(path: Path) -> list[tuple[str, str, str]]:
    """(severity, where, problem). Columns by name; keys unique; required filled; enums allowed."""
    from openpyxl import load_workbook
    b = book_for(path.name)
    if not b:
        return [("High", path.name, "not a database workbook (no description in db/schema.py)")]
    wb = load_workbook(path, read_only=True, data_only=True)
    out = []
    for s in b.sheets:
        if s.name not in wb.sheetnames:
            out.append(("High", s.name, "sheet missing"))
            continue
        rows = wb[s.name].iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows, ())]
        missing = [c.name for c in s.cols if c.name not in header]
        if missing:
            out.append(("High", s.name, f"columns not found by name: {missing}"))
            continue
        idx = {c.name: header.index(c.name) for c in s.cols}
        keys = [c.name for c in s.cols if c.key]
        seen = {}
        for n, row in enumerate(rows, 2):
            if all(v is None or str(v).strip() == "" for v in row):
                continue
            get = lambda name: row[idx[name]] if idx[name] < len(row) else None  # noqa: E731
            for c in s.cols:
                v = get(c.name)
                empty = v is None or str(v).strip() == ""
                if c.required and empty:
                    out.append(("Medium", f"{s.name}!row {n}", f"'{c.name}' is required"))
                if c.values and not empty and str(v).strip() not in c.values:
                    out.append(("Medium", f"{s.name}!row {n}", f"'{c.name}' = {v!r}; allowed: {', '.join(c.values)}"))
            if keys:
                k = tuple(str(get(x) or "").strip() for x in keys)
                if k in seen:
                    out.append(("High", f"{s.name}!row {n}", f"duplicate key {k} (first at row {seen[k]})"))
                else:
                    seen[k] = n
    wb.close()
    return out


# ---- column reference ---------------------------------------------------------------------------------------
def doc() -> str:
    lines = ["# Formulation database: tables", "",
             "Generated by `python db/schema.py doc` from `db/schema.py`. Do not edit by hand.", ""]
    for b in BOOKS:
        lines += [f"## {b.name}", "",
                  f"*{b.kind}* · folder `{b.folder}` · {'exists today' if b.status == 'existing' else 'new'}", "",
                  b.purpose + ".", "", f"- Written by: {b.written_by}", f"- Read by: {b.read_by}", ""]
        for s in b.sheets:
            if s.name in ("Read Me",) or (b.status == "existing" and s.name != "Product Master"):
                continue
            lines += [f"### {s.name}" + (" (append-only)" if s.append_only else ""), "", s.purpose + ".", "",
                      "| Column | Type | Key | Req. | Allowed / refers to | Note |", "|---|---|---|---|---|---|"]
            for c in s.cols:
                allowed = ", ".join(c.values) if c.values else (f"→ {c.ref}" if c.ref else "")
                lines.append(f"| {c.name} | {c.type} | {'●' if c.key else ''} | {'●' if c.required else ''} | "
                             f"{allowed} | {c.note} |")
            lines.append("")
    return "\n".join(lines)


def main(argv):
    import config
    if not argv or argv[0] not in ("templates", "check", "doc"):
        raise SystemExit(__doc__)
    if argv[0] == "templates":
        for p in templates(config.OUTPUT_DIR / "templates"):
            print(p)
    elif argv[0] == "check":
        path = Path(argv[1])
        config.record_read(path, "schema check")
        problems = check(path)
        for sev, where, what in problems:
            print(f"{sev:6s} {where}: {what}")
        print(f"{len(problems)} problem(s)")
        raise SystemExit(1 if any(p[0] == "High" for p in problems) else 0)
    else:
        (ROOT / "docs" / "DATABASE_TABLES.md").write_text(doc() + "\n", encoding="utf-8")
        print(ROOT / "docs" / "DATABASE_TABLES.md")


if __name__ == "__main__":
    main(sys.argv[1:])
