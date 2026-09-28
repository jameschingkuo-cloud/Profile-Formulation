"""Change control for the masters (James Kuo, 28 Sep 2026): people edit the master in Excel in SharePoint; every
edit needs an approved row in its Change Log sheet; this check finds any edit that has none and STOPS.

    python db/preflight.py check  "Formulation Master.xlsx"   compare the published master with the last accepted
                                                              version; exit 1 if any change has no approved log row
    python db/preflight.py accept "Formulation Master.xlsx"   run check; if clean, make the published master the new
                                                              accepted version (data/snapshots/<name>.json, in git)
    python db/preflight.py accept "Formulation Master.xlsx" --baseline   first time only (no snapshot yet)

A change is matched by a Change Log row added since the last accepted version with the same Sheet, Row Key (the
sheet's key columns joined by '|'), Field (column name, or '(new row)' / '(row removed)') and New Value, and a
name in Approved By. The pipeline never edits a master itself: it proposes, people approve and log.
"""
import datetime
import json
import os
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import config  # noqa: E402
from db import schema  # noqa: E402

SNAP_DIR = Path(os.environ.get('SNAP_DIR') or config.DATA_DIR / 'snapshots')   # accepted versions, kept in git
SKIP = {'Change Log', 'Issues', 'Read Me'}
NEW, GONE = '(new row)', '(row removed)'


def norm(v):
    if isinstance(v, datetime.datetime):
        v = v.date() if v.time() == datetime.time() else v
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return '' if v is None else str(v).strip()


def read(path):
    """{sheet: {row key: {column: value}}} for the master sheets, plus the Change Log rows."""
    from openpyxl import load_workbook
    book = schema.book_for(path.name)
    wb = load_workbook(path, read_only=True, data_only=True)
    data, log = {}, []
    for s in book.sheets:
        if s.name not in wb.sheetnames:
            raise SystemExit(f'STOP: sheet "{s.name}" is missing from {path.name}')
        rows = wb[s.name].iter_rows(values_only=True)
        header = [norm(h) for h in next(rows, ())]
        missing = [c.name for c in s.cols if c.name not in header]
        if missing:
            raise SystemExit(f'STOP: {path.name} / {s.name}: columns not found by name: {missing}')
        recs = [{c.name: norm(r[header.index(c.name)]) if header.index(c.name) < len(r) else '' for c in s.cols}
                for r in rows if any(v not in (None, '') for v in r)]
        if s.name == 'Change Log':
            log = recs
        elif s.name not in SKIP:
            keys = [c.name for c in s.cols if c.key] or [s.cols[0].name]
            data[s.name] = {'|'.join(r[k] for k in keys): r for r in recs}
    wb.close()
    return data, log


def diff(old, new):
    out = []
    for sheet in sorted(set(old) | set(new)):
        a, b = old.get(sheet, {}), new.get(sheet, {})
        for k in sorted(set(a) | set(b)):
            if k not in a:
                out.append((sheet, k, NEW, '', ''))
            elif k not in b:
                out.append((sheet, k, GONE, '', ''))
            else:
                for col in b[k]:
                    if a[k].get(col, '') != b[k][col]:
                        out.append((sheet, k, col, a[k].get(col, ''), b[k][col]))
    return out


def snap_path(name):
    return SNAP_DIR / (Path(name).stem + '.json')


def check(name, quiet=False):
    """(ok, report lines, current data). ok = every change since the accepted version has an approved log row."""
    pub = config.published_path(name)
    if not (pub and pub.exists()):
        raise SystemExit(f'{name} is not published at {pub}')
    config.record_read(pub, 'master (pre-flight)')
    cur, log = read(pub)
    sp = snap_path(name)
    if not sp.exists():
        return None, [f'No accepted version of {name} yet: run "accept --baseline" once after the first publish.'], (cur, log)
    snap = json.loads(sp.read_text(encoding='utf-8'))
    last_id = int(snap.get('change_log_max_id') or 0)
    new_log = [r for r in log if r['Change ID'].isdigit() and int(r['Change ID']) > last_id]
    changes = diff(snap['data'], cur)
    used, unlogged, report = set(), [], []
    for sheet, key, field, old, new in changes:
        hit = None
        for i, r in enumerate(new_log):
            if (r['Sheet'], r['Row Key'], r['Field']) == (sheet, key, field) and \
                    (field in (NEW, GONE) or r['New Value'] == new):
                hit = i
                break
        if hit is None or not new_log[hit]['Approved By']:
            why = 'no Change Log row' if hit is None else f"Change {new_log[hit]['Change ID']} has no Approved By"
            unlogged.append(f'  UNLOGGED  {sheet} | {key} | {field}: {old!r} -> {new!r}  ({why})')
        else:
            used.add(hit)
            r = new_log[hit]
            report.append(f"  ok        {sheet} | {key} | {field}: {old!r} -> {new!r}  (Change {r['Change ID']}, approved by {r['Approved By']})")
    for i, r in enumerate(new_log):
        if i not in used:
            report.append(f"  note      Change {r['Change ID']} ({r['Sheet']} | {r['Row Key']} | {r['Field']}) is logged but no such change is in the file")
    head = [f'{name}: {len(changes)} change(s) since the accepted version ({snap.get("accepted_at")}); '
            f'{len(new_log)} new Change Log row(s).']
    if unlogged:
        head.append('STOP - changed without an approved Change Log row (hard rule, James 25 Sep 2026). Nothing is built:')
    return not unlogged, head + unlogged + report, (cur, log)


def accept(name, baseline=False):
    ok, lines, (cur, log) = check(name)
    if ok is None and not baseline:
        print('\n'.join(lines)); raise SystemExit(2)
    if ok is False:
        print('\n'.join(lines)); raise SystemExit(1)
    ids = [int(r['Change ID']) for r in log if r['Change ID'].isdigit()]
    SNAP_DIR.mkdir(parents=True, exist_ok=True)
    snap_path(name).write_text(json.dumps({'workbook': name, 'accepted_at': datetime.datetime.now().isoformat(timespec='seconds'),
                                           'change_log_max_id': max(ids) if ids else 0, 'data': cur},
                                          ensure_ascii=False, indent=0, sort_keys=True), encoding='utf-8')
    print('\n'.join(lines) if ok else f'{name}: baseline taken.')
    print(f'Accepted: {snap_path(name)} (commit it).')


def main(argv):
    names = [a for a in argv[1:] if not a.startswith('--')]
    if not argv or argv[0] not in ('check', 'accept') or len(names) != 1:
        raise SystemExit(__doc__)
    if argv[0] == 'check':
        ok, lines, _ = check(names[0])
        print('\n'.join(lines))
        raise SystemExit(0 if ok else (2 if ok is None else 1))
    accept(names[0], baseline='--baseline' in argv)


if __name__ == '__main__':
    main(sys.argv[1:])
