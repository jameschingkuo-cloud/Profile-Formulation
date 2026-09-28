"""Formulation Master: built from Tech's per-line formulation calculation workbooks (SExx Formulation.xls).
Each calc block = one order's formulation: hoppers (gear ratio, screw), material, auger calibration slope
(g/min per setting unit), formula setting, g/min, and formulation %.
Usage: python calc/build_formulation_master.py   (reads work/parsed.pkl from parse_fcal.py, the daily packet JSON
(PKT_DATE, else the latest in data/packets) and the Production Formula Item log in CALC_DIR)"""
import sys as _sys, pathlib as _pl; _sys.path.insert(0, str(_pl.Path(__file__).resolve().parents[1])); import config  # repo root config
import pickle, re, json, datetime, collections, glob, xlrd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from common import split_codes, FORMULA_RE, mat_key, display_name, superseded_copies

import os
OUT = config.OUTPUT_DIR / 'Formulation Master.xlsx'
PACKET = os.environ.get('PKT_JSON') or (str(config.packet_path(os.environ['PKT_DATE'])) if os.environ.get('PKT_DATE') else str(config.latest_packet()))
config.record_read(PACKET, 'daily packet (FRM vs Calc)')
_items = sorted(p for p in config.CALC_DIR.glob('*Formula*Item*.xls') if not p.name.startswith('~$'))
ITEMLOG = str(_items[-1]) if _items else None
if ITEMLOG: config.record_read(ITEMLOG, 'Production Formula Item log')
F = 'Arial'
HDR = PatternFill('solid', fgColor='1F3864'); HFONT = Font(name=F, bold=True, color='FFFFFF', size=10)
BODY = Font(name=F, size=10); BOLD = Font(name=F, size=10, bold=True)
FLAG = PatternFill('solid', fgColor='FCE4D6'); OK = PatternFill('solid', fgColor='E2EFDA'); CUR = PatternFill('solid', fgColor='DDEBF7')
SEV = {'High': 'F8CBAD', 'Medium': 'FFE699', 'Low': 'DDEBF7', 'Info': 'EDEDED'}

B, R = pickle.load(open(config.WORK_DIR / 'parsed.pkl', 'rb'))
DROPPED = superseded_copies(B)          # older 'Copy of ...' workbooks of a line that has a newer one (common.superseded_copies)
B = [b for b in B if b['file'] not in DROPPED]
R = [r for r in R if r['file'] not in DROPPED]
SRC = {display_name(b['file']) for b in B}

def copy_no(sheet):
    m = re.search(r'\((\d+)\)\s*$', sheet); return int(m.group(1)) if m else 0
for b in B:
    b['date'] = b['period_start'] or b['order_date']
    b['date_src'] = 'Prod. Period' if b['period_start'] else ('order no.' if b['order_date'] else '')
    b['sortkey'] = (b['date'] or datetime.date(1900, 1, 1), copy_no(b['sheet']), b['row'])
    b['order_list'] = split_codes(b['orders'], 'order')
    b['product_list'] = split_codes(b.get('product'))
    b['fcode'] = (b.get('formula_code') or '').strip()
    b['id'] = f"{b['line']}|{b['sheet']}|{b['row']}"
# a hopper header overwritten by a material name (e.g. '6502A' in the H1 header): take gear ratio / screw from that line's usual H1 header
modal = collections.defaultdict(collections.Counter)
for r in R:
    if r['gear_ratio']: modal[(r['line'], r['feeder'])][(r['gear_ratio'], r['screw'])] += 1
for r in R:
    if r.get('bad_header') and not r['gear_ratio'] and modal.get((r['line'], r['feeder'])):
        r['gear_ratio'], r['screw'] = modal[(r['line'], r['feeder'])].most_common(1)[0][0]
        r['comment'] = (r['comment'] + '; ' if r['comment'] else '') + f"hopper header read '{r['bad_header']}' - ratio/screw taken from usual {r['feeder']} header"
rows_by_block = collections.defaultdict(list)
for r in R:
    rows_by_block[f"{r['line']}|{r['sheet']}|{r['row']}"].append(r)
BLK = {b['id']: b for b in B}

def fnum(v, nd=1):
    if v is None: return ''
    return f"{v:.{nd}f}".rstrip('0').rstrip('.') if isinstance(v, float) else str(v)
def recipe(bid):
    """settings string (what the floor sets) and composition string (%)"""
    rs = [r for r in rows_by_block[bid] if not r['group'].startswith('Total')]
    set_parts, pct_parts = [], []
    for r in rs:
        if r['setting'] in (None, 0) and not r['pct']: continue          # hopper listed but not used
        g = f"{r['group']}:" if r['group'] and r['group'] not in ('', 'A') or any(x['group'] for x in rs) else ''
        val = r['setting'] if r['setting'] is not None else (f"{fnum(r['pct'])}%" if r['pct'] is not None else '')
        set_parts.append(f"{g}{r['feeder']} {r['material']} {fnum(val) if not isinstance(val, str) else val}".strip())
    tot = [r for r in rows_by_block[bid] if r['group'].startswith('Total') and r['material']]
    src = tot if tot else rs
    comp = collections.OrderedDict()
    for r in src:
        if r['material'] and r['pct']: comp[r['material']] = comp.get(r['material'], 0) + r['pct']
    pct_parts = [f"{m} {fnum(p)}%" for m, p in comp.items()]
    return ' | '.join(set_parts), ' | '.join(pct_parts), sum(comp.values())
def signature(bid):
    return tuple((r['group'], r['feeder'], r['material'].upper().replace(' ', ''), round(r['setting'], 2) if r['setting'] is not None else None,
                  round(r['pct'], 1) if (r['setting'] is None and r['pct'] is not None) else None)
                 for r in rows_by_block[bid] if not r['group'].startswith('Total') and (r['material'] or r['setting'] is not None))

# ---------------------------------------------------------------- Formula library (line, code, recipe variant)
lib = collections.OrderedDict()
for b in sorted(B, key=lambda b: (b['line'], b['fcode'], b['sortkey'])):
    k = (b['line'], b['fcode'], signature(b['id']))
    lib.setdefault(k, []).append(b)
latest_by_code = {}
for (line, code, sig), bl in lib.items():
    last = bl[-1]
    if (line, code) not in latest_by_code or last['sortkey'] > latest_by_code[(line, code)]['sortkey']:
        latest_by_code[(line, code)] = last
lib_rows = []
for (line, code, sig), bl in lib.items():
    last = bl[-1]; first = bl[0]
    rec, comp, tot = recipe(last['id'])
    is_cur = latest_by_code[(line, code)] is last
    prods = list(dict.fromkeys(p for b in reversed(bl) for p in b['product_list']))
    lib_rows.append([line, code, 'Current' if is_cur else 'Earlier', rec, comp, round(tot, 1) if tot else None, len(bl),
                     first['date'], last['date'], last['date_src'], ', '.join(last['order_list'][:6]) or last['orders'],
                     ', '.join(prods[:8]) + (f' (+{len(prods) - 8})' if len(prods) > 8 else ''),
                     last.get('thk'), last.get('gsm'), last.get('grade', ''), last.get('application', ''),
                     '' if FORMULA_RE.match(code) else ('no formula code' if not code else 'code pattern'),
                     f"{display_name(last['file'])} / {last['sheet']} / row {last['row']}"])
lib_rows.sort(key=lambda x: (x[0], x[1], x[2] != 'Current', -(x[8].toordinal() if x[8] else 0)))

# ---------------------------------------------------------------- current recipes, long (one row per feeder)
cur_long = []
for (line, code), b in sorted(latest_by_code.items()):
    for r in rows_by_block[b['id']]:
        if not (r['material'] or r['setting'] is not None or r['pct'] is not None): continue
        cur_long.append([line, code, b['date'], ', '.join(b['order_list'][:4]) or b['orders'], r['group'], r['feeder'], r['gear_ratio'], r['screw'],
                         r['material'], mat_key(r['material']), r['slope'], r['setting'], r['amount_gmin'], r['pct'], r['extruder_share'], r['comment']])

# ---------------------------------------------------------------- calibration
cal = collections.defaultdict(list)
for r in R:
    if r['slope'] is None or not r['material']: continue
    b = BLK.get(f"{r['line']}|{r['sheet']}|{r['row']}")
    cal[(r['line'], r['feeder'], r['gear_ratio'], r['screw'], r['material'])].append((b['sortkey'], round(r['slope'], 4), b))
cal_cur, cal_hist = [], []
for k, lst in sorted(cal.items()):
    lst.sort(key=lambda x: x[0])
    byslope = collections.OrderedDict()
    for sk, sl, b in lst:
        byslope.setdefault(sl, []).append((sk, b))
    cur_slope = lst[-1][1]
    for sl, occ in byslope.items():
        cal_hist.append(list(k) + [sl, len(occ), occ[0][1]['date'], occ[-1][1]['date'], 'Current' if sl == cur_slope else 'Earlier',
                                   f"{occ[-1][1]['sheet']} row {occ[-1][1]['row']}"])
    occ = byslope[cur_slope]
    earlier = [f"{s} (last {fmt.isoformat() if fmt else '?'})" for s, o in byslope.items() if s != cur_slope for fmt in [o[-1][1]['date']]]
    cal_cur.append(list(k) + [mat_key(k[4]), cur_slope, len(lst), occ[0][1]['date'], lst[-1][2]['date'], len(byslope), '; '.join(earlier)])

# ---------------------------------------------------------------- product -> formula
pf = collections.defaultdict(list)
for b in B:
    for p in b['product_list']:
        pf[(p, b['line'], b['fcode'])].append(b)
pf_rows = []
for (p, line, code), bl in sorted(pf.items()):
    bl.sort(key=lambda b: b['sortkey'])
    last = bl[-1]
    pf_rows.append([p, line, code, len(bl), bl[0]['date'], last['date'], ', '.join(last['order_list'][:4]) or last['orders'], last.get('thk'), last.get('gsm'),
                    last.get('grade', ''), last.get('size', ''), last.get('application', '')])

# ---------------------------------------------------------------- calc history (long)
hist = []
for b in sorted(B, key=lambda b: (b['line'], b['sortkey'])):
    for r in rows_by_block[b['id']]:
        if not (r['material'] or r['setting'] is not None or r['pct'] is not None): continue
        hist.append([b['line'], b['date'], b['date_src'], b['orders'], b.get('product', ''), b['fcode'], b.get('thk'), b.get('gsm'), b.get('lsp'),
                     b.get('output_lbhr'), b.get('grade', ''), b.get('size', ''), b.get('application', ''), r['group'], r['feeder'], r['gear_ratio'],
                     r['screw'], r['material'], r['slope'], r['setting'], r['amount_gmin'], r['pct'], r['extruder_share'], r['comment'],
                     display_name(b['file']), b['sheet'], b['row']])

# ---------------------------------------------------------------- 24 Sep FRM vs calc
pk = json.load(open(PACKET, encoding='utf-8'))
_pd = datetime.date.fromisoformat(pk.get('packet_date') or re.search(r'(\d{4}-\d\d-\d\d)', PACKET).group(1))
LABEL = f'{_pd.day} {_pd:%b}'   # e.g. '24 Sep' (no %-d: it fails on Windows)
calc_by_order = collections.defaultdict(list)
for b in B:
    for o in b['order_list']: calc_by_order[(b['line'], o)].append(b)
FRM_FEED = lambda name: re.sub(r'^Hopper\s*', 'H', name).replace('Extruder ', '').replace(' ', '')
frm_rows = []
def vals_of(bid):
    crs = [r for r in rows_by_block[bid] if not r['group'].startswith('Total')]
    cset = sorted(r['setting'] for r in crs if r['setting'] not in (None, 0))
    cpct = sorted(round(r['pct'], 1) for r in crs if r['pct'] not in (None, 0))
    return cset, cpct
def ndiff(a, b):
    ca, cb = collections.Counter(a), collections.Counter(b)
    return sum(((ca - cb) + (cb - ca)).values())
for p in pk['frm']:
    line = p['line_code']
    for g in p['groups']:
        for fi, f in enumerate(g['formulas']):
            fset = [(k, v['material'], v['set']) for k, v in f['feeders'].items() if v['material'] or v['set']]
            fv = sorted(float(s) for _, _, s in fset if re.fullmatch(r'\d+(\.\d+)?', str(s).strip()) and float(s) != 0)
            ftxt = '; '.join(f"{k} {m} {s}" for k, m, s in fset)
            role = 'Primary' if fi == 0 else f'Alternate {fi}'
            for o in g['orders']:
                cands = calc_by_order.get((line, o), [])
                if not cands:
                    frm_rows.append([line, o, f['formula_code'], role, ftxt, '', '', None, '', 'No calc block for this order', len(cands), f"FRM p{p['scan_page']}"]); continue
                scored = []
                for b in cands:
                    cset, cpct = vals_of(b['id'])
                    d_set, d_pct = ndiff(fv, cset) if cset else 99, ndiff([round(v, 1) for v in fv], cpct) if cpct else 99
                    scored.append((min(d_set, d_pct), b['fcode'] != f['formula_code'], tuple(-x for x in (b['sortkey'][0].toordinal(), b['sortkey'][1], b['sortkey'][2])), b, d_set <= d_pct))
                scored.sort(key=lambda x: (x[0], x[1], x[2]))
                d, codediff, _, b, by_set = scored[0]
                if d == 0:
                    res = ('Same settings' if by_set else 'Same settings (calc %)') + (', different formula code' if codediff else ', same code')
                elif fi > 0:
                    res = 'Alternate not in the calc workbook (closest calc shown)'
                else:
                    res = f"Settings differ ({d} value{'s' if d > 1 else ''})" + (', different formula code' if codediff else '')
                frm_rows.append([line, o, f['formula_code'], role, ftxt, b['fcode'], recipe(b['id'])[0], b['date'],
                                 f"{display_name(b['file'])} / {b['sheet']} / row {b['row']}", res, len(cands), f"FRM p{p['scan_page']}"])

# ---------------------------------------------------------------- Production Formula Item log
wbI = xlrd.open_workbook(ITEMLOG) if ITEMLOG else None     # no item log in CALC_DIR -> the Formula Item Log sheet is empty
item_rows = []
for s in (wbI.sheets() if wbI else []):
    cur_date, cur_shift = None, ''
    for r in range(1, s.nrows):
        v = s.row_values(r)
        if not any(str(x).strip() for x in v): continue
        if v[0] != '':
            if isinstance(v[0], float):
                cur_date = xlrd.xldate_as_datetime(v[0], wbI.datemode).date()
            else:
                txt = str(v[0]).strip()
                m = re.fullmatch(r'(\d{1,2})/(\d{1,2})(\d{4})', txt)          # '6/102026' typed without the second slash
                cur_date = datetime.date(int(m.group(3)), int(m.group(1)), int(m.group(2))) if m else (cur_date if txt.lower().startswith('to') else txt)
        if v[1] != '': cur_shift = str(v[1]).strip()
        item_rows.append([s.name, cur_date, cur_shift, str(v[2]).strip(), str(v[3]).strip(), v[4] if v[4] != '' else None, str(v[5]).strip()])

# ---------------------------------------------------------------- issues
issues = []
dmax_all = max(b['date'] for b in B if b['date'])
def add(sev, line, where, check, detail): issues.append([sev, line, where, check, detail])
for code_rows in [r for r in lib_rows if r[16]]:
    pass
bad_codes = collections.Counter((r[0], r[1]) for r in lib_rows if r[16] == 'code pattern')
for (line, code), n in sorted(bad_codes.items()):
    nb = sum(len(v) for (l, c, s), v in lib.items() if l == line and c == code)
    add('Medium', line, code, 'Formula code does not fit the usual pattern',
        f'{code} used in {nb} calc block(s). Usual pattern: 2 letters + A/0/D + 3 digits + 2-letter colour + thickness (e.g. FUA152WB4); premix codes use a 3-letter colour (FU001GTW4).')
nocode = [b for b in B if not b['fcode']]
if nocode:
    add('Low', '(several)', f'{len(nocode)} blocks', 'Calc block without a formula code',
        '; '.join(sorted({f"{b['line']} {b['sheet']}" for b in nocode}))[:600])
for b in B:
    for r in rows_by_block[b['id']]:
        if r['slope'] is not None and r['setting'] is not None and r['amount_gmin'] is not None and abs(r['slope'] * r['setting'] - r['amount_gmin']) > 0.05 + 0.001 * abs(r['amount_gmin']):
            add('Medium', b['line'], f"{b['sheet']} row {b['row']}", 'g/min is not slope x setting',
                f"{r['feeder']} {r['material']}: {r['slope']} x {r['setting']} = {r['slope'] * r['setting']:.1f} but sheet shows {r['amount_gmin']:.1f}")
for (line, code), b in latest_by_code.items():
    _, _, tot = recipe(b['id'])
    if tot and abs(tot - 100) > 0.6:
        add('Low', line, f"{code} ({b['sheet']} row {b['row']})", 'Current formulation % does not add to 100', f'adds to {tot:.1f}%')
# materials that are really dies / notes
for b in B:
    for r in rows_by_block[b['id']]:
        if re.fullmatch(r'\d+(\.\d+)?', r['material'] or ''):
            add('Low', b['line'], f"{b['sheet']} row {b['row']}", 'Number in the Material cell', f"{r['feeder']}: '{r['material']}'")
for k, lst in sorted(cal.items()):             # more than one slope used for the same hopper + material in the last 12 months
    recent = [x for x in lst if x[2]['date'] and (dmax_all - x[2]['date']).days <= 365]
    sl = collections.OrderedDict()
    for sk, v, b in sorted(recent, key=lambda x: x[0]): sl.setdefault(v, []).append(b)
    if len(sl) > 1:
        add('Medium', k[0], f"{k[1]} ({k[2]}, {k[3]}) {k[4]}", 'Two calibration slopes in use for one hopper + material',
            '; '.join(f"{v} in {len(bs)} block(s), last {bs[-1]['date']} ({bs[-1]['sheet']})" for v, bs in sl.items()) +
            '. Sheets copied from older sheets may carry an old slope, so the % on those sheets is off.')
for r in frm_rows:
    if not r[9].endswith('same code'):
        sev = 'Medium' if r[9].startswith('Settings differ') else ('Low' if 'different formula code' in r[9] else 'Info')
        if r[9].startswith('Alternate'): continue
        add(sev, r[0], f'{r[1]} {r[2]} ({r[3]})', f'FRM {LABEL} vs calc workbook', f"{r[9]}. FRM: {r[4]} || closest calc {r[5]}: {r[6]}")
issues.sort(key=lambda x: (['High', 'Medium', 'Low', 'Info'].index(x[0]), x[1], x[2]))

# ================================================================ workbook
wb = Workbook()
def sheet(name, headers, rows, widths=None, first=False, table=True, datecols=(), numfmt=None):
    ws = wb.active if first else wb.create_sheet(name)
    ws.title = name
    ws.append(headers)
    for r in rows: ws.append(r)
    for c in ws[1]:
        c.font = HFONT; c.fill = HDR; c.alignment = Alignment(wrap_text=True, vertical='center')
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, min(max(len(h) + 2, 9), 30))
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'B2' if len(rows) else 'A2'
    big = len(rows) > 5000
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY
            if not big: c.alignment = Alignment(vertical='top', wrap_text=isinstance(c.value, str) and len(c.value) > 45)
    for h in datecols:
        j = headers.index(h) + 1
        for r in range(2, len(rows) + 2): ws.cell(r, j).number_format = 'yyyy-mm-dd'
    for h, fmt in (numfmt or {}).items():
        j = headers.index(h) + 1
        for r in range(2, len(rows) + 2): ws.cell(r, j).number_format = fmt
    if table and rows:
        t = Table(displayName=re.sub(r'\W', '', name)[:28] + 'T', ref=f"A1:{get_column_letter(len(headers))}{len(rows) + 1}")
        t.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True); ws.add_table(t)
    return ws

H_LIB = ['Line', 'Formula Code', 'Status', 'Recipe (feeder material setting)', 'Composition (%)', 'Total %', 'Calc Blocks', 'First Run', 'Last Run',
         'Date From', 'Last Orders', 'Products', 'Thk', 'GSM', 'Grade', 'Application / Note', 'Code Check', 'Source (last block)']
ws = sheet('Formula Library', H_LIB, lib_rows, {'Recipe (feeder material setting)': 70, 'Composition (%)': 60, 'Last Orders': 28, 'Products': 36,
                                                'Application / Note': 24, 'Source (last block)': 40, 'Formula Code': 13}, first=True,
           datecols=('First Run', 'Last Run'))
for r in range(2, len(lib_rows) + 2):
    if ws.cell(r, 3).value == 'Current': ws.cell(r, 3).fill = CUR
    if ws.cell(r, 17).value: ws.cell(r, 17).fill = FLAG
H_CUR = ['Line', 'Formula Code', 'Last Run', 'Last Orders', 'Extruder', 'Feeder', 'Gear Ratio', 'Screw', 'Material (as written)', 'Material Class',
         'Calibration Slope (g/min per unit)', 'Setting', 'g/min', 'Formulation %', 'Extruder Share', 'Comment']
sheet('Current Recipes', H_CUR, cur_long, {'Material (as written)': 24, 'Last Orders': 26, 'Comment': 24}, datecols=('Last Run',),
      numfmt={'Calibration Slope (g/min per unit)': '0.0000', 'g/min': '#,##0.0', 'Formulation %': '0.0'})
H_CAL = ['Line', 'Feeder', 'Gear Ratio', 'Screw', 'Material (as written)', 'Material Class', 'Current Slope (g/min per unit)', 'Calc Blocks',
         'Current Slope Since', 'Last Used', 'Slopes Seen', 'Earlier Slopes']
ws = sheet('Auger Calibration', H_CAL, cal_cur, {'Material (as written)': 24, 'Earlier Slopes': 50}, datecols=('Current Slope Since', 'Last Used'),
           numfmt={'Current Slope (g/min per unit)': '0.0000'})
for r in range(2, len(cal_cur) + 2):
    if (ws.cell(r, 11).value or 0) > 1: ws.cell(r, 12).fill = FLAG
H_CH = ['Line', 'Feeder', 'Gear Ratio', 'Screw', 'Material (as written)', 'Slope (g/min per unit)', 'Calc Blocks', 'First Used', 'Last Used', 'Status', 'Last Block']
sheet('Calibration History', H_CH, cal_hist, {'Material (as written)': 24, 'Last Block': 30}, datecols=('First Used', 'Last Used'),
      numfmt={'Slope (g/min per unit)': '0.0000'})
H_PF = ['Product Code', 'Line', 'Formula Code', 'Calc Blocks', 'First Run', 'Last Run', 'Last Orders', 'Thk', 'GSM', 'Grade', 'Size', 'Application / Note']
sheet('Product to Formula', H_PF, pf_rows, {'Last Orders': 26, 'Size': 20, 'Application / Note': 24}, datecols=('First Run', 'Last Run'))
H_FRM = ['Line', 'Order', 'FRM Formula Code', 'FRM Row', f'FRM Settings ({LABEL})', 'Closest Calc Formula Code', 'Calc Recipe', 'Calc Run Date', 'Calc Block',
         'Result', 'Calc Blocks for Order', 'FRM Source']
ws = sheet(f'FRM {LABEL} vs Calc', H_FRM, frm_rows, {f'FRM Settings ({LABEL})': 60, 'Calc Recipe': 60, 'Calc Block': 36, 'Result': 30},
           datecols=('Calc Run Date',))
for r in range(2, len(frm_rows) + 2):
    v = ws.cell(r, 10).value or ''
    ws.cell(r, 10).fill = OK if v.endswith('same code') else FLAG
H_IT = ['Sheet', 'Date', 'Shift', 'Line', 'Order', 'Formula Item #', 'Comments']
sheet('Formula Item Log', H_IT, item_rows, {'Comments': 24}, datecols=('Date',))
H_IS = ['Severity', 'Line', 'Where', 'Check', 'Detail', 'James / Tech response']
ws = sheet('Issues', H_IS, [i + [''] for i in issues], {'Where': 30, 'Check': 34, 'Detail': 90, 'James / Tech response': 30})
for r in range(2, len(issues) + 2): ws.cell(r, 1).fill = PatternFill('solid', fgColor=SEV[ws.cell(r, 1).value])
H_H = ['Line', 'Run Date (est.)', 'Date From', 'Orders', 'Product(s)', 'Formula Code', 'Thk', 'GSM', 'LSP (m/min)', 'Output (lb/hr)', 'Grade', 'Size',
       'Application / Note', 'Extruder', 'Feeder', 'Gear Ratio', 'Screw', 'Material', 'Slope', 'Setting', 'g/min', 'Formulation %', 'Extruder Share',
       'Comment', 'Workbook', 'Sheet', 'Row']
sheet('Calc History', H_H, hist, {'Orders': 26, 'Product(s)': 26, 'Material': 22, 'Workbook': 26, 'Sheet': 22}, datecols=('Run Date (est.)',))

# Read Me with live counts
rm = wb.create_sheet('Read Me', 0)
dmin = min(b['date'] for b in B if b['date'] and b['date'].year > 2008); dmax = max(b['date'] for b in B if b['date'])
lines = ['Formulation Master - built from Tech\'s formulation calculation workbooks',
         '',
         f'Source: {len(SRC)} calc workbooks read from {config.CALC_DIR} on {datetime.date.today():%d %b %Y} (' + ', '.join(sorted(SRC)) + ').'
         + (f' Left out as older copies of a line that has a newer workbook: {", ".join(sorted(display_name(f) for f in DROPPED))}.' if DROPPED else '')
         + ' Lines 7, 12, 13 and 16 are weight blenders (no slopes); there is no SE42 workbook.',
         f'Each workbook has one sheet per product type; each sheet holds calc blocks, one per order: product code, formula code, Thk, GSM, line speed, hoppers (gear ratio, screw size), material, auger calibration slope, formula setting, g/min and formulation %. {len(B):,} blocks were read, {dmin:%b %Y} to {dmax:%b %Y}.',
         'How the numbers relate (checked on every block): g/min = calibration slope x formula setting; formulation % = g/min / total g/min. The "Set" printed on the daily FRM page is the formula setting (auger dial), not a percentage, on the hopper lines; on SE24 the FRM Set is the % itself.',
         'Run Date (est.) = the Prod. Period start written on the block where there is one; otherwise the date read from the order number (RP26811 = 2026-08-11; H69A... = Sep 2026; older H + letter year, K = 2009 ... W = 2021). The order-number reading is inferred from the data: it agreed with the written Prod. Period on all but 29 of 1,700 blocks.',
         '',
         'Sheets:',
         '  Formula Library - one row per line + formula code + recipe. "Current" = the most recent calc for that code on that line; "Earlier" = recipes it had before.',
         '  Current Recipes - the current recipe of every line + formula code, one row per feeder (material, calibration slope, setting, g/min, %). This is the table the formulation automation reads.',
         '  Auger Calibration - current slope (g/min per setting unit) per line, feeder (gear ratio, screw) and material, with any earlier slopes. Calibration History lists every slope value and when it was used.',
         '  Product to Formula - every product code in the calcs with the line and formula code it ran with, and when.',
         f'  FRM {LABEL} vs Calc - each formula row on the {LABEL} daily formulation report compared with the calc block for the same order.',
         '  Formula Item Log - the "Production Formula Item" sheet (date, shift, line, order, formula item #).',
         '  Issues - codes that break the pattern, calc arithmetic errors, FRM/calc differences.',
         '  Calc History - every block x feeder as read (the raw extract; filter it).',
         '',
         'Material names are kept as written in the calc (e.g. "Talc MB-TL460", "CaCO3MB-CA410", "W26038A"). Material Class is a broad grouping used only for comparisons; it is not a master name.',
         '',
         'Summary']
for i, l in enumerate(lines, 1):
    rm.cell(i, 1, l).font = Font(name=F, size=12 if i == 1 else 10, bold=i in (1, len(lines)))
    rm.cell(i, 1).alignment = Alignment(wrap_text=True, vertical='top')
n0 = len(lines) + 1
summ = [('Calc History rows (block x feeder)', f"=COUNTA('Calc History'!A:A)-1", ''),
        ('Formula codes (line + code)', f"=COUNTIFS('Formula Library'!C:C,\"Current\")", ''),
        ('Recipe variants', f"=COUNTA('Formula Library'!A:A)-1", ''),
        ('Products with a formula', f"=SUMPRODUCT(1/COUNTIF('Product to Formula'!A2:A{len(pf_rows) + 1},'Product to Formula'!A2:A{len(pf_rows) + 1}))", ''),
        ('Calibrations (line + feeder + material)', f"=COUNTA('Auger Calibration'!A:A)-1", ''),
        ('Calibrations that changed over time', f"=COUNTIF('Auger Calibration'!K:K,\">1\")", ''),
        (f'FRM {LABEL} rows: same settings, same code', f"=COUNTIF('FRM {LABEL} vs Calc'!J:J,\"*same code\")", ''),
        (f'FRM {LABEL} rows: same settings, different code', f"=COUNTIF('FRM {LABEL} vs Calc'!J:J,\"Same settings*different formula code\")", ''),
        (f'FRM {LABEL} rows: settings differ', f"=COUNTIF('FRM {LABEL} vs Calc'!J:J,\"Settings differ*\")", ''),
        (f'FRM {LABEL} rows checked', f"=COUNTA('FRM {LABEL} vs Calc'!A:A)-1", ''),
        ('Issues', f"=COUNTA(Issues!A:A)-1", '')]
for k, (a, f, _) in enumerate(summ):
    rm.cell(n0 + k, 1, a).font = BODY; c = rm.cell(n0 + k, 2, f); c.font = BODY; c.number_format = '#,##0'
rm.column_dimensions['A'].width = 130; rm.column_dimensions['B'].width = 12
wb.save(OUT)
print('saved', len(B), 'blocks;', len(lib_rows), 'library rows;', len(latest_by_code), 'line+codes;', len(cal_cur), 'calibrations;', len(pf_rows),
      'product-formula;', len(frm_rows), 'FRM rows;', len(issues), 'issues;', len(item_rows), 'item log rows;', len(hist), 'history rows')
print(collections.Counter(r[9] for r in frm_rows))
