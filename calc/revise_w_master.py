"""W% Master Formulation, revision 2: James Kuo's notes of 23 Oct 2026 applied to the 30 Sep / 1 Oct draft
(James, 23 Oct 2026: "keep this excel in your MD file and produce a revision").

    python calc/revise_w_master.py            -> out/W% Master Formulation (draft v2).xlsx

Reads data/w_master/W Master Formulation draft (version 1) - James notes 2026-10-23.xlsx (his file, kept as received)
and applies the Notes column, nothing else:
  - Main / Secondary rows stay, under the name James gives ("change name to", "replace name with", "change to");
    the old name is kept as "Printed as" so Tech's pages still match;
  - "replace by / with ..." rows fold into their Main or Secondary (named in the note, or the one of the same colour,
    thickness, prefix and grade with the same words); they are listed on "Replaced" with the composition they had;
  - archive / remove / obsolete / "???" rows go to "Archived", and so does a VOIDFORM row with reclaim (R10);
  - what could not be resolved goes to "Open for James", with the questions still open from 23 Oct.
Compositions are James's rounded values. The rules (R1-R10) are only noted here; docs/FORMULA_ASSIGNMENT_RULES.md has them.
"""
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT)); sys.path.insert(0, str(ROOT / 'daily'))
import config  # noqa: E402
import assign  # noqa: E402

SRC = assign.W_MASTER
OUT = config.OUTPUT_DIR / 'W% Master Formulation (draft v2).xlsx'
W_COLS = assign.W_COLS
RUN_COLS = ['Run: Run Date', 'Run: Line', 'Run: Order', 'Run: Prod Code', 'Run: Grade', 'Run: Spec', 'Run: Colors', 'Run: Thk (mm)',
            'Run: GSM', 'Run: Special Instructions']
HEAD = PatternFill('solid', fgColor='1F3864'); HF = Font(color='FFFFFF', bold=True); WRAP = Alignment(wrap_text=True, vertical='top')
GREEN = PatternFill('solid', fgColor='E2EFDA'); AMBER = PatternFill('solid', fgColor='FFF2CC')
# the indicator each prefix stands for (James, 23 Oct 2026: "capture all the indicators like OP, R1, A1 etc"): the Spec the
# production schedule prints on the order (R1R1R1 ...), as docs/FORMULA_ASSIGNMENT_RULES.md R3
INDICATOR = {'FU': 'R1/R2/R3 (standard; R3 triple wall)', 'RU': 'R4 (reclaim)', 'FX': 'OP (extra opaque)', 'FS': 'S1 (ultra smooth)',
             'FA': 'A1 (anti-static)', 'FE': 'U1 (extra UV)', 'FR': 'F1/F4 (flame retardant)', 'BF': 'R6/RD/RM (bubble line)',
             'HU': 'HDPE', 'FH': 'U1 (extra UV, old letter)'}


def words(note):
    """James's words for what the formulation is for, in his order."""
    n = (note or '').lower()
    out = []
    for w, label in (('graphic', 'graphic grade'), ('packag', 'packaging grade'), ('packing', 'packaging grade'), ('void', 'VOIDFORM'),
                     ('corn', 'corn box'), ('sign blank', 'sign blank'), ('spa', 'sign blank (SPA)'), ('summer', 'summer'), ('winter', 'winter'),
                     ('r4', 'spec R4'), ('r2', 'spec R2'), ('r3', 'spec R3'), ('opaque', 'extra opaque (OP)'), ('ultra smooth', 'ultra smooth (S1)'),
                     ('extra uv', 'extra UV (U1)'), ('anti', 'anti-static'), ('flam', 'flame retardant'), ('hdpe', 'HDPE'), ('pe f', 'PE'),
                     ('roll', 'roll'), ('orbis', 'customer Orbis'), ('carter', 'customer Carter'), ('tzu', 'customer Tzu-Chi'),
                     ('vistech', 'customer Vistech'), ('nissan', 'customer Nissan'), ('se43', 'SE43'), ('6.3', '6.3 mm'), ('4.6', '4.6 mm triple wall'),
                     ('ltl', 'LTL wrap'), ('batten', 'roof batten'), ('cabbage', 'cabbage box'), ('bb line', 'bubble line'), ('customer specific', 'customer specific')):
        if w in n and label not in out:
            out.append(label)
    return ', '.join(out)


def main():
    config.record_read(SRC, "W% Master with James's notes (revision 2)")
    ws = load_workbook(SRC, read_only=True, data_only=True)['Formulation + Latest Run']
    rows = list(ws.iter_rows(values_only=True))
    R = assign.apply_later_notes([dict(zip(rows[0], r)) for r in rows[1:] if r[0]])
    for i, r in enumerate(R):
        r['_i'] = i; r['_level'] = assign.level(r['Notes']); r['_final'] = assign.final_code(r)
    keep = [r for r in R if r['_level'] in ('main', 'secondary')]
    # R10 (James, 23 Oct 2026): "void form should not use any reclaim" - "There are some old void form formula that have
    # reclaim. I don't want them": a VOIDFORM row with reclaim leaves the Master for "Archived"
    void_rc = [r for r in keep if 'void' in assign.quals(r['Notes']) and (r.get('Reclaim %') or 0) >= assign.RECLAIM_MIN]
    # fold "replace by" rows into their target: daily/assign.py replace_target, the same logic the daily run uses
    mains = assign.rules()['mains']
    by_main = {}                      # assign's main dict -> our row
    for m in mains:
        for r in keep:
            if r['_final'] == m['code'] and r['Application'] == m['application'] and r['Notes'] == m['note'] and id(m) not in by_main:
                by_main[id(m)] = r
    target, same = {}, {}
    for r in R:
        if r['_level'] != 'replace':
            continue
        m = assign.replace_target(r, mains)
        if m is not None and id(m) in by_main:
            target[r['_i']] = by_main[id(m)]
            same[r['_i']] = assign.same_composition(r, m['w'])
    printed = defaultdict(set)
    for r in keep:
        if r['Formula Code'] != r['_final']:
            printed[r['_i']].add(r['Formula Code'])
    for i, t in target.items():
        if R[i]['Formula Code'] != t['_final'] and same[i]:
            printed[t['_i']].add(R[i]['Formula Code'])

    wb = Workbook()
    def sheet(name, header, data, widths=None, first=False):
        s = wb.active if first else wb.create_sheet(name)
        s.title = name; s.append(header)
        for c in s[1]:
            c.fill, c.font, c.alignment = HEAD, HF, WRAP
        for row in data:
            s.append(row)
        for i, h in enumerate(header, 1):
            s.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, 14)
        s.freeze_panes = 'C2'; s.auto_filter.ref = s.dimensions
        return s

    # ---- Read Me
    rm = sheet('Read Me', ['Item', 'Value'], [], {'Item': 24, 'Value': 120}, first=True)
    for item, val in [
        ('Workbook', 'W% Master Formulation - DRAFT, revision 2 (23 Oct 2026). Pure formulation (weight % by material group); no auger or line data.'),
        ('Revision 1', '30 Sep / 1 Oct 2026, built from Tech\'s calc workbooks and every FRM page on file (370 formulations, 242 codes).'),
        ('James\'s notes', 'James Kuo, 23 Oct 2026, in "W Master Formulation draft (version 1).xlsx" (kept as received in data/w_master): Main / Secondary / replace by / archive / remove / rename on every row; compositions rounded. '
                          '"There are too many new product code which make the one to one match not very useful. To make it work, we need to build logic system so that you assign formulation to product base on logic."'),
        ('Revision 2', 'Those notes applied, nothing else: sheet "W% Master v2" keeps the Main and Secondary formulations under the names James gives (old names as "Printed as", so Tech\'s pages still match); '
                       '"Replaced" lists every row folded into a Main or Secondary; "Archived" the archive / remove / obsolete rows and the VOIDFORM rows with reclaim (R10); "Open for James" what the notes do not settle.'),
        ('Later notes', 'Notes James gave after his 23 Oct file, applied the same way (daily/assign.py LATER_NOTES): '
                        + '; '.join(f'{code} "{old}" -> {new}' for code, old, new in assign.LATER_NOTES) + '.'),
        ('Rules', 'R1 colour other than KS/WB is graphic grade (third character A); the bubble line (BF) stays 0. R2 reclaim only in KS/WB. R3 the schedule\'s Spec decides the prefix '
                  '(R1R1R1 FU, R4R4R4 RU, OPOPOP FX, S1S1S1 FS, A1A1A1 FA, U1U1U1 FE - "Extra UV is FE", F.. FR, bubble line BF). R4 the middle digits are a version number. '
                  'R5 summer formulation 1 Apr - 31 Oct, winter 1 Nov - 31 Mar. R6 print James\'s names, never the old codes. R7 slitter / SPA / SPP orders take the sign-blank Main. '
                  'R8 roll orders take the Main roll, LTL wrap orders RU0001WB4. R9 a formulation with a Primary and a Secondary prints both: production runs the Primary '
                  '(reclaim) and switches to the virgin Secondary if reclaim runs out. R10 VOIDFORM never runs reclaim ("The reclaim calcium is not stable ... very difficult to control the weight"). '
                  'Full text and tests: docs/FORMULA_ASSIGNMENT_RULES.md in the formulation-pipeline repo.'),
        ('Used by', 'daily/assign.py: the daily Word formulation proposes the Main for a new order from these rows (rows marked PROPOSED BY RULE).'),
        ('Status', 'DRAFT. Nothing here is approved or in the Formulation Master; James reviews "Open for James" first.'),
    ]:
        rm.append([item, val])
    for row in rm.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP

    # ---- W% Master v2
    hdr = (['Formula Code', 'Printed as (old names)', 'Status', 'For', 'Prefix', 'Indicator (from the prefix)', 'Indicator (on its orders)',
            'Spec on its orders', 'Grade on its orders', 'Grade', 'Colour', 'Thk (mm)'] + W_COLS
           + ['Total', 'Materials', 'Last used', 'Evidence (calc / FRM)'] + RUN_COLS + ["James's note (23 Oct 2026)", 'Change in revision 2'])
    data = []
    for r in sorted((r for r in keep if r not in void_rc), key=lambda r: ((r['Colour'] or ''), r['Thk (mm)'] or 0, r['_final'], r['_level'])):
        code = r['_final']
        pre = assign.canon_prefix(code[:2]) if len(code) >= 2 else ''
        change = []
        if r['Formula Code'] != code:
            change.append(f"renamed from {r['Formula Code']}")
        if (r.get('Reclaim %') or 0) >= assign.RECLAIM_MIN and 'void' in str(r.get('Run: Special Instructions') or '').lower():
            change.append('R10: never on a VOIDFORM order (it has reclaim; its latest run was one)')
        absorbed = sorted({R[i]['Formula Code'] + f" ({R[i]['Application']}{' ' + R[i]['Version'] if R[i]['Version'] else ''}"
                                                   f"{'' if same[i] else ', other composition'})" for i, t in target.items() if t['_i'] == r['_i']})
        if absorbed:
            change.append('replaces ' + ', '.join(absorbed))
        total = sum((r.get(c) or 0) for c in W_COLS)
        spec = r.get('Run: Spec') or ''
        data.append([code, ', '.join(sorted(printed[r['_i']])), r['_level'].capitalize(), words(r['Notes']),
                     pre, INDICATOR.get(pre, ''), spec[:2] if spec else '', spec, r.get('Run: Grade') or '',
                     {'A': 'graphic', '0': 'packaging'}.get(code[2] if len(code) == 9 else '', ''), r['Colour'], r['Thk (mm)']]
                    + [r.get(c) for c in W_COLS] + [round(total, 1), r.get('Materials'), r.get('Last used'), r.get('Evidence (calc / FRM)')]
                    + [r.get(c) for c in RUN_COLS] + [r['Notes'], '; '.join(change)])
    s = sheet('W% Master v2', hdr, data, {'Formula Code': 13, 'Printed as (old names)': 22, 'For': 26, 'Materials': 34, "James's note (23 Oct 2026)": 40,
                                          'Change in revision 2': 40, 'Run: Special Instructions': 40, 'Evidence (calc / FRM)': 12,
                                          'Indicator (from the prefix)': 20, 'Spec on its orders': 12})
    # the indicator the prefix implies against the Spec the schedule printed for its latest run
    for row in s.iter_rows(min_row=2):
        ind, spec = row[hdr.index('Indicator (from the prefix)')].value, row[hdr.index('Indicator (on its orders)')].value
        if ind and spec and spec not in ind.split('/'):
            row[hdr.index('Indicator (on its orders)')].fill = AMBER
    for row in s.iter_rows(min_row=2):
        if row[2].value == 'Main':
            row[0].fill = GREEN
        for c in row:
            if isinstance(c.value, str) and len(c.value) > 30:
                c.alignment = WRAP
        if hasattr(row[hdr.index('Last used')].value, 'strftime'):
            row[hdr.index('Last used')].number_format = 'yyyy-mm-dd'
    # ---- Replaced
    data = []
    for i, t in sorted(target.items(), key=lambda kv: (R[kv[0]]['Colour'] or '', R[kv[0]]['Thk (mm)'] or 0, R[kv[0]]['Formula Code'])):
        r = R[i]
        data.append([r['Formula Code'], r['Application'], r['Version'], t['_final'], t['_level'].capitalize(),
                     'same formulation (within 2.5 points): its line settings carry over' if same[i]
                     else "other composition: where a line issued it, that is the line's own recipe of the Main (24 Oct 2026)",
                     r['Colour'], r['Thk (mm)']] + [r.get(c) for c in W_COLS] + [r.get('Materials'), r.get('Last used'), r['Notes']])
    s = sheet('Replaced', ['Old code', 'Application', 'Version', 'Replaced by', 'Status of that', 'Same composition?', 'Colour', 'Thk (mm)'] + W_COLS + ['Materials (old)', 'Last used', "James's note"],
              data, {'Materials (old)': 34, "James's note": 40, 'Same composition?': 44})
    # ---- Archived
    data = [[r['Formula Code'], r['Application'], r['Version'], r['_level'], r['Colour'], r['Thk (mm)']] + [r.get(c) for c in W_COLS] + [r.get('Materials'), r.get('Last used'), r['Notes']]
            for r in R if r['_level'] in ('archive', 'remove', 'other') or (r['Notes'] or '').strip().startswith('?')]
    data += [[r['Formula Code'], r['Application'], r['Version'], 'R10: VOIDFORM with reclaim (James, 23 Oct 2026: "void form should not use any reclaim")',
              r['Colour'], r['Thk (mm)']] + [r.get(c) for c in W_COLS] + [r.get('Materials'), r.get('Last used'), r['Notes']] for r in void_rc]
    s = sheet('Archived', ['Formula Code', 'Application', 'Version', 'Why', 'Colour', 'Thk (mm)'] + W_COLS + ['Materials', 'Last used', "James's note"],
              data, {'Materials': 34, "James's note": 40})
    # ---- Open for James
    unresolved = [r for r in R if r['_level'] == 'replace' and r['_i'] not in target]
    several = defaultdict(list)
    for r in keep:
        if r['_level'] == 'main' and not assign.quals(r['Notes']) and not assign.restricted(r['Notes']) and len(r['_final']) == 9:
            several[(assign.canon_prefix(r['_final'][:2]), r['_final'][2], (r['Colour'] or '')[:2], r['Thk (mm)'])].append(r['_final'])
    data = []
    for k, codes in sorted(several.items()):
        fe = {c for c in codes if c[:2] == 'FE'}
        if len(set(codes)) > 1 and not (k[0] == 'FE' and len(fe) == 1):     # FE key: the FE-named one (James, 26 Oct: "U1 is for FEA")
            data.append(['Two Mains for one key', f'{k[2]} {k[3]:g} mm {k[0]} {"graphic" if k[1] == "A" else "packaging"}: {", ".join(sorted(set(codes)))}', 'Which is the Main?'])
    for r in unresolved:
        data.append(['Replace-by note without a target', f"{r['Formula Code']} {r['Application']} {r['Version'] or ''}: \"{r['Notes']}\"", 'Which code replaces it?'])
    for r in keep:
        if len(r['_final']) != 9:
            data.append(['Code length', f"{r['Formula Code']} -> {r['_final']} ({r['Notes']})", 'Where does the A go / how is it spelt?'])
    data += [
        ['Extra UV is FE', 'FHA101WB3 and FHA012WB4 keep FH in the names James gave; FH0001NSA has two Mains (anti-static -> FHA001NSA, extra UV -> FEA001NSA)', 'Rename FHA.. to FEA..? Anti-static is FA by the Spec rule'],
        ['Weight-blender lines', 'SE43 prints FU0151WB3 (WB 3 mm packaging; note: Secondary), FU0010GS3 / FU0010BD3 / FUA000BD3 (coloured 3 mm); SE42 prints FU0151KS3 ("Main SE43 only")', 'Own Mains on SE42/SE43, or the general ones?'],
        ['GT premix', 'DPP40WB1700 runs FU001GTW4 / FU0021GT4 (note "???")', 'Keep as an engineer decision?'],
        ['3 mm VOIDFORM (James, 23 Oct: "Check old record for 3mm void form formula")',
         'Tech\'s records: SE11 issues FU0022WB3 "(For VOIDFORM order only)" on every 3 mm WB VOIDFORM order since 23 Sep (and the calc library notes it "Void Form" since 2023) - '
         'it is marked Main Voidform already. SE43 issues FU0151WB3 for its 3 mm WB orders, the VOIDFORM one (H69A038-1 DPP30WB1020, 28-30 Sep) and the plain packaging ones '
         '(H69A183-2/-3, H69A330-3, 6-7 Oct) alike; revision 1 labelled FU0151WB3 "Standard" (its latest run was the VOIDFORM order), so the VOIDFORM use was not visible and '
         'the note became "Secondary Packing grade". Also labelled Standard / run-out though their latest run was a VOIDFORM order: FU0011WB5 1 of 2, FUA011WB5 2 of 6, '
         'FU0001WBA, FU0041WBA, FU0001WBD, FU0061WBD (the VOIDFORM-specific rows exist beside them).',
         'FU0151WB3 = "Main SE43" (3 mm WB, void form and packaging), like FU0151KS3 and FU0151WB4?'],
        ['VOIDFORM without reclaim (R10)',
         'Tech has run VOIDFORM orders on reclaim formulations: FU0011WB5 on SE25 (H64A244-1, H66A116-1, H68A090-1) and FU0061WBD on SE22 (H68A053-1). Archived: '
         + ', '.join(f"{r['Formula Code']} \"{r['Notes']}\" ({r.get('Reclaim %'):g}% reclaim)" for r in void_rc)
         + '. The daily run gives a VOIDFORM order only a Main and Secondary without reclaim: WB 5 mm FU0051WB5, WB 13 mm FU0001WBD. Neither has run on the line '
         'the 23 Oct schedule puts it on (FU0051WB5 on SE25, FU0001WBD on SE32): those orders print "ENGINEER TO SET".',
         'Settings for FU0051WB5 on SE25 and FU0001WBD on SE32 (auger calculation)?'],
        # James, 26 Oct 2026: "FU0041WB4 as Secondary Corn box, make sure that is on your new W by W excel. will review them later"
        ['Corn box Secondary (26 Oct)',
         'This file has FU0041WB4 "Secondary Corn box" (virgin 77.1, F1203K 9.7, talc 5.5, CaCO3 5, colour 2.8). Tech’s corn box run-out is the '
         'virgin-only FU0021WB4 "(CORN BOX FORMULA) (Use this formula in case PP WB Reclaim is run out)" (SE23, no CaCO3). James, 26 Oct: "there '
         'should be a PPvirgin only formulation for corn box under the old FU0021WB4 ... should stay within cornbox formulation". The daily run '
         'now prints Tech’s corn box run-out as the Secondary where Tech has one (SE23 RP26902-1); FU0041WB4 only where it has none (SE42 RP26410-1). '
         'James, 26 Oct: "can you see is the name FU0022WB4 used? if its not use, set it up as corn box secondary formula" - FU0022WB4 was used '
         'nowhere (this file, Tech’s calc workbooks, pages on file, masters): Tech’s virgin corn box row now prints as FU0022WB4 '
         '(daily/assign.py RUN_OUT_NAMES). "the new secondary cornbox formulation is correct".',
         'Add FU0022WB4 "Secondary Corn box" (W% from Tech’s SE23 row); FU0041WB4 -> archive or keep?'],
        ['FW 4 mm flame retardant (26 Oct)',
         'No FW 4 mm row. Tech issues FR0021FW4 on SE21 (RP26826-6 RPP40FW2, WB:IS premix, FR-GPP3003MP); James: "ok we need to change it to FRA '
         'as well" - printed as FRA021FW4 (daily/assign.py LATER_RENAMES). FR0021WB4 / FRA021WB4 is the row whose orders are FW.',
         'Add FRA021FW4 as the FW 4 mm Main (and FR0021WB4 / FRA021WB4 -> ?)'],
        ['Black graphic 10 mm Secondary (26 Oct)',
         'FUA051KSA (reclaim) has no Secondary row. James: "i recall black graphic has option to go virgin PP with no reclaim for secondary". The '
         'daily run takes Tech’s SE22 run-out FUA001KSA (virgin, "in case PP MIX Reclaim is run out") as the Secondary.',
         'Add FUA001KSA as "Secondary, Graphic" for KS 10 mm?'],
    ]
    sheet('Open for James', ['Question', 'Detail', 'To decide'], data, {'Question': 30, 'Detail': 110, 'To decide': 40})
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f'{OUT}: {len(keep) - len(void_rc)} Main/Secondary rows, {len(target)} replaced, {len(unresolved)} unresolved, '
          f'{sum(1 for r in R if r["_level"] in ("archive", "remove", "other")) + len(void_rc)} archived ({len(void_rc)} VOIDFORM with reclaim)')


if __name__ == '__main__':
    main()
