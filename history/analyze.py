"""What the system's past schedules (history/prod_instr.py -> work/history/ext_history.csv) say about the database's
rules and masters. Read-only: prints findings; nothing is changed here.

    python history/analyze.py [section ...]     sections: orders, codes, lines, master, pairs (default: all)
"""
import csv
import re
import sys
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import config  # noqa: E402
import product_code  # noqa: E402

HIST = config.WORK_DIR / 'history' / 'ext_history.csv'
MONTH = {**{str(i): i for i in range(1, 10)}, 'A': 10, 'B': 11, 'C': 12}


def load():
    return list(csv.DictReader(open(HIST, encoding='utf-8')))


def shape(s):
    return ''.join('9' if ch.isdigit() else 'A' for ch in s)


def order_ym(base):
    """(year, month) an order number was opened, from its code; None when the code does not say.
    RP26821 -> (2026, 8); RP24A11 -> (2024, 10); H69A039 -> (2026, 9) when the year digit is read in the 2020s;
    H3CA288 -> (2023, 12); HV3A221 -> letter year (V=2020, W=2021 ... read from the data, not assumed)."""
    if base.startswith('RP') and len(base) == 7 and base[2:4].isdigit() and base[4] in MONTH:
        return 2000 + int(base[2:4]), MONTH[base[4]]
    if base.startswith('H') and len(base) == 7 and base[2] in MONTH:
        y = base[1]
        if y.isdigit():
            return 2020 + int(y), MONTH[base[2]]
        return ('letter ' + y), MONTH[base[2]]
    return None


def orders(rows):
    by = Counter((r['date'][:4], shape(r['order_base']), r['order_base'][:2] if r['order_base'][0] in 'RS' else r['order_base'][0]) for r in rows)
    print('== order number shapes by schedule year')
    for k, v in sorted(by.items()):
        print('  ', k, v)
    # H year character by schedule year
    yc = Counter((r['date'][:4], r['order_base'][1]) for r in rows if r['order_base'].startswith('H'))
    print('== H orders: year character by schedule year')
    for k, v in sorted(yc.items()):
        print('  ', k, v)
    # position 2 / 4 (month) values
    print('== month characters: H', sorted(Counter(r['order_base'][2] for r in rows if r['order_base'].startswith('H')).items()))
    print('== month characters: RP', sorted(Counter(r['order_base'][4] for r in rows if r['order_base'].startswith('RP')).items()))
    print('== H position 4 (after the month):', sorted(Counter(r['order_base'][3] for r in rows if r['order_base'].startswith('H')).items()))
    # order opened vs schedule date (digit-year era only)
    late, age = [], Counter()
    for r in rows:
        ym = order_ym(r['order_base'])
        if not ym or isinstance(ym[0], str):
            continue
        d = date.fromisoformat(r['date'])
        months = (d.year - ym[0]) * 12 + d.month - ym[1]
        if months < 0:
            late.append((r['date'], r['order'], r['line']))
        age[min(months, 36) if months >= 0 else -1] += 1
    print('== orders dated after their schedule (should be none):', len(late), late[:10])
    print('== order age in months at schedule (36 = 36+):', sorted(age.items()))
    sfx = Counter(len(r['suffix']) for r in rows)
    print('== suffix length:', sorted(sfx.items()), 'max', max(int(r['suffix']) for r in rows))


def codes(rows):
    probs = Counter()
    ex = defaultdict(set)
    for r in rows:
        for p in product_code.problems(r['prod_code']):
            k = re.sub(r'"[^"]*"', '"…"', p) if 'not in' not in p else p
            probs[k] += 1
            ex[k].add(r['prod_code'])
    print('== product codes breaking the rule (system history):')
    for k, v in probs.most_common():
        print(f'   {v:6} {k}   e.g. {sorted(ex[k])[:8]}')
    fams = Counter(r['prod_code'][:3] for r in rows)
    cols = Counter(r['prod_code'][5:7] for r in rows if product_code.CODE_RX.match(r['prod_code']))
    print('== families:', fams.most_common())
    print('== colours:', cols.most_common())
    thk = Counter()
    bad = defaultdict(set)
    for r in rows:
        t = product_code.thickness_mm(r['prod_code'])
        if t is None or r['thk'] in ('', None):
            continue
        ok = abs(t - float(r['thk'])) < 0.01
        thk[ok] += 1
        if not ok:
            bad[(r['prod_code'], r['thk'])].add(r['date'][:4])
    print('== code thickness vs Thk column: same', thk[True], 'different', thk[False])
    for (c, t), ys in sorted(bad.items())[:30]:
        print('   ', c, 'Thk', t, sorted(ys))
    # colour in the code vs the colour column (3 layers)
    cc = Counter()
    for r in rows:
        if product_code.CODE_RX.match(r['prod_code']) and r['colour']:
            cc[r['prod_code'][5:7] in r['colour'].split()] += 1
    print('== code colour among the 3 layer colours:', dict(cc))


def lines(rows):
    by = defaultdict(Counter)
    for r in rows:
        by[r['date'][:4]][r['line']] += 1
    print('== lines by year:')
    for y in sorted(by):
        print('  ', y, dict(sorted(by[y].items())))


def pairs(rows):
    by = defaultdict(set)
    for r in rows:
        by[r['date'][:4]].add((r['order'], r['prod_code']))
    print('== distinct order+product pairs by year:', {y: len(v) for y, v in sorted(by.items())})
    days = sorted({r['date'] for r in rows})
    print('== schedule days:', len(days), days[0], '->', days[-1])


def master(rows):
    from openpyxl import load_workbook
    p = config.published_path('Product Master.xlsx')
    ws = load_workbook(p, read_only=True)['Product Master']
    it = ws.iter_rows(values_only=True)
    head = list(next(it))
    pm = {r[0]: dict(zip(head, r)) for r in it if r[0]}
    hist = defaultdict(list)
    for r in rows:
        hist[r['prod_code']].append(r)
    in_h = set(hist)
    print('== Product Master', len(pm), 'products; system history', len(in_h), 'products')
    print('   in history, not in the PM:', len(in_h - set(pm)))
    print('   in the PM, never in history (2020-2026):', len(set(pm) - in_h))
    recent = {c for c, rs in hist.items() if max(x['date'] for x in rs) >= '2025-01-01'}
    print('   in history since 2025, not in the PM:', len(recent - set(pm)), sorted(recent - set(pm))[:20])
    blank = Counter()
    for c in in_h & set(pm):
        for h in ('Material', 'Spec', 'Colours (3 layers)', 'Thk (mm)', 'GSM', 'Width (in)', 'Length (in)', 'EXT Pack Code', 'PCs/Stack'):
            if pm[c].get(h) in (None, ''):
                blank[h] += 1
    print('   PM products with history but a blank column:', dict(blank))


if __name__ == '__main__':
    rows = load()
    want = sys.argv[1:] or ['orders', 'codes', 'lines', 'pairs', 'master']
    for s in want:
        globals()[s](rows)
