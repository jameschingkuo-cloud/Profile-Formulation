"""Print-ready formulation document (Word) for the day's extrusion schedule (James Kuo, 29 Sep 2026: "operator will use
the paper copy. So whenever i or any other engineer scan you the production schedule. you will product a word
formulation document for us to print out"; "you can use the current one (and maybe improve a little) as template").

    PKT_DATE=2026-09-30 python daily/resolve.py        (the FRM Draft: every order resolved, or an Exception)
    PKT_DATE=2026-09-30 python daily/render_frm.py     -> out/FRM Formulation <date>.docx

Layout follows Tech's FRM pages (one landscape page per line: "Line N (SExx) Formulations", orders, formula code,
feeders, note, footnotes). Improvements: what "Set" means on that line; product under each order; every formulation of
an order in run order ("Run first", "If reclaim runs out"); replaced materials printed as what to load; orders with no
formulation shown as ENGINEER TO COMPLETE; a cover page with the exceptions and the issue block.

IWPFO055 §5.3: Technical personnel issue the formulation. This file is PREPARED by the pipeline; it is a DRAFT until an
engineer completes the exceptions and signs the issue block. Nothing is guessed: an Exception prints blank, never a
suggestion.
"""
import ast
import datetime
import json
import os
import re
import sys
from collections import OrderedDict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / 'daily'))
import config  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.section import WD_ORIENT  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt, Inches, RGBColor  # noqa: E402
from resolve import split_feeder  # noqa: E402
import assign  # noqa: E402

PKT = os.environ.get('PKT_DATE')
WEIGHT = {'SE24', 'SE42', 'SE43', 'SE61'}
LINE_NO = {'SE11': 1, 'SE12': 2, 'SE13': 3, 'SE21': 4, 'SE22': 5, 'SE23': 6, 'SE24': 7, 'SE31': 8, 'SE32': 9, 'SE25': 10,
           'SE42': 12, 'SE43': 13, 'SE61': 16}
FONT = 'Times New Roman'          # as on Tech's pages
# Black and white only: production prints on a black-and-white printer (James Kuo, 27 Oct 2026: "no BLACK, make everything
# black and white"). Emphasis is bold / italic, never colour or shading.
BLACK = RGBColor(0, 0, 0)


# Printed as decided by James (29 Sep 2026), without a "(replaces ...)" note: spelling / naming only, same material.
PRINT_AS = {
    'FOAM – Bergen XO-256': 'FOAM – Bergen X0-256',   # "change them to X0" (master Change 46)
    'PP Virgin': 'PP Virgin (F6502A)',                 # "PP Virgin is F6502A" (master Changes 21-23)
    'White Reclaim': 'PP WB Reclaim (White Reclaim)',  # "same as PP WB Reclaim" (master Changes 15-20)
}
# The black (CR400K) prints its sources as IWPFT062 lists them (rev 18: NPC PE-B9000F; approved second sources Avient
# 3800 Black PEC, MDI Black PE-500, Ampacet 19501 Black PE MB). Tech's pages still name Spartech B60009 and PolyOne
# LD-250, which the list does not, and spell the NPC grade PE90000F. James, 27 Oct 2026: "we should follow it".
# Printed short: James, 27 Oct 2026: "the black take up too much space. just put KS NPC PEB90000F" (IWPFT062's spelling
# of the grade is PE-B9000F; the second sources stay on IWPFT062)
KS_NPC = KS_MDI = 'KS NPC PE-B9000F'
PRINT_AS.update({
    'KS- NPC PE90000F (or Spartech B60009 or MDI PE-500)': KS_NPC,
    'KS-MDI PE-500 (or PolyOne LD-250 or NPC PE90000F': KS_MDI,
    'KS-MDI PE-500 (or Spartech B60009 or NPC PE90000F)': KS_MDI,
    'KS-MDI PE-500 (or Spartech B60009)': KS_MDI,
})


def replaced():
    tree = ast.parse((ROOT / 'daily' / 'checks.py').read_text(encoding='utf-8'))
    for node in ast.walk(tree):
        if isinstance(node, ast.Assign) and any(getattr(t, 'id', '') == 'REPLACED' for t in node.targets):
            return {k: v[0].split(' (')[0] for k, v in ast.literal_eval(node.value).items()}
    return {}


def num(s):
    try:
        return float(str(s).replace(',', ''))
    except (TypeError, ValueError):
        return None


# ---- data ----------------------------------------------------------------------------------------------------------
def load():
    draft_path = config.OUTPUT_DIR / f'FRM Draft {PKT}.xlsx'
    if not draft_path.exists():
        raise SystemExit(f'No {draft_path.name}: run PKT_DATE={PKT} python daily/resolve.py first')
    config.record_read(draft_path, 'FRM Draft')
    wb = load_workbook(draft_path, read_only=True)

    def rows(name):
        r = list(wb[name].iter_rows(values_only=True))
        return [dict(zip(r[0], x)) for x in r[1:]]
    draft, exc = rows('Draft'), rows('Exceptions')
    packets = {}
    files = {}
    for p in config.packet_files():
        d = json.loads(p.read_text(encoding='utf-8'))
        packets[d['packet_date']] = d; files[d['packet_date']] = p
    today = packets[PKT]
    config.record_read(files[PKT], 'daily packet (schedule)')
    # page layout per line: latest issued FRM page on or before this date
    page = {}
    for d in sorted(x for x in packets if x <= PKT):
        for pg in packets[d]['frm']:
            page[pg['line_code']] = pg
    return today, draft, exc, page


def build_lines(today, draft, exc, page):
    repl = replaced()
    forms = OrderedDict()       # (line, order) -> {row: {'code', 'note', 'feeders': {(ext, feeder): (mat, set, printed)}}}
    for r in draft:
        f = forms.setdefault((r['Line Code'], r['Order']), {}).setdefault(r['Formula Row'], {'code': r['Formula Code'], 'note': r['Note'] or '', 'feeders': {}})
        mat = r['Material (as issued)'] or ''
        # an obsolete number Tech's page still prints (F1102K, Q1203K) prints as what is loaded now, with no note: James,
        # 26 Oct 2026: "No need to talk about replacing F1102K that number is obsolete, engineer just never had the
        # chance to clean it up"
        shown = repl.get(mat) or PRINT_AS.get(mat, mat)
        if r['Feeder'] is None:          # a Primary / Secondary with no settings on the line: printed "ENGINEER TO SET"
            continue
        f['feeders'][(r['Extruder'] or '', r['Feeder'])] = (shown, r['Set'], None)
    exc_by = {(e['Line Code'], e['Order']): e for e in exc}
    lines = []
    for e_pg in sorted(today['ext'], key=lambda p: p['scan_page']):
        pass
    by_line = OrderedDict()
    for pg in sorted(today['ext'], key=lambda p: p['scan_page']):
        for r in pg['rows']:
            by_line.setdefault(pg['line'], []).append(r)
    for line in sorted(by_line, key=lambda c: LINE_NO.get(c, 99)):
        pg = page.get(line)
        cols = [split_feeder(c) for c in (pg['feeder_columns'] if pg else [])]
        groups = OrderedDict()   # formula signature -> {'orders': [...], 'forms': [...]}
        for r in by_line[line]:
            f = forms.get((line, r['order']))
            # the same formulation shares one cell (James, 26 Oct 2026, RP26717-1 and RP24C18-4, both FUA151WB4
            # 48/25/48/15: "since these two are the same formula, they should share the same cell")
            if f:
                fl = [f[k] for k in sorted(f)]
                sig = json.dumps([(x['code'], x['note'], sorted((list(k), list(v)) for k, v in x['feeders'].items())) for x in fl])
            else:
                fl, sig = None, 'EXC:' + r['order']
            g = groups.setdefault(sig, {'orders': [], 'forms': fl, 'exc': [] if f else [exc_by.get((line, r['order']))]})
            g['orders'].append(r)
        used = {k for g in groups.values() if g['forms'] for x in g['forms'] for k in x['feeders']}
        if not cols:
            cols = sorted(used)
        elif len(cols) > 6:     # co-ex / V-layouts: print only the feeders in use today
            cols = [c for c in cols if c in used] or cols[:4]
        lines.append({'line': line, 'no': LINE_NO.get(line), 'cols': cols, 'groups': list(groups.values()),
                      'ac': (pg or {}).get('header_note', ''), 'notes': (pg or {}).get('footnotes', []),
                      'effective': (pg or {}).get('effective_date', ''), 'weight': line in WEIGHT})
    return lines


# ---- weight % copy -------------------------------------------------------------------------------------------------
# James, 27 Oct 2026: "make two copy of the formulation in the future, one weight % and one normal. This will help us
# troubleshoot faster" - "word file". The weight % copy shows each feeder's share of the sheet: on an auger line the
# setting x the hopper's calibration slope (g/min per unit), over the total; a weight-blender line is already weight %.
def slope_classes(mat):
    m = (mat or '').upper()
    if 'TALC' in m:
        return ['TALC']
    if '1203K' in m or '1102K' in m:
        return ['1203K', '1102K']
    if 'VIRGIN' in m or '6502' in m:
        return ['VIRGIN PP']
    if 'CACO3' in m:
        return ['CACO3']
    if 'RECLAIM' in m:
        return (['MIX RECLAIM'] if 'MIX' in m else ['WB RECLAIM'] if 'WB' in m else []) + ['RECLAIM']
    if 'VISTAMAXX' in m:
        return ['VISTAMAXX']
    if 'GPP' in m:              # FR masterbatch: printed "FR-GPP30003MP", the calc book "FR-GPP3003 MP" (RP26826-6, SE21)
        return ['FR-GPP']
    codes = re.findall(r'[A-Z]?\d{5}[A-Z]?|[A-Z]{1,2}-?[A-Z]?\d{5}[A-Z]?', m)
    # then the bare 5-digit number: Tech's calc book names SE31 H5 "NPC 22151" for the printed "NPC PE-W22151" (H69A059-1)
    codes += [n for n in re.findall(r'(?<!\d)\d{5}(?!\d)', m) if n not in codes]
    return codes + ([m[:2] + ' COLOR'] if re.match(r'^[A-Z]{2}[- ]', m) else []) + ['COLOR']


def weight_pct(line, feeders):
    """{feeder key: (material, set, printed)} -> ({key: 'x.x'}, '') or (None, why) on an auger line."""
    w = {}
    for key, (mat, st, _) in feeders.items():
        v = num(st)
        if v is None:
            continue
        digit = re.search(r'(\d)\s*$', key[1])
        k = next((sl for c in slope_classes(mat) for sl in [assign.class_slope(line, 'H' + digit.group(1), c)[0]] if sl),
                 None) if digit else None
        if not k:
            return None, f"no calibration for {key[1]} {mat}"
        w[key] = v * k
    t = sum(w.values())
    if not t:
        return None, 'no settings'
    return {key: f"{x * 100 / t:.1f}" for key, x in w.items()}, ''


# ---- docx helpers --------------------------------------------------------------------------------------------------
def para(cell_or_doc, text='', size=10, bold=False, color=None, align=None, italic=False, space_after=0):
    p = cell_or_doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if text:
        run(p, text, size, bold, color, italic)
    return p


def run(p, text, size=10, bold=False, color=None, italic=False):
    r = p.add_run(text)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    return r


def cell_text(cell, text, size=10, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.CENTER, italic=False):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    if text:
        run(p, text, size, bold, color, italic)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return p


def set_widths(table, widths):
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader'); el.set(qn('w:val'), 'true'); trPr.append(el)


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:cantSplit'); el.set(qn('w:val'), 'true'); trPr.append(el)


def page_field(p):
    for code in ('PAGE',):
        for kind, txt in (('begin', None), (None, code), ('end', None)):
            r = p.add_run(); r.font.name = FONT; r.font.size = Pt(8)
            if kind:
                fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), kind); r._r.append(fc)
            else:
                it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = f' {txt} '; r._r.append(it)


# ---- the document --------------------------------------------------------------------------------------------------
RECLAIM_OUT = re.compile(r'reclaim\b[^()]*?\brun\s*out', re.I)


def when_label(forms, i):
    """Forms are in run order (resolve.run_order: reclaim first). A formula without reclaim after one with reclaim runs when
    the reclaim runs out (James Kuo, 30 Sep 2026: "we always want to use up our scrap first before using Virgin PP"); so
    does James's Secondary, also when its Primary has no settings yet (R9, 23 Oct 2026)."""
    uses = lambda f: any('reclaim' in (m or '').lower() and (num(s) or 0) > 0 for (m, s, _) in f['feeders'].values())
    f = forms[i]
    # only Tech's reclaim run-out note: "use W26038A if NPC PE-W22151 is run out" is a colour swap (James Kuo, 28 Oct 2026,
    # H69A059-1: "there is only one primary. why do you write 'if reclaim run out?' we dont use reclaim for that order")
    if RECLAIM_OUT.search(f['note']) or re.search(r'(?:^|· )Secondary\b', f['note']):
        return 'If reclaim runs out'
    if len(forms) == 1:
        return ''
    if i == 0:
        return 'Run first'
    if uses(f):
        return 'Next'
    return 'If reclaim runs out' if any(uses(x) for x in forms[:i]) else 'Alternative'


def render(today, lines, exc, out, weight_copy=False):
    day = datetime.date.fromisoformat(PKT)
    us = f'{day.month}/{day.day}/{day:%y}'
    n_orders = sum(len(g['orders']) for ln in lines for g in ln['groups'])
    n_exc = sum(len(g['orders']) for ln in lines for g in ln['groups'] if not g['forms'])
    n_set = sum(len(g['orders']) for ln in lines for g in ln['groups'] if g['forms'] and any(not x['feeders'] for x in g['forms']))
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Inches(11), Inches(8.5)
    for side in ('left_margin', 'right_margin'):
        setattr(sec, side, Inches(0.45))
    sec.top_margin, sec.bottom_margin = Inches(0.45), Inches(0.5)
    st = doc.styles['Normal']; st.font.name = FONT; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    status = 'DRAFT' if n_exc or n_set else 'READY TO ISSUE'
    # header / footer on every page
    hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(hp, f'{status} · not for production until issued by Technical (IWPFO055 §5.3)' if n_exc or n_set else
        'Prepared by the formulation pipeline · issued only when signed by Technical (IWPFO055 §5.3)', 8, color=BLACK)
    fp = sec.footer.paragraphs[0]
    run(fp, f"Schedule {today.get('source_scan', '')} · packet {PKT}{' (step 1: read by the scan reader)' if today.get('stage') == 1 else ''} · formulation by the rules of 23 Oct 2026 (James Kuo's W% Master), settings as Technical last issued them on the line · page", 8, color=BLACK)
    page_field(fp)

    # ---- cover
    para(doc, 'Inteplast Profile Plant · WPJK · Technical', 10, color=BLACK)
    para(doc, f'Extrusion Formulations · {us}', 26, bold=True, space_after=2)
    para(doc, f"{len(lines)} lines · {n_orders} orders · from the extrusion schedule {today.get('source_scan', '')}", 12, space_after=8)
    box = doc.add_table(rows=1, cols=1); box.alignment = WD_TABLE_ALIGNMENT.LEFT
    c = box.rows[0].cells[0]
    cell_text(c, f'{status}', 16, True, BLACK, WD_ALIGN_PARAGRAPH.LEFT)
    para(c, (f'{n_exc} order{"s" if n_exc != 1 else ""} have no formulation yet' + (f'; {n_set} have a Primary or Secondary with no settings on its line' if n_set else '') +
             '. An engineer completes them (table below), then Technical signs the issue block. Until then, do not run from this copy.') if n_exc or n_set else
         'Every order has its formulation. Technical checks and signs the issue block, then it goes to the floor.', 11)
    para(doc, '', space_after=4)
    para(doc, 'How to read the line pages', 12, bold=True, space_after=2)
    tips = ['Auger lines: Set is the auger motor speed, 0 to 100. It is not a percentage.',
            'Weight blender lines (Line 7, 12, 13, 16): Set is weight %; each extruder adds to 100; Auto takes the balance.',
            'Where an order lists more than one formulation, they are printed in run order: reclaim first (use up the scrap before virgin PP), then down the list (James Kuo, 29 and 30 Sep 2026).',
]
    tips.append("Primary and Secondary (James Kuo, 23 Oct 2026): where an order shows both, both are its formulation. Production decides by the "
                "material on hand: when reclaim runs out, the reclaim formula is replaced by the virgin-resin one. VOIDFORM orders never run reclaim (weight).")
    n_calc = sum(len(g['orders']) for ln in lines for g in ln['groups'] if g['forms'] and any('CALCULATED SETTINGS' in x['note'] for x in g['forms']))
    if n_calc:
        tips.append(f'CALCULATED SETTINGS ({n_calc} order{"s" if n_calc != 1 else ""}): never issued on the line; worked out from the line\'s own auger '
                    'settings for the same kind of product and its auger calibration (R4 black: 98% black reclaim, 2% black colour - James Kuo, '
                    '24 Oct 2026). An engineer checks them before the order runs.')
    for t in tips:
        para(doc, '•  ' + t, 10.5, space_after=1)
    if n_exc or n_set:
        para(doc, '', space_after=4)
        para(doc, 'Orders an engineer must complete', 12, bold=True, space_after=2)
        t = doc.add_table(rows=1, cols=6); t.style = 'Table Grid'
        hdr = ['Line', 'Order', 'Product', 'Why', 'Suggestion (not used)', 'Formula decided / by']
        for i, hd in enumerate(hdr):
            cell_text(t.rows[0].cells[i], hd, 9.5, True)
        repeat_header(t.rows[0])
        for ln in lines:
            for g in ln['groups']:
                if g['forms']:
                    for x in g['forms']:
                        if x['feeders']:
                            continue
                        for o in g['orders']:
                            rr = t.add_row(); no_split(rr)
                            vals = [ln['line'], o['order'], o['prod_code'], f"Settings for {x['code']} to be set on this line ({x['note'].split(' · ')[-1]})", '', '']
                            for i, v in enumerate(vals):
                                cell_text(rr.cells[i], str(v), 9.5, i == 1, align=WD_ALIGN_PARAGRAPH.LEFT)
                    continue
                for o in g['orders']:
                    e = (g['exc'] or [None])[0] or {}
                    rr = t.add_row(); no_split(rr)
                    vals = [ln['line'], o['order'], o['prod_code'], e.get('Reason') or 'No formulation found', e.get('Suggestion (not used)') or '', '']
                    for i, v in enumerate(vals):
                        cell_text(rr.cells[i], str(v), 9.5, i == 1, align=WD_ALIGN_PARAGRAPH.LEFT)
        set_widths(t, [Inches(0.6), Inches(1.0), Inches(1.2), Inches(2.4), Inches(2.6), Inches(2.3)])
    para(doc, '', space_after=6)
    para(doc, 'Issue (IWPFO055 §5.3 – §5.5)', 12, bold=True, space_after=2)
    it = doc.add_table(rows=2, cols=4); it.style = 'Table Grid'
    for i, hd in enumerate(['Issued by (Technical)', 'Signature', 'Date / time', 'Copy in "Schedule" binder (§5.4)']):
        cell_text(it.rows[0].cells[i], hd, 9.5, True)
    it.rows[1].height = Inches(0.45)
    cell_text(it.rows[1].cells[3], '☐', 14)
    set_widths(it, [Inches(2.6), Inches(2.6), Inches(2.0), Inches(2.9)])

    # ---- one page per line
    for ln in lines:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        tp = para(doc, '', align=WD_ALIGN_PARAGRAPH.CENTER)
        run(tp, f"Line {ln['no']} ({ln['line']})  Formulations" + ('  ·  Weight %' if weight_copy else ''), 22, True)
        mp = para(doc, '', align=WD_ALIGN_PARAGRAPH.RIGHT)
        run(mp, us, 11, True)          # Tech's "AC = 1 / 90" not printed (James, 27 Oct 2026: "lets take it out. not sure what it is")
        para(doc, ('WEIGHT % COPY, for troubleshooting - run from the normal copy. ' + ('Weight blender: as set (weight %).' if ln['weight'] else
                   "Each material's share of the sheet: auger setting x the hopper's calibration slope, over the total.")
                   if weight_copy else
                   ('Weight blender: Set = weight %. Each extruder adds to 100; Auto = the balance.' if ln['weight'] else
                   'Auger feeders: Set = auger speed, 0 to 100 (not %).') + '  Several formulations: run the first while reclaim lasts, then the next.'),
             9.5, italic=True, color=BLACK, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
        cols = ln['cols']
        ncol = 3 + len(cols)
        t = doc.add_table(rows=1, cols=ncol); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        heads = ['Order #', 'Formula Code'] + [(f'Extruder {e} · {f}' if e else f) for e, f in cols] + ['Note']
        for i, hd in enumerate(heads):
            cell_text(t.rows[0].cells[i], hd, 9.5 if len(cols) > 6 else 11, True)
        repeat_header(t.rows[0])
        for g in ln['groups']:
            first_row = None
            forms = g['forms'] or [None]
            for i, f in enumerate(forms):
                rr = t.add_row(); no_split(rr)
                first_row = first_row or rr
                if f is None:   # no formulation: one wide cell to write it in by hand, never a guess
                    rr.height = Inches(0.55)
                    wide = rr.cells[1].merge(rr.cells[ncol - 2])
                    cell_text(wide, 'ENGINEER TO COMPLETE: formula code and settings', 10, True, BLACK, WD_ALIGN_PARAGRAPH.LEFT)
                    cell_text(rr.cells[ncol - 1], (g['exc'][0] or {}).get('Reason', '') if g['exc'] else '', 8.5, color=BLACK)
                    continue
                if not f['feeders']:            # Primary / Secondary with no settings on this line: written in by the engineer
                    rr.height = Inches(0.45)
                    cell_text(rr.cells[1], f['code'], 11)
                    lab = when_label(forms, i)
                    if lab:
                        para(rr.cells[1], lab, 8.5, bold=True, color=BLACK, align=WD_ALIGN_PARAGRAPH.CENTER)
                    if len(cols) >= 1:
                        wide = rr.cells[2].merge(rr.cells[ncol - 2])
                        cell_text(wide, f"ENGINEER TO SET: settings for {f['code']} on this line", 10, True, BLACK, WD_ALIGN_PARAGRAPH.LEFT)
                    cell_text(rr.cells[ncol - 1], f['note'], 8.5, color=BLACK)
                    continue
                pct, why = weight_pct(ln['line'], f['feeders']) if weight_copy and not ln['weight'] else (None, '')
                cp = cell_text(rr.cells[1], f['code'], 11)
                lab = when_label(forms, i)
                if lab:
                    para(rr.cells[1], lab, 8.5, bold=True, color=BLACK, align=WD_ALIGN_PARAGRAPH.CENTER)
                for ci, key in enumerate(cols, start=2):
                    v = f['feeders'].get(key)
                    if not v:
                        cell_text(rr.cells[ci], '')
                        continue
                    mat, st_, printed = v
                    cell_text(rr.cells[ci], mat, 9 if len(cols) > 6 else 10, bool(printed))
                    if printed:
                        para(rr.cells[ci], f'(replaces {printed})', 7.5, italic=True, color=BLACK, align=WD_ALIGN_PARAGRAPH.CENTER)
                    sv = str(st_ or '')
                    if ln['weight'] and sv.lower() == 'auto':
                        ext = key[0]
                        fixed = sum(num(s) or 0 for (e2, _), (m2, s, _) in f['feeders'].items() if e2 == ext and str(s).lower() != 'auto')
                        sv = f'Auto ({100 - fixed:g})'
                    if weight_copy and not ln['weight']:
                        sv = f"{pct[key]} %" if pct and key in pct else ('-' if not pct else '')
                    para(rr.cells[ci], sv, 12, True, align=WD_ALIGN_PARAGRAPH.CENTER)
                cell_text(rr.cells[ncol - 1], f['note'] + (f" · weight % not worked out: {why}" if weight_copy and why else ''), 8.5)
            # order # cell, merged down the group's rows
            last = t.rows[-1]
            oc = first_row.cells[0].merge(last.cells[0]) if last is not first_row else first_row.cells[0]
            oc.text = ''
            for k, o in enumerate(g['orders']):
                p = oc.paragraphs[0] if k == 0 else oc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
                run(p, o['order'], 10.5, True)
                q = oc.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.space_after = Pt(1)
                run(q, o['prod_code'], 7.5, color=BLACK)
            oc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        usable = 10.1
        w_order, w_code, w_note = 1.55, 1.25, 1.45
        w_feed = (usable - w_order - w_code - w_note) / max(1, len(cols))
        set_widths(t, [Inches(w_order), Inches(w_code)] + [Inches(w_feed)] * len(cols) + [Inches(w_note)])
        for n in ln['notes']:
            para(doc, n, 10, space_after=0)
        fp2 = para(doc, '', space_after=0)
        run(fp2, 'Tech. Department', 10); run(fp2, f"        Effective Date: {ln['effective']}" if ln['effective'] else '', 10)
    doc.save(out)
    return n_orders, n_exc + n_set


def main():
    if not PKT:
        raise SystemExit('Set PKT_DATE=YYYY-MM-DD')
    today, draft, exc, page = load()
    lines = build_lines(today, draft, exc, page)
    out = config.OUTPUT_DIR / f'FRM Formulation {PKT}.docx'
    n, e = render(today, lines, exc, out)
    print(out, f'- {len(lines)} lines, {n} orders, {e} for an engineer')
    wout = config.OUTPUT_DIR / f'FRM Formulation {PKT} (weight %).docx'     # the troubleshooting copy
    render(today, lines, exc, wout, weight_copy=True)
    print(wout)


if __name__ == '__main__':
    main()
