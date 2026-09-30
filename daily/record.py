"""The three history records: Formulation Report Record, Extrusion Production Record, Converting Production Record.

    python daily/record.py 2026-09-23 2026-09-24 ...      (append these packet dates, in the order given)
    python daily/record.py --history work/history/ext_history.csv [more.csv ...] [--replace-history]
        (Extrusion Production Record: the system's own past schedules, history/prod_instr.py; --replace-history first
         takes out the rows an earlier --history run added - never a row from a daily scan)

Append-only (db/schema.py kind 'record'; James Kuo, 28 Sep 2026). Each record starts from the published copy in
PUBLISH_DIR (read and recorded, per the hard rule) or, if none is published yet, from nothing. For every date:
  - not in the record yet        -> its rows are appended;
  - already there, rows the same -> skipped;
  - already there, rows differ   -> STOP: nothing is written. Past rows are never edited.
Rows come from data/packets/packet_<date>.json, values as printed. The workbooks go to out/; publish.py publishes.
"""
import datetime
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import config  # noqa: E402
from db import schema  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
sys.path.insert(0, str(Path(__file__).resolve().parent))
from resolve import split_feeder, variant  # noqa: E402

DOSING = {'SE24': 'WEIGHT', 'SE42': 'WEIGHT', 'SE43': 'WEIGHT', 'SE61': 'WEIGHT'}   # the rest are auger lines (load.DOSING)
PLTS_DONE = re.compile(r'(\d[\d,]*)\s*PLTS?\s+DONE', re.I)
X_OF_Y = re.compile(r'^\s*(\d[\d,]*)\s+OF\s+(\d[\d,]*)\s*$', re.I)
RECORDS = {'FRM': ('Formulation Report Record.xlsx', 'Issued'),
           'EXT': ('Extrusion Production Record.xlsx', 'Orders by Day'),
           'CNV': ('Converting Production Record.xlsx', 'Orders by Day')}


def num(s):
    try:
        v = float(str(s).replace(',', '').strip())
    except (TypeError, ValueError):
        return None
    return int(v) if v.is_integer() else v


def as_int_or_text(s):
    """int when the print is a number, else the text as printed (e.g. '######' overflow); None when blank."""
    if s in (None, ''):
        return None
    v = num(s)
    return v if v is not None else str(s)


def scan_ref(pk, kind, page):
    scan = pk.get('frm_source_scan') if kind == 'frm' and pk.get('frm_source_scan') else pk.get('source_scan', '')
    return f"{scan} p{page}"


# ---- rows per record, from one packet -------------------------------------------------------------------------
def earlier_products(date):
    """order -> (product, packet date) from the EXT pages of packets before `date`, latest wins. For an order still on
    the FRM after it left the schedule (H69A139-1 on 25 Sep): its product code is printed on an earlier EXT."""
    out = {}
    for p in sorted(config.PACKETS_DIR.glob('packet_*.json')):
        d = p.stem.split('_', 1)[1]
        if d < date:
            config.record_read(p, 'daily packet (earlier EXT, product lookup)')
            for pg in json.loads(p.read_text(encoding='utf-8'))['ext']:
                for r in pg['rows']:
                    out[r['order']] = (r['prod_code'], d)
    return out


def frm_rows(pk):
    day = datetime.date.fromisoformat(pk['packet_date'])
    prod = {r['order']: r['prod_code'] for pg in pk['ext'] for r in pg['rows']}
    missing = {o for pg in pk['frm'] for g in pg['groups'] for o in g['orders'] if o not in prod}
    if missing:
        early = earlier_products(pk['packet_date'])
        for o in missing:
            if o in early:
                prod[o] = early[o][0]
                prod[o + '#from'] = early[o][1]
    out = []
    for pg in pk['frm']:
        line = pg['line_code']
        for g in pg['groups']:
            for i, f in enumerate(g['formulas']):
                var = variant(f.get('note', ''), i == 0)
                used = [(c, v) for c, v in f['feeders'].items() if v.get('material') or v.get('set')]
                pct = weight_pct(line, used)
                for o in g['orders']:
                    for c, v in used:
                        ext, feeder = split_feeder(c)
                        out.append({'Issue Date': day, 'Line Code': line, 'Order': o, 'Product Code': prod.get(o),
                                    'Formula Code': f['formula_code'], 'Variant': var, 'Extruder': ext or None,
                                    'Feeder': feeder, 'Formula Row': i + 1, 'Material ID': None,
                                    'Material (as printed)': v.get('material') or None, 'Set': v.get('set') or None,
                                    'Weight %': pct.get(c), 'Note': f.get('note') or None, 'Source': 'Tech FRM',
                                    'Source Scan': scan_ref(pk, 'frm', pg['scan_page']) + (
                                        f"; not on this day's EXT, product code from EXT {prod[o + '#from']}"
                                        if o + '#from' in prod else ''),
                                    'Signed By': None, 'Signed At': None})
    return out


def weight_pct(line, used):
    """Weight lines: Set is weight %; 'Auto' takes the balance of its extruder to 100. Auger lines: none."""
    if DOSING.get(line) != 'WEIGHT':
        return {}
    by_ext = {}
    for c, v in used:
        by_ext.setdefault(split_feeder(c)[0], []).append((c, v.get('set', '')))
    pct = {}
    for ext, items in by_ext.items():
        autos = [c for c, s in items if str(s).strip().lower() == 'auto']
        fixed = {c: num(s) for c, s in items if c not in autos}
        for c, s in fixed.items():
            if s is not None:
                pct[c] = s
        if len(autos) == 1 and all(s is not None for s in fixed.values()):
            pct[autos[0]] = round(100 - sum(fixed.values()), 4)
    return pct


def ext_rows(pk):
    day = datetime.date.fromisoformat(pk['packet_date'])
    out = []
    for pg in pk['ext']:
        for r in pg['rows']:
            si = r.get('special_instructions') or ''
            m = PLTS_DONE.search(si)
            sheets = [num(c.get('total_sheets')) for c in r.get('cut_rows', [])]
            out.append({'Schedule Date': day, 'Line Code': pg['line'], 'Order': r['order'], 'Product Code': r['prod_code'],
                        'Total Sheets': sum(s for s in sheets if s is not None) if any(s is not None for s in sheets) else None,
                        'Weight (LBs)': num(r.get('weight_lbs')), 'Plts Done (EXT)': num(m.group(1)) if m else None,
                        'Plts Ordered': as_int_or_text(r.get('num_plt')), 'Special Instructions': si or None,
                        'Handwritten': r.get('handwritten') or None, 'Source Scan': scan_ref(pk, 'ext', pg['scan_page'])})
    return out


def cnv_pages(pk):
    """A sheet scanned twice (same line, same page-of, same rows) is kept once, as daily/load.py does."""
    seen, pages = set(), []
    for p in sorted(pk['cnv'], key=lambda p: p['scan_page']):
        k = (p['line_code'], p.get('page_of'),
             json.dumps([{a: b for a, b in r.items() if a != 'unclear'} for r in p['rows']], sort_keys=True))
        if k not in seen:
            seen.add(k)
            pages.append(p)
    return pages


def cnv_rows(pk):
    day = datetime.date.fromisoformat(pk['packet_date'])
    out, count = [], {}
    for pg in cnv_pages(pk):
        for r in pg['rows']:
            k = (pg['line_code'], r['order'])
            count[k] = count.get(k, 0) + 1
            m = X_OF_Y.match(r.get('extrusion_status') or '')
            out.append({'Schedule Date': day, 'Converting Line': pg['line_code'], 'Order': r['order'], 'Row': count[k],
                        'Product Code': r.get('product_code'), 'Extrusion Status (printed)': r.get('extrusion_status') or None,
                        'Plts Extruded': num(m.group(1)) if m else None, 'Plts Ordered': num(m.group(2)) if m else None,
                        'Semi Size': r.get('semi_size') or None, 'Die #': r.get('die_no') or None,
                        'Die Status': r.get('die_status') or None, 'Plate Status': r.get('plate_status') or None,
                        'Ink Color': r.get('ink_color') or None, 'Total Sheets': as_int_or_text(r.get('total_sheets')),
                        'Pack Code': r.get('pack_code') or None, '# of Plts': as_int_or_text(r.get('num_plts')),
                        'Pc/Plt': as_int_or_text(r.get('pc_per_plt')), 'Req. Date': r.get('req_date') or None,
                        'Done Note': r.get('done_note') or None, 'Handwritten': r.get('handwritten') or None,
                        'Source Scan': scan_ref(pk, 'cnv', pg['scan_page'])})
    return out


BUILD = {'FRM': frm_rows, 'EXT': ext_rows, 'CNV': cnv_rows}


def history_rows(csv_paths):
    """Extrusion Production Record rows from the system's own past schedules (history/prod_instr.py; James Kuo,
    30 Sep 2026: "i got the past production schedule ... refine your data base"). A line printed again later the same
    day (061620-2.pdf ...) is a revision: per day and line, the latest print is the day's schedule."""
    import csv
    by = {}
    for csv_path in csv_paths:
        config.record_read(csv_path, 'system schedule history (Production Instruction PDFs)')
        for r in csv.DictReader(open(csv_path, encoding='utf-8')):
            by.setdefault((r['date'], r['line']), {}).setdefault((r['run_time'], r['file']), []).append(r)
    out = {}
    for (d, line), prints in sorted(by.items()):
        for r in prints[max(prints)]:
            si = r['special'] or ''
            m = PLTS_DONE.search(si)
            out.setdefault(d, []).append({
                'Schedule Date': datetime.date.fromisoformat(d), 'Line Code': line, 'Order': r['order'],
                'Product Code': r['prod_code'], 'Total Sheets': num(r['total_sheets']), 'Weight (LBs)': num(r['weight_lbs']),
                'Plts Done (EXT)': num(m.group(1)) if m else None, 'Plts Ordered': as_int_or_text(r['plts']),
                'Special Instructions': si or None, 'Handwritten': r.get('handwritten') or None,
                'Source Scan': f"Production Instruction/{r['file']} p{r['page']} " + (
                    f"({r['source']})" if r.get('source') else f"(system PDF, run {r['run_date']} {r['run_time']})")})
    return out


HIST_SRC = 'Production Instruction/'


def main_history(csv_paths, replace=False):
    """Append the system's past schedules to the Extrusion Production Record: days not recorded yet are added; a day
    already recorded (from the daily scan) stays as issued, and its line/order/product rows are compared and reported.
    replace: the rows an earlier history run added (Source Scan 'Production Instruction/...') are taken out first and
    built again from the CSVs (30 Sep 2026: the first run counted only the first cut row's sheets and lost instructions
    that run on to the next page). Rows from the daily scans are never touched."""
    name, sheet = RECORDS['EXT']
    header = cols(name, sheet)
    existing, pub = load_existing(name, sheet)
    n_before, replaced = len(existing), 0
    if replace:
        keep = [r for r in existing if not str(r.get('Source Scan') or '').startswith(HIST_SRC)]
        replaced, existing = len(existing) - len(keep), keep
    have = {}
    for r in existing:
        have.setdefault(norm(r.get(header[0])), []).append(r)
    hist = history_rows(csv_paths)
    added, log, rows = 0, [], []
    for d, new in sorted(hist.items()):
        if d in have:
            k = lambda rs: sorted((norm(r['Line Code']), norm(r['Order']), norm(r['Product Code'])) for r in rs)
            same = k(have[d]) == k(new)
            log.append(f"{d}: already recorded from the daily scan ({len(have[d])} rows) - kept as issued; the system's "
                       f"schedule has {'the same' if same else 'DIFFERENT'} lines, orders and products")
            continue
        rows += new
        added += len(new)
    rows = sorted(rows + existing, key=lambda r: norm(r.get(header[0])))
    days = sorted(set(hist) - set(have))
    log.insert(0, f'{len(days)} days appended from the system history ({days[0]} to {days[-1]}), {added} rows' if days else 'no new days')
    if replace:
        log.insert(1, f'{replaced} rows added by the earlier history run taken out and built again (Total Sheets = every cut row, '
                      'as printed in the line totals; instructions that run on to the next page kept)')
    notes = [['Workbook', f'{name}: {schema.BY_NAME[name].purpose}.'],
             ['Rule', 'Append-only. Rows are added, never edited (db/schema.py). Values as printed.'],
             ['Based on', f'{pub} ({n_before} rows)' if pub else 'nothing published yet: started empty'],
             ['Rows', len(rows)],
             ['Days in record', f"{len({norm(r[header[0]]) for r in rows})} ({norm(rows[0][header[0]])} to {norm(rows[-1][header[0]])})"],
             ['System history', "Days before the pipeline's first daily scan come from the system's own schedules "
                                '(Technical Engineering Team/Production Instruction PDFs, read as text: '
                                'history/prod_instr.py). Source Scan names the PDF, page and print time. A line printed '
                                'again later the same day is the revision: the latest print of each line is recorded. '
                                'Total Sheets = all the order\'s cut rows; every line agrees with the line total printed '
                                'on the report. Days kept only as image scans (8 May 2020, 13-14 Nov 2023) were read by '
                                'the scan reader and checked; Source Scan says so.'],
             ['Plts Ordered', '# Plt as printed. The field caps at 999; a hand-corrected count is in Handwritten.']]
    notes += [['This run', x] for x in log]
    notes += [['Run', datetime.date.today().isoformat()]]
    out = write(name, sheet, header, rows, notes)
    print(out, '-', log[0], '|', sum('DIFFERENT' in x for x in log), 'recorded day(s) differ from the system')
    for x in log[1:]:
        print('  ', x)


# ---- workbook in / out ----------------------------------------------------------------------------------------
def cols(name, sheet):
    return [c.name for c in next(s for s in schema.BY_NAME[name].sheets if s.name == sheet).cols]


def norm(v):
    if isinstance(v, datetime.datetime):
        v = v.date() if v.time() == datetime.time() else v
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return '' if v is None else str(v)


def load_existing(name, sheet):
    pub = config.published_path(name)
    if not (pub and pub.exists()):
        return [], None
    config.record_read(pub, 'published record')
    ws = load_workbook(pub, read_only=True)[sheet]
    rows = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(rows)]
    return [dict(zip(header, r)) for r in rows if any(v not in (None, '') for v in r)], pub


def write(name, sheet, header, rows, notes):
    F = 'Arial'
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = Font(name=F, bold=True, color='FFFFFF', size=10), PatternFill('solid', fgColor='1F3864')
    body = Font(name=F, size=10)
    for r in rows:
        ws.append([r.get(h) for h in header])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = body
            if isinstance(c.value, datetime.date):
                c.number_format = 'yyyy-mm-dd'
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    wide = {'Special Instructions': 60, 'Material (as printed)': 36, 'Note': 40, 'Handwritten': 40, 'Source Scan': 34,
            'Extrusion Status (printed)': 16, 'Done Note': 30, 'Semi Size': 20}
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = wide.get(h, max(10, min(22, len(h) + 3)))
    rm = wb.create_sheet('Read Me')
    rm.append(['Item', 'Value'])
    for c in rm[1]:
        c.font, c.fill = Font(name=F, bold=True, color='FFFFFF', size=10), PatternFill('solid', fgColor='1F3864')
    for item in notes:
        rm.append(item)
    for row in rm.iter_rows(min_row=2):
        for c in row:
            c.font = body
    rm.column_dimensions['A'].width, rm.column_dimensions['B'].width = 30, 120
    out = config.OUTPUT_DIR / name
    wb.save(out)
    return out


def main(dates):
    if not dates:
        raise SystemExit(__doc__)
    packets = {}
    for d in dates:
        p = config.packet_path(d)
        if not p.exists():
            raise SystemExit(f'No packet for {d}: {p}')
        config.record_read(p, 'daily packet')
        packets[d] = json.loads(p.read_text(encoding='utf-8'))
    plan, stop = {}, []
    for kind, (name, sheet) in RECORDS.items():
        header = cols(name, sheet)
        existing, pub = load_existing(name, sheet)
        date_col = header[0]
        have = {}
        for r in existing:
            have.setdefault(norm(r.get(date_col)), []).append(r)
        rows, log = list(existing), []
        for d in dates:
            new = BUILD[kind](packets[d])
            if d in have:
                same = sorted(tuple(norm(r.get(h)) for h in header) for r in have[d]) == \
                       sorted(tuple(norm(r.get(h)) for h in header) for r in new)
                if same:
                    log.append(f'{d}: already recorded, identical ({len(new)} rows) - skipped')
                else:
                    stop.append(f'{name}: {d} is already recorded with different rows ({len(have[d])} there, '
                                f'{len(new)} from the packet). Past rows are never edited; nothing written.')
                continue
            if not new:
                log.append(f'{d}: no rows in the packet (e.g. no FRM pages) - nothing appended')
                continue
            rows += new
            log.append(f'{d}: {len(new)} rows appended')
        plan[kind] = (name, sheet, header, rows, log, pub, len(existing))
    if stop:
        print('STOP - nothing written:')
        for s in stop:
            print('  ' + s)
        raise SystemExit(1)
    reads = config.reads()
    for kind, (name, sheet, header, rows, log, pub, n_old) in plan.items():
        notes = [['Workbook', f'{name}: {schema.BY_NAME[name].purpose}.'],
                 ['Rule', 'Append-only. Rows are added, never edited (db/schema.py). Values as printed on the scan.'],
                 ['Based on', f'{pub} ({n_old} rows)' if pub else 'nothing published yet: started empty'],
                 ['Rows', len(rows)],
                 ['Dates in record', ', '.join(sorted({norm(r[header[0]]) for r in rows}))]]
        notes += [['This run', x] for x in log]
        if kind == 'FRM':
            notes += [['Set', 'As printed. Auger lines (all but SE24, SE42, SE43, SE61): motor speed 0-100, not %. '
                              'Weight lines: Set is weight %; Weight % shows it, and the balance to 100 for Auto.'],
                      ['Variant', 'From the row note (reclaim run-out, VOIDFORM, sign blank, corn box, roll); '
                                  'first formula listed = Primary, others Other. Same rules as the FRM Draft.'],
                      ['Material ID', 'Blank until the Materials table (Formulation Master) exists.']]
        if kind == 'EXT':
            notes += [['Plts Ordered', '# Plt as printed. The field caps at 999; a hand-corrected count is in Handwritten.']]
        notes += [['Run', datetime.date.today().isoformat()]]
        notes += [[f'Source read: {Path(k).name}', f"{v['modified']} · sha256 {v['sha256'][:16]}"]
                  for k, v in reads.items() if 'packet_' in k or Path(k).name == name]
        out = write(name, sheet, header, rows, notes)
        print(out, '-', '; '.join(log))


def main_history_cnv(json_path):
    """Append converting schedules read from scans (history/cnv_scans.py: days in the daily packet's CNV layout) to the
    Converting Production Record: days not recorded yet are added with cnv_rows(); a recorded day stays as issued."""
    name, sheet = RECORDS['CNV']
    header = cols(name, sheet)
    existing, pub = load_existing(name, sheet)
    have = {norm(r.get(header[0])) for r in existing}
    config.record_read(json_path, 'converting schedules read from scans')
    days = json.loads(Path(json_path).read_text(encoding='utf-8'))
    rows, log = list(existing), []
    for pk in days:
        d = pk['packet_date']
        if d in have:
            log.append(f'{d}: already recorded - kept as issued')
            continue
        new = cnv_rows(pk)
        for r in new:
            r['Source Scan'] += ' (scan read by eye; checked against the same day\'s extrusion schedule)'
        rows += new
        log.append(f'{d}: {len(new)} rows appended from {pk["source_scan"]}')
    rows.sort(key=lambda r: norm(r.get(header[0])))
    notes = [['Workbook', f'{name}: {schema.BY_NAME[name].purpose}.'],
             ['Rule', 'Append-only. Rows are added, never edited (db/schema.py). Values as printed on the scan.'],
             ['Based on', f'{pub} ({len(existing)} rows)' if pub else 'nothing published yet: started empty'],
             ['Rows', len(rows)],
             ['Dates in record', ', '.join(sorted({norm(r[header[0]]) for r in rows}))],
             ['Older days', 'Days before the daily scans come from converting sheets kept in older copier scans, read by eye '
                            '(history/cnv_scans.py); every order and product matches an extrusion schedule.']]
    notes += [['This run', x] for x in log] + [['Run', datetime.date.today().isoformat()]]
    out = write(name, sheet, header, rows, notes)
    print(out, '-', '; '.join(log))


if __name__ == '__main__':
    if sys.argv[1:2] == ['--history-cnv']:
        main_history_cnv(sys.argv[2])
    elif sys.argv[1:2] == ['--history']:
        main_history([a for a in sys.argv[2:] if not a.startswith('--')], replace='--replace-history' in sys.argv)
    else:
        main(sys.argv[1:])
