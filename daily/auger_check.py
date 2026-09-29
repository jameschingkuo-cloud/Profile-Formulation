"""Auger check on the day's FRM Draft: does each hopper on an auger line carry a material its hardware is meant for?

    PKT_DATE=2026-09-29 python daily/auger_check.py

Reads out/FRM Draft <date>.xlsx (daily/resolve.py). For auger lines only (weight lines SE24/SE42/SE43/SE61 have no
hoppers to check), each fed hopper's material is given a role (calc/auger_rules.role) and compared with that line's
hopper rules (calc/auger_rules.RULES: DRAFT, §10 Q13, awaiting James). A material outside the hopper's allowed roles
is listed; nothing is changed. Slopes are NOT checked yet: there is no Auger Calibration master until Tech's calc
workbooks are in inputs/calc (CLAUDE.md "Never", §7.13).
"""
import os
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'calc'))
import config  # noqa: E402
from auger_rules import DOSING, RULES, role  # noqa: E402
from common import mat_key  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


def main():
    pkt = os.environ.get('PKT_DATE') or sys.exit('Set PKT_DATE')
    src = config.OUTPUT_DIR / f'FRM Draft {pkt}.xlsx'
    config.record_read(src, 'FRM Draft (auger check)')
    ws = load_workbook(src, read_only=True)['Draft']
    rows = list(ws.iter_rows(values_only=True))
    h = {k: i for i, k in enumerate(rows[0])}
    seen, findings, checked = set(), [], 0
    for r in rows[1:]:
        line, code, var, feeder, mat, st = (r[h[k]] for k in ('Line Code', 'Formula Code', 'Variant', 'Feeder', 'Material (as issued)', 'Set'))
        key = (line, code, var, feeder)
        if DOSING.get(line) != 'AUGER' or key in seen or not mat:
            continue
        seen.add(key)
        m = re.fullmatch(r'Hopper (\d)', str(feeder or '').strip())
        rule = RULES.get(line, {}).get(f'H{m.group(1)}') if m else None
        if not rule:
            continue
        checked += 1
        hw, main_role, allowed = rule
        rl = role(mat_key(mat), mat)
        if rl not in allowed:
            findings.append((line, code, var, feeder, mat, st, rl, hw, main_role, ', '.join(sorted(allowed))))
    print(f'{checked} hopper settings checked on auger lines ({len({k[:3] for k in seen})} formulas); {len(findings)} outside the draft hopper rules')
    for f in findings:
        print('  ' + ' | '.join(str(x) for x in f))
    return findings


if __name__ == '__main__':
    main()
