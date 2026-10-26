"""Tech's FRM pages, read by eye into a short text form, into the packet's FRM part (James Kuo, 8 Oct 2026: a copier scan
of Tech's issued formulation pages for 2, 5, 6 and 7 Oct, the days that came only as the system's PDFs).

    python daily/frm_from_text.py data/frm_text/frm_<date>.txt --scan <copier pdf name>          check only
    python daily/frm_from_text.py data/frm_text/frm_<date>.txt --scan <copier pdf name> --write  into the packet

The text (one file per issue date):
    ## <scan page> | <line code> | <line no> | <date as printed> | <header note> | <columns> | <effective date>
    foot: <footnote> | <footnote>          hand: <handwriting>          unclear: <what was hard to read>
    > <order>, <order>                     a group: the orders that share the formulas below
    <formula code> ; <feeder> ; <feeder> ... [; # <note>]
A feeder is "material @ set", "-" for an empty cell (Hopper pages, in column order), or "<column>=material @ set" on the
wide pages (columns not named are empty). Columns: H5 = Hopper 1-5; EX13 = Extruder A V1-V5, B V1-V4, C V1-V4 (SE24);
ABC13 = the same as "A V1".. (SE31, SE32); V9 = V1-V9 (SE42, SE43); SE61 = A 1-6, B 1, C 1-6, D 1.
A formula code printed in a form the code structure does not allow is written as printed with "!" after it: kept as
printed (never corrected or guessed) and listed for James.

Checks (nothing is changed by a check; every value stays as read from the page):
  - every order on a page is on that day's extrusion schedule on the same line, and every scheduled order of the line is
    on the page (Tech's page lists what the line runs);
  - against the day's Word formulation (out/FRM Draft <date>.xlsx): each order's formulas, feeder by feeder, the same or
    not; orders the draft left to the engineer now have Tech's formula.
"""
import argparse
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import config  # noqa: E402

COLUMNS = {
    'H5': [f'Hopper {i}' for i in range(1, 6)],
    'EX13': [f'Extruder A V{i}' for i in range(1, 6)] + [f'Extruder B V{i}' for i in range(1, 5)] + [f'Extruder C V{i}' for i in range(1, 5)],
    'ABC13': [f'A V{i}' for i in range(1, 6)] + [f'B V{i}' for i in range(1, 5)] + [f'C V{i}' for i in range(1, 5)],
    'V9': [f'V{i}' for i in range(1, 10)],
    'SE61': [f'A {i}' for i in range(1, 7)] + ['B 1'] + [f'C {i}' for i in range(1, 7)] + ['D 1'],
}
ORDER = re.compile(r'^[A-Z0-9]{7}-\d{1,2}$')
FORMULA = re.compile(r'^[A-Z]{2}[0-9A-Z]{4}[A-Z]{2}[0-9A-Z]$')     # FU0022WB3, FUA152WB4, BF0000EB5, FUA060WBA
ODD = []                                                          # (scan page, line, code) printed codes that do not fit


def feeder(cell):
    if cell.strip() == '-':
        return {'material': '', 'set': ''}
    if '@' not in cell:
        raise ValueError(f'feeder without "@": {cell!r}')
    m, s = cell.rsplit('@', 1)
    return {'material': m.strip(), 'set': s.strip()}


def parse(text):
    pages, page, group = [], None, None
    for n, raw in enumerate(text.splitlines(), 1):
        line = raw.strip()
        if not line or (line.startswith('#') and not line.startswith('## ')):
            continue
        try:
            if line.startswith('## '):
                sp, code, no, date, head, cols, eff = [x.strip() for x in line[3:].split('|')]
                page = {'scan_page': int(sp), 'title': f'Line {no} ({code}) Formulations', 'line_no': no, 'line_code': code,
                        'date': date, 'header_note': head, 'feeder_columns': COLUMNS[cols], 'footnotes': [],
                        'footer': 'Tech. Department', 'effective_date': eff, 'handwritten': '', 'unclear': '', 'groups': []}
                pages.append(page); group = None
            elif line.startswith('foot:'):
                page['footnotes'] = [x.strip() for x in line[5:].split(' | ') if x.strip()]
            elif line.startswith('hand:'):
                page['handwritten'] = line[5:].strip()
            elif line.startswith('unclear:'):
                page['unclear'] = line[8:].strip()
            elif line.startswith('>'):
                group = {'orders': [o.strip() for o in line[1:].split(',') if o.strip()], 'formulas': []}
                bad = [o for o in group['orders'] if not ORDER.match(o)]
                if bad:
                    raise ValueError(f'not an order number: {bad}')
                page['groups'].append(group)
            else:
                parts = [p.strip() for p in line.split(';')]
                note = ''
                if parts[-1].startswith('#'):
                    note = parts.pop()[1:].strip()
                code, cells = parts[0], parts[1:]
                if code.endswith('!'):            # printed so on the page, does not fit the code structure: kept, flagged
                    code = code[:-1]
                    ODD.append((page['scan_page'], page['line_code'], code))
                elif not FORMULA.match(code):
                    raise ValueError(f'not a formula code: {code!r} (if the page prints it so, end it with "!")')
                cols = page['feeder_columns']
                feeders = {c: {'material': '', 'set': ''} for c in cols}
                named = [c for c in cells if '=' in c.split('@')[0]]
                if named:
                    for c in cells:
                        col, rest = c.split('=', 1)
                        if col.strip() not in feeders:
                            raise ValueError(f'unknown column {col!r} for {page["line_code"]}')
                        feeders[col.strip()] = feeder(rest)
                else:
                    if len(cells) > len(cols):
                        raise ValueError(f'{len(cells)} feeders for {len(cols)} columns')
                    for c, cell in zip(cols, cells):
                        feeders[c] = feeder(cell)
                group['formulas'].append({'formula_code': code, 'feeders': feeders, 'note': note})
        except (ValueError, KeyError, AttributeError, TypeError) as e:
            raise SystemExit(f'line {n}: {e}\n  {raw}')
    return pages


def norm(s):
    return re.sub(r'[\s\-–—()]', '', (s or '').lower())


def tech_formulas(pages):
    """(line, order) -> [(code, {column: (material, set)})] in page order."""
    out, seen = {}, {}
    for pg in pages:
        for g in pg['groups']:
            fs = [(f['formula_code'], {c: (v['material'], v['set']) for c, v in f['feeders'].items() if v['material'] or v['set']})
                  for f in g['formulas']]
            for o in g['orders']:
                # a line on two pages (SE21, 23 Oct 2026): the second page's rows continue the order's formulas
                out[(pg['line_code'], o)] = out.get((pg['line_code'], o), []) + fs if seen.get((pg['line_code'], o)) != pg['scan_page'] else fs
                seen[(pg['line_code'], o)] = pg['scan_page']
    return out


def draft_formulas(date):
    """The day's Word formulation, from out/FRM Draft <date>.xlsx: (line, order) -> [(code, {column: (material, set)})],
    and the orders it left to the engineer."""
    import openpyxl
    f = config.OUTPUT_DIR / f'FRM Draft {date}.xlsx'
    if not f.exists():
        return None, None
    wb = openpyxl.load_workbook(f, read_only=True)
    rows = list(wb['Draft'].iter_rows(values_only=True))
    h = rows[0]
    out, to_set = {}, set()
    for r in rows[1:]:
        d = dict(zip(h, r))
        key = (d['Line Code'], d['Order'])
        if not d['Feeder']:                  # an ENGINEER TO SET row (a decided formulation with no settings on the line)
            to_set.add(key)
            continue
        fs = out.setdefault(key, {})
        col = (f"{d['Extruder']} {d['Feeder']}" if d['Extruder'] else d['Feeder']).strip()
        fs.setdefault((d['Formula Row'], d['Formula Code']), {})[col] = (d['Material (as issued)'] or '', str(d['Set'] or ''))
    exc = {(r[0], r[1]) for r in list(wb['Exceptions'].iter_rows(values_only=True))[1:]} | (to_set - set(out))
    return {k: [(code, cols) for (_, code), cols in sorted(v.items())] for k, v in out.items()}, exc


def same_cols(a, b):
    """Compare feeders by column; the draft names a wide-page column without its extruder word ('A V1')."""
    key = lambda c: c.replace('Extruder ', '')  # noqa: E731
    a2 = {key(c): (norm(m), s.strip()) for c, (m, s) in a.items()}
    b2 = {key(c): (norm(m), s.strip()) for c, (m, s) in b.items()}
    return a2 == b2, a2, b2


def check(date, pages):
    pk = json.loads(config.packet_path(date).read_text(encoding='utf-8'))
    sched = {}
    for pg in pk['ext']:
        for r in pg['rows']:
            sched.setdefault(pg['line'], []).append(r['order'])
    report = []
    on_pages = {}
    for pg in pages:
        for g in pg['groups']:
            for o in g['orders']:
                on_pages.setdefault(pg['line_code'], []).append(o)
    for line in sorted(set(sched) | set(on_pages)):
        s, p = set(sched.get(line, [])), set(on_pages.get(line, []))
        for o in sorted(p - s):
            other = [ln for ln, os_ in sched.items() if o in os_]
            report.append(('CHECK', line, o, 'on Tech page, not on the schedule' + (f' (scheduled on {other[0]})' if other else '')))
        for o in sorted(s - p):
            report.append(('info', line, o, 'scheduled, not on Tech page'))
    tech = tech_formulas(pages)
    draft, exc = draft_formulas(date)
    same = diff = new = 0
    if draft is not None:
        for key, tfs in sorted(tech.items()):
            if key in exc:
                new += 1
                report.append(('new', *key, 'left to the engineer in the Word file; Tech issued ' + ', '.join(c for c, _ in tfs)))
                continue
            dfs = draft.get(key)
            if dfs is None:
                continue
            if [c for c, _ in dfs] != [c for c, _ in tfs] and sorted(c for c, _ in dfs) != sorted(c for c, _ in tfs):
                diff += 1
                report.append(('DIFF', *key, f"formulas: Word {[c for c, _ in dfs]} vs Tech {[c for c, _ in tfs]}"))
                continue
            bad = []
            for code, tcols in tfs:
                match = [dcols for dc, dcols in dfs if dc == code]
                if not any(same_cols(tcols, dc)[0] for dc in match):
                    _, a2, b2 = same_cols(tcols, match[0] if match else {})
                    cols = sorted(set(a2) | set(b2))
                    bad.append(code + ': ' + '; '.join(f"{c} Tech {a2.get(c)} vs Word {b2.get(c)}" for c in cols if a2.get(c) != b2.get(c)))
            if bad:
                diff += 1
                report.append(('DIFF', *key, ' | '.join(bad)))
            else:
                same += 1
    return report, (same, diff, new, draft is not None)


def main(argv):
    ap = argparse.ArgumentParser()
    ap.add_argument('text'); ap.add_argument('--scan', required=True)
    ap.add_argument('--write', action='store_true'); ap.add_argument('--force', action='store_true')
    a = ap.parse_args(argv)
    t = Path(a.text)
    date = re.search(r'(\d{4}-\d{2}-\d{2})', t.name).group(1)
    pages = parse(t.read_text(encoding='utf-8'))
    groups = sum(len(p['groups']) for p in pages)
    orders = sum(len(g['orders']) for p in pages for g in p['groups'])
    print(f'{date}: {len(pages)} pages, {groups} groups, {orders} orders, '
          f'{sum(len(g["formulas"]) for p in pages for g in p["groups"])} formula rows')
    report, (same, diff, new, had) = check(date, pages)
    for r in report:
        print('  ' + ' | '.join(r))
    for sp, line, code in ODD:
        print(f'  ODD CODE | {line} | scan p{sp} | {code!r} printed; does not fit the formula code structure - for James, not corrected')
    if had:
        print(f'  vs the Word formulation: {same} orders the same, {diff} different, {new} the engineer had to complete')
    if a.write:
        pp = config.packet_path(date)
        pk = json.loads(pp.read_text(encoding='utf-8'))
        if pk.get('frm') and not a.force:
            raise SystemExit(f'{pp.name} already has FRM pages (--force to replace them before they are recorded)')
        pk['frm'], pk['frm_source_scan'] = pages, a.scan
        pp.write_text(json.dumps(pk, ensure_ascii=False, indent=1), encoding='utf-8')
        print(f'written: {pp}')


if __name__ == '__main__':
    main(sys.argv[1:])
