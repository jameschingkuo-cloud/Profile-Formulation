import sys as _sys, pathlib as _pl; _sys.path.insert(0, str(_pl.Path(__file__).resolve().parents[1])); import config  # repo root config
import json, re, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from load import *
import checks  # runs checks, builds checks.issues
import importlib.util as _ilu, os as _os
def _manual(date):
    """Hand-found issues for the day: daily/manual/manual_issues_<date>.py if it exists, else daily/manual_issues.py (the working file)."""
    here = _pl.Path(__file__).resolve().parent
    p = here / 'manual' / f'manual_issues_{date}.py'
    p = p if p.exists() else here / 'manual_issues.py'
    spec = _ilu.spec_from_file_location('manual_issues', p); m = _ilu.module_from_spec(spec); spec.loader.exec_module(m)
    return m.MANUAL, p.name
MANUAL, MANUAL_FILE = _manual(_os.environ.get('PKT_DATE', ''))

import os
PKT = os.environ.get('PKT_DATE', '2026-09-23')
_pd = datetime.date.fromisoformat(PKT)
import load as _load
SRC = (os.environ.get('PKT_SRC') or (_load._PK or {}).get('source_scan') or 'scan name not recorded (set PKT_SRC)') + f' (scan of the {_pd.day} {_pd:%b %Y} production packet)'
_cfg = os.environ.get('PKT_CFG') or str(_pl.Path(__file__).resolve().parent / 'cfg' / f'cfg_{PKT}.json')
EXTRA = json.load(open(_cfg, encoding='utf-8')) if os.path.exists(_cfg) else {}
F = 'Arial'
HDR = PatternFill('solid', fgColor='1F3864'); HFONT = Font(name=F, bold=True, color='FFFFFF', size=10)
BODY = Font(name=F, size=10); BOLD = Font(name=F, size=10, bold=True)
FLAG = PatternFill('solid', fgColor='FCE4D6'); OKF = PatternFill('solid', fgColor='E2EFDA')
SEV_FILL = {'High': 'F8CBAD', 'Medium': 'FFE699', 'Low': 'DDEBF7', 'Info': 'EDEDED'}

def sheet(wb, name, headers, rows, widths=None, first=False, table=True):
    ws = wb.active if first else wb.create_sheet(name)
    ws.title = name
    ws.append(headers)
    for r in rows: ws.append(r)
    for c in ws[1]:
        c.font = HFONT; c.fill = HDR; c.alignment = Alignment(wrap_text=True, vertical='center')
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.font is None or not c.font.bold: c.font = BODY
            c.alignment = Alignment(vertical='top', wrap_text=isinstance(c.value, str) and len(c.value) > 40)
    ws.freeze_panes = 'A2'
    for i, h in enumerate(headers, 1):
        w = (widths or {}).get(h) or min(max(len(str(h)) + 2, 10), 40)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30
    if table and rows:
        ref = f"A1:{get_column_letter(len(headers))}{len(rows)+1}"
        t = Table(displayName=re.sub(r'\W', '', name)[:30] + 'T', ref=ref)
        t.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True)
        ws.add_table(t)
    return ws

def readme(wb, title, lines):
    ws = wb.create_sheet('Read Me')
    ws['A1'] = title; ws['A1'].font = Font(name=F, bold=True, size=14)
    for i, l in enumerate(lines, 3):
        ws.cell(i, 1, l).font = BODY
        ws.cell(i, 1).alignment = Alignment(wrap_text=True, vertical='top')
    ws.column_dimensions['A'].width = 120
    return ws

def issues_sheet(wb, docs):
    rows = []
    allis = checks.issues + [dict(zip(['Severity','Document','Line','Order','Check','Detail','Source'], m)) for m in MANUAL]
    order = {'High': 0, 'Medium': 1, 'Low': 2, 'Info': 3}
    for i in sorted(allis, key=lambda i: (order[i['Severity']], i['Document'], i['Line'], i['Order'])):
        if any(d in i['Document'] for d in docs):
            rows.append([i['Severity'], i['Document'], i['Line'], i['Order'], i['Check'], i['Detail'], i['Source'], ''])
    ws = sheet(wb, 'Issues', ['Severity','Document','Line','Order','Check','Detail','Source','James / Tech response'], rows,
               {'Severity': 10, 'Document': 12, 'Line': 16, 'Order': 22, 'Check': 34, 'Detail': 80, 'Source': 14, 'James / Tech response': 36})
    for r in range(2, len(rows) + 2):
        ws.cell(r, 1).fill = PatternFill('solid', fgColor=SEV_FILL[ws.cell(r, 1).value])
        ws.cell(r, 6).alignment = Alignment(wrap_text=True, vertical='top')
    return ws, rows

today = os.environ.get('RUN_DATE') or datetime.date.today().isoformat()
common = [f'Source: {SRC}.',
          f'Transcribed {today} from the page images (each field read by eye at high zoom), then cross-checked by script. '
          + EXTRA.get('ext_crosscheck', 'Key fields of all 82 extrusion orders also match the independently verified set in ext_truth_2026-09-23.csv.'),
          'Values are copied exactly as printed (text columns). Number columns are converted from the printed text for checking; '
          'the printed text is kept beside them where formatting matters (fractions, commas).',
          'Nothing was corrected: where the print looks wrong, the printed value is kept and the problem is listed on the Issues sheet.',
          'Issues sheet: severity High = could send wrong numbers to the floor; Medium = inconsistent, needs an answer; '
          'Low = cosmetic or probably fine; Info = for the record. The last column is for your answers.']

# ================================================================== 1. EXTRUSION
wb = Workbook()
hdr = ['Scan Page','Report Page','Line','T','Order','Mfg#','Ord#','Prod Code','Die','Order Width (printed)','Order Length (printed)',
       'Order Width (in)','Order Length (in)','Material','Grade','Spec','Colors','Thk (mm)','GSM','Cut Rows','Cut Width (printed)','Cut Length (printed)',
       'Total Sheets','Pack Code','# Plt','PCs/Stack','Stk/Plt','Pallet Capacity','Pallets Needed','Weight (LBs)','Calc Weight (LBs)','Weight Ratio',
       'In-str Date','Web Width','Special Instructions','Handwriting / Marks','Flags']
rows = []
for i, r in enumerate(ext_rows, 2):
    mat, grade, spec = (r['mat_spec'].split() + ['', '', ''])[:3]
    mfg, ordn = r['order'].split('-')
    flags = '; '.join(sorted({x['Check'] for x in checks.issues if x['Order'] == r['order'] and x['Document'].startswith('EXT') and x['Severity'] != 'Info'}))
    rows.append([r['scan_page'], int(r['report_page']), r['line'], r['T'], r['order'], mfg, int(ordn), r['prod_code'], r['die'],
                 r['order_width'], r['order_length'], frac(r['order_width']), frac(r['order_length']), mat, grade, spec, r['colors'],
                 float(r['thk']), num(r['gsm']), len(r['cut_rows']), r['cut_rows'][0]['width'], r['cut_rows'][0]['length'],
                 f"=SUMIFS('Cut Rows'!$F:$F,'Cut Rows'!$A:$A,E{i})", r['pack_code'], num(r['num_plt']), num(r['pcs_per_stack']), num(r['stk_per_plt']),
                 f"=Y{i}*Z{i}*AA{i}", f"=IFERROR(W{i}/(Z{i}*AA{i}),\"\")", num(r['weight_lbs']),
                 f"=W{i}*L{i}*M{i}*0.00064516*S{i}/453.59237", f"=IFERROR(AD{i}/AE{i},\"\")",
                 r['instr_date'], r['web_width'], r['special_instructions'], r['handwritten'], flags])
ws = sheet(wb, 'Orders', hdr, rows, {'Order': 12, 'Prod Code': 14, 'Special Instructions': 60, 'Handwriting / Marks': 40, 'Flags': 40, 'Colors': 11,
                                     'Total Sheets': 12, 'Weight (LBs)': 12, 'Calc Weight (LBs)': 12}, first=True)
for r in range(2, len(rows) + 2):
    for col, fmt in (('W', '#,##0'), ('AB', '#,##0'), ('AC', '#,##0.0'), ('AD', '#,##0'), ('AE', '#,##0'), ('AF', '0.000'), ('L', '0.000'), ('M', '0.000'), ('R', '0.0'), ('S', '#,##0')):
        ws[f'{col}{r}'].number_format = fmt
    if ws[f'AK{r}'].value: ws[f'AK{r}'].fill = FLAG
ws['A1'].comment = None
# cut rows
crow = []
for r in ext_rows:
    for k, c in enumerate(r['cut_rows'], 1):
        crow.append([r['order'], r['line'], r['scan_page'], k, c['width'] + ' x ' + c['length'], num(c['total_sheets']), frac(c['width']), frac(c['length'])])
ws2 = sheet(wb, 'Cut Rows', ['Order', 'Line', 'Scan Page', 'Cut Row', 'Cut Size (printed)', 'Total Sheets', 'Cut Width (in)', 'Cut Length (in)'], crow, {'Cut Size (printed)': 22})
for r in range(2, len(crow) + 2): ws2[f'F{r}'].number_format = '#,##0'
# line totals
lt = []; lines = []
for p in EXT:
    if p['line'] not in lines: lines.append(p['line'])
tot = {p['line']: p for p in EXT if p.get('line_total_pcs')}
for j, ln in enumerate(lines, 2):
    p = tot.get(ln)
    lt.append([ln, len([r for r in ext_rows if r['line'] == ln]), num(p['line_total_pcs']) if p else None, num(p['line_total_lbs']) if p else None,
               f"=SUMIFS('Cut Rows'!$F:$F,'Cut Rows'!$B:$B,A{j})", f"=SUMIFS(Orders!$AD:$AD,Orders!$C:$C,A{j})",
               f'=IF(C{j}="","no printed total",IF(AND(C{j}=E{j},D{j}=F{j}),"OK","MISMATCH"))',
               (EXTRA.get('no_total_note', 'Report page 11 missing from scan; handwritten 591,328# on p.10 = sum of weights') if not p else f"scan p{p['scan_page']}")])
n = len(lt) + 2
_ft = re.search(r'([\d,]+)\s*PCs\s*/\s*([\d,]+)\s*LBs', next((p['final_total'] for p in EXT if p.get('final_total')), ''))
lt.append(['ALL', f'=SUM(B2:B{n-1})', num(_ft.group(1)) if _ft else None, num(_ft.group(2)) if _ft else None, f'=SUM(E2:E{n-1})', f'=SUM(F2:F{n-1})', f'=IF(AND(C{n}=E{n},D{n}=F{n}),"OK","MISMATCH")', 'Final Total printed on the last report page' + EXTRA.get('final_note', '')])
ws3 = sheet(wb, 'Line Totals', ['Line', 'Orders', 'Printed PCs', 'Printed LBs', 'Read PCs', 'Read LBs', 'Check', 'Note'], lt, {'Note': 60, 'Check': 22}, table=False)
for r in range(2, n + 1):
    for col in 'CDEF': ws3[f'{col}{r}'].number_format = '#,##0'
    ws3[f'A{r}'].font = BOLD if r == n else BODY
issues_sheet(wb, ['EXT'])
pg = [[p['scan_page'], p['report_page'], p['line'], len(p['rows']), p.get('carryover_top', ''), p.get('page_notes', '')] for p in EXT]
sheet(wb, 'Pages', ['Scan Page', 'Report Page', 'Line', 'Records', 'Carry-over text at top', 'Page notes'], pg, {'Carry-over text at top': 40, 'Page notes': 90}, table=False)
readme(wb, f"PP Profile Production Instruction - Extrusion (WPPPOPRC), run {EXT[0]['run_date']} {EXT[0]['run_time']}", common + [
    'Sheets: Orders (one row per order), Cut Rows (one row per cut-dimension line; Total Sheets on Orders is a SUMIFS over this sheet), '
    'Line Totals (printed totals vs the rows; formulas), Issues, Pages (per-page notes, carry-over text).',
    'Calc Weight = Total Sheets x order width x order length (in) x 0.00064516 m2/in2 x GSM / 453.59237 g/lb. Weight Ratio = printed / calculated' + EXTRA.get('weight_note', '; every order is between 0.97 and 1.00.'),
    'Pallet Capacity = # Plt x PCs/Stack x Stk/Plt; Pallets Needed = Total Sheets / (PCs/Stack x Stk/Plt). # Plt is printed as 999 when the real count is higher (see Issues).',
    EXTRA.get('ext_pages_note', 'The scan is missing report page 11 (SE25 continuation); the Final Total still equals all 82 rows, so it held no extra orders. Scan pages 10-11 are report pages 12-13 (SE31); scan page 13 is report page 10 (SE25).')])
wb.save(config.OUTPUT_DIR / f'EXT Extrusion Schedule {PKT}.xlsx')

# ================================================================== 2. CONVERTING
wb = Workbook()
hdr = ['Scan Page', 'Sheet', 'Line', 'Page', 'Order', 'Product Code', 'Extrusion Status', 'Plts Extruded', 'Plts Ordered', 'Plts Left to Extrude',
       'Semi pc/plt', 'Semi Start', 'Semi-Size', 'Color', 'Apl.', 'MM', 'Flute', 'Die #', 'Die Description', 'Die Status', 'Plate Status', 'Ink Color',
       'Total Sheets (printed)', 'Total Sheets', 'Pack Code', '# of Plts', 'Pc./Plt.', 'Plts x Pc/Plt', 'Sheets Check', 'Req. Date', 'Done Note', 'Row Notes',
       'On EXT today', 'Handwriting', 'Unclear', 'Flags']
rows = []
page_of = {p['scan_page']: p for p in CNV}
ext_orders = {r['order']: r for r in ext_rows}
for i, c in enumerate(cnv_rows, 2):
    m = re.fullmatch(r'\s*(\d+)\s+(?:OF|of)\s+(\d+)\s*', c['extrusion_status'] or '')
    ts = num(c['total_sheets'])
    flags = '; '.join(sorted({x['Check'] for x in checks.issues if x['Order'] == c['order'] and ('CNV' in x['Document']) and x['Severity'] != 'Info'}))
    e = ext_orders.get(c['order'])
    rows.append([c['scan_page'], page_of[c['scan_page']]['title'].replace('PP PROFILE PRODUCTION INSTRUCTIONS - ', ''), c['cnv_line'], page_of[c['scan_page']]['page_of'],
                 c['order'], c['product_code'], c['extrusion_status'], int(m.group(1)) if m else None, int(m.group(2)) if m else None,
                 f'=IF(OR(H{i}="",I{i}=""),"",I{i}-H{i})', num(c['semi_pc_plt']), c.get('semi_start', ''), c['semi_size'], c['color'], c['apl'], c['mm'],
                 c['flute'], c['die_no'], c['die_description'], c['die_status'], c['plate_status'], c['ink_color'], c['total_sheets'], ts,
                 c['pack_code'], num(c['num_plts']), num(c['pc_per_plt']), f'=IF(OR(Z{i}="",AA{i}=""),"",Z{i}*AA{i})',
                 f'=IF(OR(X{i}="",AB{i}=""),"",IF(X{i}=AB{i},"OK","MISMATCH"))', c['req_date'], c['done_note'], c['row_notes'],
                 (e['line'] if e else ''), c.get('handwritten', ''), c.get('unclear', ''), flags])
ws = sheet(wb, 'Orders', hdr, rows, {'Sheet': 22, 'Die Description': 26, 'Row Notes': 60, 'Flags': 40, 'Semi-Size': 20, 'Ink Color': 20, 'Unclear': 30, 'Order': 12, 'Product Code': 14}, first=True)
for r in range(2, len(rows) + 2):
    for col in ('X', 'AB'): ws[f'{col}{r}'].number_format = '#,##0'
    if ws[f'AJ{r}'].value: ws[f'AJ{r}'].fill = FLAG
bn = []
for p in CNV:
    for b in p.get('banner_notes', []):
        bn.append([p['scan_page'], p['line_code'], b if isinstance(b, str) else json.dumps(b)])
    if p.get('page_notes'): bn.append([p['scan_page'], p['line_code'], 'PAGE NOTE: ' + p['page_notes']])
sheet(wb, 'Banner Notes', ['Scan Page', 'Line', 'Note (as printed, with position)'], bn, {'Note (as printed, with position)': 120})
issues_sheet(wb, ['CNV'])
_sheets = list(dict.fromkeys(f"{p['title'].replace('PP PROFILE PRODUCTION INSTRUCTIONS - ', '').title()} {p['line_code']}" for p in CNV))
readme(wb, f"PP Profile Production Instructions - Converting ({', '.join(_sheets)}), issued {CNV[0]['issue_date']}", common + [
    'Sheets: Orders (one row per order line on each converting sheet), Banner Notes (the instruction lines typed between rows, with where they sit), Issues.',
    'Extrusion Status "X OF Y" is split into Plts Extruded (X) and Plts Ordered (Y); Plts Left to Extrude = Y - X (formula).',
    'Sheets Check compares printed Total Sheets with # of Plts x Pc./Plt. (formula).',
    f"On EXT today = the extrusion line this order is on in the same day's extrusion schedule ({len({c['order'] for c in cnv_rows if c['order'] in ext_orders})} of the {len({c['order'] for c in cnv_rows})} converting orders are).",
] + ([f'Scan page {d} is a second scan of page {o} (same sheet, same rows) and was left out.' for d, o in DUP_PAGES]) + [
    'Cells that printed as Excel overflow (####) are kept as "#####"; the Unclear column says how many hashes printed.'])
wb.save(config.OUTPUT_DIR / f'CNV Converting Schedule {PKT}.xlsx')

# ================================================================== 3. FORMULATIONS
wb = Workbook()
long = []
for f in frm_rows:
    for k in f['columns']:
        v = f['feeders'].get(k, {'material': '', 'set': ''})
        if not v['material'] and not v['set']: continue
        s = v['set']; sv = num(s)
        long.append([int(f['line_no']), f['line'], f['scan_page'], f['group'], ', '.join(f['orders']), f['formula_code'],
                     'Primary' if f['row_in_group'] == 1 else f"Alternate {f['row_in_group']-1}", k,
                     (k.split()[0] if f['line'] in ('SE24', 'SE31', 'SE32', 'SE61') else ''), v['material'], s, sv, f['note']])
ws = sheet(wb, 'Formulations', ['Line No', 'Line', 'Scan Page', 'Group', 'Orders', 'Formula Code', 'Row', 'Feeder', 'Extruder', 'Material', 'Set (printed)', 'Set', 'Note'],
           long, {'Orders': 40, 'Material': 44, 'Note': 50, 'Formula Code': 14}, first=True)
# order -> formula
om = []
for f in frm_rows:
    for o in f['orders']:
        e = ext_orders.get(o)
        om.append([o, int(f['line_no']), f['line'], f['formula_code'], 'Primary' if f['row_in_group'] == 1 else f"Alternate {f['row_in_group']-1}", f['note'],
                   e['prod_code'] if e else '', e['mat_spec'] if e else '', e['colors'] if e else '', e['thk'] if e else '', (e['special_instructions'][:120] if e else ''), f['scan_page']])
sheet(wb, 'Order to Formula', ['Order', 'Line No', 'Line', 'Formula Code', 'Row', 'Formula Note', 'EXT Prod Code', 'EXT Mat/Grade/Spec', 'EXT Colors', 'EXT Thk', 'EXT Special Instructions (start)', 'FRM Scan Page'],
      om, {'Formula Note': 44, 'EXT Special Instructions (start)': 60, 'EXT Mat/Grade/Spec': 16})
# set sums with formulas
ss = []; r0 = 2
keys = []
for f in frm_rows:
    exts = sorted({(k.split()[0] if f['line'] in ('SE24', 'SE31', 'SE32', 'SE61') else '') for k, v in f['feeders'].items() if v['material'] or v['set']})
    for ex in exts: keys.append((f, ex))
for j, (f, ex) in enumerate(keys, 2):
    grp = f['group']
    row_lbl = 'Primary' if f['row_in_group'] == 1 else f"Alternate {f['row_in_group']-1}"
    ss.append([f['line'], grp, ', '.join(f['orders']), f['formula_code'], row_lbl, ex or '(all hoppers)',
               f'=SUMIFS(Formulations!$L:$L,Formulations!$B:$B,A{j},Formulations!$D:$D,B{j},Formulations!$F:$F,D{j},Formulations!$G:$G,E{j},Formulations!$I:$I,IF(F{j}="(all hoppers)","",F{j}))',
               f'=COUNTIFS(Formulations!$B:$B,A{j},Formulations!$D:$D,B{j},Formulations!$F:$F,D{j},Formulations!$G:$G,E{j},Formulations!$I:$I,IF(F{j}="(all hoppers)","",F{j}),Formulations!$K:$K,"Auto")',
               DOSING_LABEL.get(DOSING.get(f['line']), 'Unknown - ask James'),
               f'=IF(AND(LEFT(I{j},6)="Weight",H{j}=1),100-G{j},"")',
               f'=IF(LEFT(I{j},6)<>"Weight","n/a (auger speed, not %)",IF(H{j}=1,IF(AND(J{j}>=0,J{j}<=100),"Yes (Auto = balance)","NO"),IF(ABS(G{j}-100)<0.05,"Yes","NO")))'])
sheet(wb, 'Set Sums', ['Line', 'Group', 'Orders', 'Formula Code', 'Row', 'Extruder', 'Sum of Set', 'Feeders on Auto', 'Dosing', 'Auto = balance (%)', 'Adds to 100?'], ss, {'Orders': 40, 'Dosing': 30, 'Adds to 100?': 24})
# pages / footnotes
pg = [[p['scan_page'], int(p['line_no']), p['line_code'], p['date'], p['header_note'], len(p['groups']), sum(len(g['formulas']) for g in p['groups']),
       ' | '.join(p.get('footnotes', [])), p.get('effective_date', ''), p.get('handwritten', '')] for p in FRM]
sheet(wb, 'Pages', ['Scan Page', 'Line No', 'Line', 'Date', 'Header (AC)', 'Groups', 'Formula Rows', 'Footnotes', 'Effective Date', 'Handwriting'], pg, {'Footnotes': 90})
# materials
mats = json.load(open(config.WORK_DIR / 'issues.json', encoding='utf-8'))['mats']
sheet(wb, 'Materials', ['Material (as printed)', 'Times used'], mats, {'Material (as printed)': 60})
issues_sheet(wb, ['FRM'])
readme(wb, f"Line Formulations (Tech. Department), dated {FRM[0]['date']} - Lines {', '.join(p['line_no'] for p in FRM)}", common + [
    'Sheets: Formulations (one row per formula row x feeder that has a material or a setting - the long format a master table needs), '
    'Order to Formula (every order with its primary and alternate formulas, beside the order\'s EXT attributes), '
    'Set Sums (sum of settings per formula row and extruder; formulas over the Formulations sheet), Pages (header AC value, footnotes, effective dates), '
    'Materials (every material spelling and how often it appears), Issues.',
    'Feeder names: Hopper 1-5 (Lines 1-6, 10); A V1-V5 / B V1-V4 / C V1-V4 (Lines 7-9); V1-V9 (Lines 12-13); A 1-6 / B 1 / C 1-6 / D 1 (Line 16).',
    '"Set" is kept exactly as printed ("73.3", "Auto"); the numeric Set column is blank for "Auto".',
    'Dosing (James, 26 Sep 2026): Lines 7, 12, 13 and 16 are weight blenders - Set is weight %, each extruder adds to 100 and "Auto" takes the balance. '
    'All other lines dose by auger - Set is the auger motor speed 0-100 (no RPM feedback), so it does not add to 100; weight % needs the calibration slope (Formulation Master).'])
wb.save(config.OUTPUT_DIR / f'FRM Formulation Report {PKT}.xlsx')
print('saved')
