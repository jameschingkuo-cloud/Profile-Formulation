"""How close an FRM Draft is to what Tech issued, order by order (James Kuo, 26 Oct 2026: "do a comparison see how close
are we"). Used by ui/build.py (the page's numbers when the day has Tech's pages) and work/compare_draft_tech.py (every day,
the draft replayed as that morning).

Formulation: the draft's Primary against Tech's first formula, under James's names - a code Tech printed counts as the
same when it is an old name of the Main or one of its printed names (the same composition). Not the same: "Tech's Primary
is our Secondary" (R9 prints both; Tech printed only the run-out one), "James replaced the code Tech printed" (another
composition, his cleanup of 23 Oct), "VOIDFORM rule (R10: no reclaim)", "different", or no formulation (an Exception).
Settings: each formula in both, every feeder's material and set; a spelling difference (dashes, spaces, 6502A / F6502A,
Q1203K / F1102K for F1203K) is not a difference; a formulation printed ENGINEER TO SET counts on its own."""
import re
from collections import defaultdict
from functools import lru_cache

import assign
from db.seed_master import MAP
from openpyxl import load_workbook
from resolve import split_feeder

SAME = ('same', 'same Primary, other members')
BY_JAMES = ('James replaced the code Tech printed', 'VOIDFORM rule (R10: no reclaim)')


@lru_cache(maxsize=None)
def MAIN():
    return {m['code']: m for m in assign.rules()['mains']}


def num(s):
    try:
        return float(str(s).strip().rstrip('%'))
    except ValueError:
        return None


def same_code(d, t):
    """Draft code d (James's name) and Tech's printed code t are one formulation: the same code, Tech's an old name of it,
    or one of the Main's printed names (the same composition)."""
    m = MAIN().get(d)
    return d == t or assign.current_name(t) == d or bool(m and t in m['printed'])


def replaced(d, t):
    """Tech printed a code James replaced by d, a Main of another composition (his cleanup of 23 Oct 2026)."""
    m = MAIN().get(d)
    return bool(m and t in m.get('replaces', set()))


def read_draft(path):
    """FRM Draft workbook -> ({(line, order): [formula {'code', 'how', 'feeders'}]}, exceptions, orders by rule, to set)."""
    wb = load_workbook(path, read_only=True)
    rows = list(wb['Draft'].iter_rows(values_only=True)); h = {k: i for i, k in enumerate(rows[0])}
    out = defaultdict(dict)                 # (line, order) -> {row: {'code', 'how', 'feeders'}}
    for r in rows[1:]:
        f = out[(r[h['Line Code']], r[h['Order']])].setdefault(r[h['Formula Row']], {
            'code': r[h['Formula Code']], 'how': r[h['How Resolved']] or '', 'feeders': {}})
        f['feeders'][(r[h['Extruder']] or '', r[h['Feeder']])] = (r[h['Material (as issued)']] or '', r[h['Set']] or '')
    exc = {(r[0], r[1]): r[8] for r in list(wb['Exceptions'].iter_rows(values_only=True))[1:]}
    new = {(r[0], r[1]) for r in list(wb['By Rule'].iter_rows(values_only=True))[1:]}
    to_set = defaultdict(list)
    if 'Settings To Set' in wb.sheetnames:
        for r in list(wb['Settings To Set'].iter_rows(values_only=True))[1:]:
            to_set[(r[0], r[1])].append(r[3])
    return {k: [v[i] for i in sorted(v)] for k, v in out.items()}, exc, new, to_set


REPLACED = {'50-3963-301': '50-3963-019'}      # Q1203K / F1102K printed: F1203K is the material (James, 30 Sep 2026)


def same_mat(a, b):
    """One material however it is spelled: the Formulation Master's spelling map, else the text without dashes / spaces."""
    ia, ib = MAP.get(a), MAP.get(b)
    if ia and ib:
        return REPLACED.get(ia, ia) == REPLACED.get(ib, ib)
    n = lambda s: re.sub(r'[\s\-–]', '', s or '').lower()     # noqa: E731
    return n(a) == n(b)


def settings_diff(df, tf):
    """-> (n feeders compared, n different, text of the differences). A spelling difference is not a difference."""
    keys = set(df['feeders']) | set(tf)
    diff = []
    for k in sorted(keys):
        a, b = df['feeders'].get(k, ('', '')), tf.get(k, ('', ''))
        if a == b:
            continue
        na, nb = num(a[1]), num(b[1])
        if same_mat(a[0], b[0]) and (na == nb if na is not None else a[1] == b[1]):
            continue
        col = ' '.join(x for x in k if x)
        diff.append(f"{col}: {a[0] or '-'} @ {a[1] or '-'} -> {b[0] or '-'} @ {b[1] or '-'}")
    return len(keys), len(diff), '; '.join(diff)


def compare(day, P, draft_path):
    """Every scheduled order on Tech's FRM pages of the day, the draft against Tech's page (see the module docstring)."""
    draft, exc, new, to_set = read_draft(draft_path)
    tech = {}
    for pg in P['frm']:
        for g in pg['groups']:
            fl = [{'code': f['formula_code'], 'note': f.get('note') or '',
                   'feeders': {split_feeder(c): (v['material'], v['set']) for c, v in f['feeders'].items() if v['material'] or v['set']}}
                  for f in g['formulas']]
            for o in g['orders']:
                tech.setdefault((pg['line_code'], o), fl)
    out = []
    for e in P['ext']:
        for r in e['rows']:
            k = (e['line'], r['order'])
            t = tech.get(k)
            if t is None:
                continue                                    # not on Tech's pages (dropped, or a page missing)
            d = draft.get(k)
            kind = 'new order (by rule)' if k in new else ('repeat order' if d else 'new order')
            row = {'day': day, 'line': k[0], 'order': k[1], 'product': r['prod_code'], 'kind': kind,
                   'tech': ', '.join(f['code'] for f in t), 'draft': ', '.join(f['code'] for f in d) if d else '',
                   'how': (d[0]['how'] if d else (exc.get(k) or '')), 'n_feeders': 0, 'n_diff': 0, 'diff': ''}
            if not d:
                row['formulation'] = 'engineer (exception)'
                out.append(row)
                continue
            p_same = same_code(d[0]['code'], t[0]['code'])
            all_t = all(any(same_code(x['code'], y['code']) for x in d) for y in t)
            if p_same and all_t and len(d) == len(t):
                row['formulation'] = 'same'
            elif p_same:
                row['formulation'] = 'same Primary, other members'
            elif any(same_code(x['code'], t[0]['code']) for x in d):
                row['formulation'] = "Tech's Primary is our Secondary"
            elif replaced(d[0]['code'], t[0]['code']):
                row['formulation'] = 'James replaced the code Tech printed'
            elif 'void' in (MAIN().get(d[0]['code'], {}).get('quals') or ()):
                row['formulation'] = 'VOIDFORM rule (R10: no reclaim)'
            else:
                row['formulation'] = 'different'
            # settings of every formula both have: each draft formula against the closest of Tech's with that code not
            # paired yet (a code printed twice = a reclaim row and its run-out row)
            used = set()
            for x in d:
                if not x['feeders'] or all(not v[1] for v in x['feeders'].values()):
                    if any(same_code(x['code'], y['code']) for y in t):
                        row['n_diff'] += 1; row['diff'] += (' | ' if row['diff'] else '') + f"{x['code']}: ENGINEER TO SET"
                    continue
                cand = [(settings_diff(x, y['feeders']), i) for i, y in enumerate(t) if i not in used and same_code(x['code'], y['code'])]
                if not cand:
                    continue
                (n, nd, txt), i = min(cand, key=lambda c: (c[0][1], c[1]))
                used.add(i)
                row['n_feeders'] += n; row['n_diff'] += nd
                if txt:
                    row['diff'] += (' | ' if row['diff'] else '') + f"{x['code']}: {txt}"
            row['to_set'] = ', '.join(to_set.get(k, []))
            row['settings'] = ('n/a' if row['formulation'] in ('different', 'James replaced the code Tech printed') else
                               'same' if row['n_diff'] == 0 else 'differ')
            out.append(row)
    return out


def summary(rows):
    """Counts for the page: orders compared, formulation same, by James's decision, Tech's Primary is our Secondary,
    different, an Exception; of the same formulations, settings the same, to set by an engineer, different."""
    c = defaultdict(int)
    for r in rows:
        c[r['formulation']] += 1
    same_f = [r for r in rows if r['formulation'] in SAME]
    eng = sum(1 for r in same_f if 'ENGINEER TO SET' in r['diff'])
    return {'compared': len(rows), 'formula_same': len(same_f), 'by_james': sum(c[k] for k in BY_JAMES),
            'tech_secondary': c["Tech's Primary is our Secondary"], 'different': c['different'],
            'exception': c['engineer (exception)'], 'settings_same': sum(1 for r in same_f if r['settings'] == 'same'),
            'settings_to_set': eng, 'settings_differ': sum(1 for r in same_f if r['settings'] == 'differ') - eng}
