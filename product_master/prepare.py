"""Prepare the Product Master as the base table, from the data we have now.

Reads `Product Master.xlsx` (columns by name, never by position) and writes a
separate workbook, `Product Master - Prepared <date>.xlsx`. The Product Master
itself is never changed here.

What it adds:
  * Code-derived columns (material, grade, nominal thickness, colour), grey,
    marked "from code - not verified". The packet/calc columns are untouched.
  * Completeness per product (which core fields are still blank) and activity
    (from Formula Last Run), so Tech can verify the products that run first.
  * Fix List: code pattern errors, thickness conflicts, code vs data conflicts.
  * Colour Codes seen in product codes, for James/Tech to name (HANDOFF Q2).
  * Import Map: every column with its fill count and a blank "source field", ready
    for the full data pull.

Thickness rule (James Kuo, 28 Sep 2026): "we have product that is in between such
as 3.3mm. Unfortunately our product code doesnt capture that and one will have to
look at the spec." So the thickness in the code is nominal only. A difference under
1 mm is not an error; it goes on the Fix List as "confirm from spec". Neither value
is ever corrected automatically.

Usage:
    python product_master/prepare.py "<path>/Product Master.xlsx" --out out/ [--asof 2026-09-28]
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import os
import re
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

SHEET = "Product Master"
KEY = "Product Code"

# Columns the prepare step needs. Checked by name before anything is built.
REQUIRED = [KEY, "Material", "Grade", "End Use", "Spec", "Colours (3 layers)", "Thk (mm)",
            "GSM", "Width (in)", "Length (in)", "Cut Size (in)", "Formula Code(s)",
            "Formula Last Run", "Source", "Check", "Status", "Last Updated"]

# The fields a product needs before it can be the base for a formulation.
CORE = ["Material", "Grade", "Spec", "Colours (3 layers)", "Thk (mm)", "GSM",
        "Cut Size (in)", "End Use"]

# Product code shape (HANDOFF §7.6 R4): letter, material letter, grade letter,
# thickness (digit or letter, then a digit), 2-letter colour, sequence number.
CODE_RE = re.compile(r"^([A-Z])([A-Z])([A-Z])([0-9A-Z])([0-9])([A-Z]{2})(\d+)$")

# Code letter -> field value. Taken from the 53 products that carry full packet data
# (23-24 Sep 2026): every PPP product has P in position 2, every BBB product has B;
# position 3 is the grade. Anything else is left blank and reported.
MATERIAL_BY_LETTER = {"P": "PPP", "B": "BBB"}
GRADE_BY_LETTER = {"P": "P", "A": "A"}

# Colours confirmed on the extrusion schedule (HANDOFF §7.6 R7).
KNOWN_COLOURS = {"WB", "KS", "BL", "WM", "GT", "EB", "NS"}

# Thickness in the code vs the data: under this gap it is an in-between thickness
# the code can't show (confirm from spec); at or above it, the two disagree.
THK_CONFLICT_MM = 1.0

ACTIVE_DAYS = 365
RECENT_DAYS = 3 * 365

DERIVED_FILL = PatternFill("solid", fgColor="D9D9D9")
HEADER_FONT = Font(bold=True)
SEVERITY_FILL = {"High": PatternFill("solid", fgColor="F4B084"),
                 "Medium": PatternFill("solid", fgColor="FFE699"),
                 "Low": PatternFill("solid", fgColor="E2EFDA")}


def decode_thickness(t1: str, t2: str) -> float | None:
    """`40` -> 4.0, `33` -> 3.3, `A0` -> 10.0, `D0` -> 13.0 (letter = 10 + n)."""
    if not t2.isdigit():
        return None
    if t1.isdigit():
        return int(t1) + int(t2) / 10
    return 10 + (ord(t1) - ord("A")) + int(t2) / 10


def decode_code(code: str) -> dict:
    """Split a product code into what it states. Empty dict if it doesn't fit R4."""
    m = CODE_RE.match(code)
    if not m:
        return {}
    _, mat, grade, t1, t2, colour, _ = m.groups()
    return {"material": MATERIAL_BY_LETTER.get(mat, ""),
            "material_letter": mat,
            "grade": GRADE_BY_LETTER.get(grade, ""),
            "grade_letter": grade,
            "thk": decode_thickness(t1, t2),
            "colour": colour}


def suggest_code_fix(code: str) -> str:
    """A letter O typed for a zero in the thickness position (`AO` for `A0`)."""
    if len(code) > 4 and code[4] == "O":
        fixed = code[:4] + "0" + code[5:]
        if CODE_RE.match(fixed):
            return fixed
    return ""


def numbers(value) -> list[float]:
    """Numbers in a cell that may list several values (`4 | 4.3`)."""
    if value is None or value == "":
        return []
    if isinstance(value, (int, float)):
        return [float(value)]
    out = []
    for part in str(value).split("|"):
        try:
            out.append(float(part.strip().replace(",", "")))
        except ValueError:
            pass
    return out


def as_date(value) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if value:
        try:
            return dt.date.fromisoformat(str(value).strip()[:10])
        except ValueError:
            return None
    return None


def activity(last_run: dt.date | None, asof: dt.date) -> str:
    if last_run is None:
        return "Unknown"
    age = (asof - last_run).days
    if age <= ACTIVE_DAYS:
        return "Active"
    if age <= RECENT_DAYS:
        return "Recent"
    return "Dormant"


def blank(value) -> bool:
    return value is None or str(value).strip() == ""


def file_facts(path: Path) -> dict:
    data = path.read_bytes()
    return {"path": str(path), "bytes": len(data),
            "modified": dt.datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds"),
            "sha256": hashlib.sha256(data).hexdigest()[:16]}


def read_master(path: Path) -> tuple[list[str], list[dict]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        raise SystemExit(f"Pre-flight failed: no sheet '{SHEET}' in {path}")
    rows = wb[SHEET].iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    missing = [c for c in REQUIRED if c not in header]
    if missing:
        raise SystemExit(f"Pre-flight failed: columns not found by name: {missing}")
    records = []
    for row in rows:
        rec = dict(zip(header, row))
        if not blank(rec.get(KEY)):
            rec[KEY] = str(rec[KEY]).strip()
            records.append(rec)
    wb.close()
    return header, records


def analyse(records: list[dict], asof: dt.date) -> dict:
    """Per-product derived values, the fix list and the summaries."""
    derived, fixes, colours = {}, [], Counter()
    codes = {r[KEY] for r in records}
    agree = Counter()

    for r in records:
        code = r[KEY]
        d = decode_code(code)
        last_run = as_date(r.get("Formula Last Run"))
        missing = [c for c in CORE if blank(r.get(c))]
        derived[code] = {
            "Code Material": d.get("material", ""),
            "Code Grade": d.get("grade", ""),
            "Code Thk (nominal)": d.get("thk"),
            "Code Colour": d.get("colour", ""),
            "Code Fits Pattern": "Yes" if d else "No",
            "Core Fields Filled": f"{len(CORE) - len(missing)}/{len(CORE)}",
            "Missing": ", ".join(missing),
            "Activity": activity(last_run, asof),
        }

        def fix(severity, category, detail, action):
            fixes.append({"Severity": severity, "Product Code": code, "Category": category,
                          "Detail": detail, "Action": action,
                          "Activity": derived[code]["Activity"],
                          "Formula Last Run": last_run})

        if not d:
            sugg = suggest_code_fix(code)
            if sugg:
                also = " (that code is also in the master: merge the two rows)" if sugg in codes else ""
                fix("High", "Code pattern",
                    f"Letter O typed for zero in the thickness position: should be {sugg}{also}",
                    "Confirmed by James (28 Sep 2026). Correct the code in the calc sheet that carries it")
            else:
                fix("High", "Code pattern", "Code does not fit letter-letter-letter + thickness + colour + number",
                    "Check the AS400 item")
            continue

        colours[d["colour"]] += 1
        if d["material_letter"] not in MATERIAL_BY_LETTER:
            fix("Medium", "Code letter unknown", f"Position 2 is '{d['material_letter']}' (known: P = PPP, B = BBB)",
                "James/Tech: what material does this letter mean?")
        if d["grade_letter"] not in GRADE_BY_LETTER:
            fix("Medium", "Code letter unknown", f"Position 3 is '{d['grade_letter']}' (known: P, A)",
                "James/Tech: what grade does this letter mean?")

        # Code vs data, only where the data is there.
        mat = str(r.get("Material") or "").strip()
        if mat and d["material"]:
            ok = d["material"] in [m.strip() for m in mat.split("|")]
            agree["material", ok] += 1
            if not ok:
                fix("Medium", "Material: code vs data", f"Code says {d['material']}, data says {mat}",
                    "Confirm the material from the product spec")
        grade = str(r.get("Grade") or "").strip()
        if grade and d["grade"]:
            ok = d["grade"] in [g.strip() for g in grade.split("|")]
            agree["grade", ok] += 1
            if not ok:
                fix("Medium", "Grade: code vs data", f"Code says {d['grade']}, data says {grade}",
                    "Confirm the grade from the product spec")
        col = str(r.get("Colours (3 layers)") or "").split()
        if col:
            ok = d["colour"] == col[0]
            agree["colour", ok] += 1
            if not ok:
                fix("Medium", "Colour: code vs data", f"Code says {d['colour']}, layers are {' '.join(col)}",
                    "Confirm the colour from the product spec")

        thks = numbers(r.get("Thk (mm)"))
        if thks and d["thk"] is not None:
            gap = max(abs(t - d["thk"]) for t in thks)
            agree["thickness", gap < 1e-9] += 1
            shown = " | ".join(f"{t:g}" for t in thks)
            if gap >= THK_CONFLICT_MM:
                fix("High", "Thickness conflict",
                    f"Code {d['thk']:g} mm vs Thk {shown} mm ({gap:g} mm apart)",
                    "Too far apart for an in-between thickness: check the product spec and the calc block")
            elif gap > 1e-9:
                fix("Low", "Thickness: confirm from spec",
                    f"Code {d['thk']:g} mm vs Thk {shown} mm",
                    "In-between thickness the code can't show (James, 28 Sep 2026): take Thk from the product spec")

        # Checks already on the master that aren't about thickness stay open.
        for part in str(r.get("Check") or "").split(";"):
            part = part.strip()
            if part and not part.startswith("code thk"):
                fix("Medium", "Existing check", part, "Carried from the Product Master Check column")

    return {"derived": derived, "fixes": fixes, "colours": colours, "agree": agree}


def write_sheet(ws, header, rows, widths=None):
    ws.append(header)
    for c in ws[1]:
        c.font = HEADER_FONT
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, (dt.date, dt.datetime)):
                c.number_format = "yyyy-mm-dd"
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(10, min(40, len(h) + 2)))


def build(master: Path, out_dir: Path, asof: dt.date) -> Path:
    facts = file_facts(master)
    header, records = read_master(master)
    dupes = [c for c, n in Counter(r[KEY] for r in records).items() if n > 1]
    res = analyse(records, asof)
    derived = res["derived"]
    extra = list(next(iter(derived.values())).keys()) if derived else []

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    write_sheet(ws, header + extra,
                [[r.get(h) for h in header] + [derived[r[KEY]][e] for e in extra] for r in records],
                {KEY: 16, "Missing": 40})
    for i in range(len(header) + 1, len(header) + len(extra) + 1):
        ws.cell(1, i).fill = DERIVED_FILL

    order = {"High": 0, "Medium": 1, "Low": 2}
    act = {"Active": 0, "Recent": 1, "Dormant": 2, "Unknown": 3}
    fixes = sorted(res["fixes"], key=lambda f: (order[f["Severity"]], act[f["Activity"]], f["Category"], f["Product Code"]))
    fh = ["Severity", "Product Code", "Category", "Detail", "Action", "Activity", "Formula Last Run"]
    ws = wb.create_sheet("Fix List")
    write_sheet(ws, fh, [[f[h] for h in fh] for f in fixes], {"Detail": 60, "Action": 60, "Category": 26})
    for row in ws.iter_rows(min_row=2, max_col=1):
        row[0].fill = SEVERITY_FILL[row[0].value]

    vh = [KEY, "Formula Last Run", "Formula Code(s)", "Core Fields Filled", "Missing", "Status"]
    active = [r for r in records if derived[r[KEY]]["Activity"] == "Active"]
    active.sort(key=lambda r: as_date(r.get("Formula Last Run")) or dt.date.min, reverse=True)
    ws = wb.create_sheet("Verify First")
    write_sheet(ws, vh + ["Verified by", "Date"],
                [[r.get(KEY), as_date(r.get("Formula Last Run")), r.get("Formula Code(s)"),
                  derived[r[KEY]]["Core Fields Filled"], derived[r[KEY]]["Missing"], r.get("Status"), None, None]
                 for r in active], {"Formula Code(s)": 40, "Missing": 50})

    ws = wb.create_sheet("Colour Codes")
    write_sheet(ws, ["Colour Code", "Products", "On EXT list (R7)", "Meaning (James/Tech)"],
                [[c, n, "Yes" if c in KNOWN_COLOURS else "No", None] for c, n in res["colours"].most_common()],
                {"Meaning (James/Tech)": 40})

    ws = wb.create_sheet("Import Map")
    write_sheet(ws, ["Column", "Products filled", "Of", "Source field in the full data pull", "Notes"],
                [[h, sum(1 for r in records if not blank(r.get(h))), len(records), None,
                  "Derived by this step; not imported" if h in extra else None]
                 for h in header], {"Column": 26, "Source field in the full data pull": 36, "Notes": 36})

    ag = res["agree"]
    sev = Counter(f["Severity"] for f in res["fixes"])
    cat = Counter(f["Category"] for f in res["fixes"])
    acts = Counter(d["Activity"] for d in derived.values())
    ws = wb.create_sheet("Read Me", 0)
    lines = [
        ["Product Master - Prepared", None],
        [None, None],
        ["The Product Master prepared as the base table from the data on hand. The Product Master itself is not changed.", None],
        ["Grey columns on the Product Master sheet come from the product code and are NOT verified. The code's thickness is nominal: in-between thicknesses (e.g. 3.3 mm) are not in the code, so the product spec decides (James Kuo, 28 Sep 2026).", None],
        ["Draft until Tech verifies. Nothing guessed replaces a packet or calc value.", None],
        [None, None],
        ["Source file", facts["path"]],
        ["Source modified", facts["modified"]],
        ["Source bytes", facts["bytes"]],
        ["Source SHA-256 (first 16)", facts["sha256"]],
        ["As of", asof],
        ["Pre-flight", "Passed: all required columns found by name" + (f"; DUPLICATE CODES {dupes}" if dupes else "; no duplicate codes")],
        [None, None],
        ["Products", len(records)],
        ["All core fields filled", sum(1 for d in derived.values() if d["Missing"] == "")],
        ["Code fits pattern", sum(1 for d in derived.values() if d["Code Fits Pattern"] == "Yes")],
    ]
    lines += [[f"Activity: {k}", acts[k]] for k in ("Active", "Recent", "Dormant", "Unknown")]
    lines += [[f"Code agrees with data: {k}", f"{ag[k, True]} of {ag[k, True] + ag[k, False]}"]
              for k in ("material", "grade", "colour", "thickness")]
    lines += [[f"Fix List: {k}", sev[k]] for k in ("High", "Medium", "Low")]
    lines += [[f"  {k}", n] for k, n in cat.most_common()]
    lines += [[None, None],
              ["Sheets", "Product Master (+ grey derived columns) · Fix List · Verify First (products run in the last 12 months, newest first) · Colour Codes · Import Map (fill the source field for the full data pull)"]]
    for ln in lines:
        ws.append(ln)
    ws["B11"].number_format = "yyyy-mm-dd"
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 90

    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"Product Master - Prepared {asof.isoformat()}.xlsx"
    wb.save(out)
    return out


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("master", type=Path, help="Product Master.xlsx")
    ap.add_argument("--out", type=Path, default=Path(os.environ.get("OUT_DIR", "out")))
    ap.add_argument("--asof", type=dt.date.fromisoformat, default=dt.date.today())
    a = ap.parse_args(argv)
    print(build(a.master, a.out, a.asof))


if __name__ == "__main__":
    main()
