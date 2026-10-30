"""Build the Profile Formulation UI prototype (published as a claude.ai Artifact, HANDOFF §7.28).

    PKT_DATE=2026-09-29 python ui/build.py  -> out/Profile Formulation <date>.html (+ out/profile-formulation.html)
    python publish.py "Profile Formulation <date>.html"  -> Daily Formulation Report/Interface Copy/

A copy, refreshed after each daily run, not live data (James Kuo, 29 Sep 2026: "lets keep a copy create a separate
folder in the fomulation record for now to hold these file. No Tech sign off require"). Reads the packets, the
PUBLISHED Formulation Master (pre-flighted) and records, and out/FRM Draft <date>.xlsx for the draft-vs-Tech figures.
"""
import datetime
import glob
import json
import os
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / 'daily'))
sys.path.insert(0, str(ROOT / 'calc'))
import config  # noqa: E402
import product_code  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from db.seed_master import MAP  # noqa: E402

LN = {'SE11': 1, 'SE12': 2, 'SE13': 3, 'SE21': 4, 'SE22': 5, 'SE23': 6, 'SE24': 7, 'SE31': 8, 'SE32': 9, 'SE25': 10,
      'SE42': 12, 'SE43': 13, 'SE61': 16}
DOS = {'SE24': 'WEIGHT', 'SE42': 'WEIGHT', 'SE43': 'WEIGHT', 'SE61': 'WEIGHT'}


def real_pallets(r):
    """AS400 caps # Plt at 999 (James Kuo, 29 Sep 2026): above that, sheets / (pcs per stack x stacks per pallet)."""
    def n(s):
        try:
            return float(str(s).replace(',', ''))
        except (TypeError, ValueError):
            return None
    plt, pcs, stk = n(r.get('num_plt')), n(r.get('pcs_per_stack')), n(r.get('stk_per_plt'))
    sheets = sum(n(c.get('total_sheets')) or 0 for c in r.get('cut_rows', []))
    if plt == 999 and pcs and stk and sheets / (pcs * stk) > 999.5:
        return round(sheets / (pcs * stk))
    return None


def replaced():
    """checks.REPLACED (printed text -> what to load, and why), read from the source without running the checks."""
    import ast
    tree = ast.parse((ROOT / 'daily' / 'checks.py').read_text(encoding='utf-8'))
    for node in ast.walk(tree):
        if isinstance(node, ast.Assign) and any(getattr(t, 'id', '') == 'REPLACED' for t in node.targets):
            return {k: {'now': v[0], 'why': v[1]} for k, v in ast.literal_eval(node.value).items()}
    return {}


REPL_NOW = {k: v['now'] for k, v in replaced().items()}


def current(mat):
    """An obsolete number a page still prints (F1102K, Q1203K) -> the material loaded now; anything else as it is. The
    page shows only the current name (James, 26 Oct 2026: "No need to talk about replacing F1102K that number is
    obsolete, engineer just never had the chance to clean it up")."""
    r = REPL_NOW.get(mat)
    return r.split(' (')[0] if r else mat


def draft_vs_tech(day, P):
    """Orders the FRM Draft proposed, how many equal Tech's issue field for field, and the Exceptions."""
    from resolve import split_feeder
    src = config.OUTPUT_DIR / f'FRM Draft {day}.xlsx'
    if not src.exists():
        return {'drafted': 0, 'draft_match': 0, 'exceptions': 0, 'to_set': 0, 'rules': False}
    config.record_read(src, 'FRM Draft (interface copy)')
    wb = load_workbook(src, read_only=True)
    rows = list(wb['Draft'].iter_rows(values_only=True)); h = {k: i for i, k in enumerate(rows[0])}
    draft = {}
    for r in rows[1:]:
        k = (r[h['Line Code']], r[h['Order']])
        draft.setdefault(k, {}).setdefault((r[h['Formula Row']], r[h['Formula Code']]), set()).add(
            (r[h['Extruder']] or '', r[h['Feeder']], r[h['Material (as issued)']] or '', r[h['Set']] or ''))
    draft = {k: [(c, v) for (_, c), v in sorted(d.items())] for k, d in draft.items()}
    tech = {}
    for pg in P['frm']:
        for g in pg['groups']:
            for f in g['formulas']:
                s = {split_feeder(c) + (v['material'], v['set']) for c, v in f['feeders'].items() if v['material'] or v['set']}
                for o in g['orders']:
                    tech.setdefault((pg['line_code'], o), []).append((f['formula_code'], s))
    exc = len(list(wb['Exceptions'].iter_rows(values_only=True))) - 1
    to_set = len(list(wb['Settings To Set'].iter_rows(values_only=True))) - 1 if 'Settings To Set' in wb.sheetnames else 0
    rules = any('by rule' in str(r[h['How Resolved']] or '').lower() for r in rows[1:])     # the rules' draft (23 Oct 2026 on)
    # how close (James Kuo, 26 Oct 2026: "do a comparison see how close are we"): formulation under James's names, then
    # the settings of the same formulations, spelling aside (daily/compare_tech.py)
    import compare_tech
    close = compare_tech.summary(compare_tech.compare(day, P, src)) if P['frm'] else None
    return {'drafted': len(draft), 'draft_match': sum(1 for k, v in draft.items() if tech.get(k) == v), 'exceptions': exc,
            'to_set': to_set, 'rules': rules, 'close': close}


def draft_lines(day, P, meta, src=None):
    """A day with no Tech FRM pages yet (2 Oct 2026: the schedule came alone, as the system's PDFs): the lines from the
    FRM Draft - each order's formulation by rule since 23 Oct 2026 (Primary and Secondary, R9), the same as the Word file
    (daily/render_frm.py). Orders sharing the same formulas are one group; an Exception has no group: the page shows it
    as 'an engineer completes it', with the reason. A formulation whose settings are not set yet has no feeders (its note
    says ENGINEER TO SET). Feeder columns and notes from Tech's latest page for the line."""
    from resolve import split_feeder
    src = src or config.OUTPUT_DIR / f'FRM Draft {day}.xlsx'
    if not src.exists():
        raise SystemExit(f'No FRM pages and no {src.name}: run PKT_DATE={day} python daily/resolve.py first')
    config.record_read(src, 'FRM Draft (interface copy)')
    wb = load_workbook(src, read_only=True)
    rows = list(wb['Draft'].iter_rows(values_only=True))
    h = {k: i for i, k in enumerate(rows[0])}
    exc = {(r[0], r[1]): r[8] for r in list(wb['Exceptions'].iter_rows(values_only=True))[1:]}
    forms = {}                                   # (line, order) -> {formula row: {'code', 'note', 'role', 'feeders': {(ext, feeder): (mat, set)}}}
    for r in rows[1:]:
        f = forms.setdefault((r[h['Line Code']], r[h['Order']]), {}).setdefault(
            r[h['Formula Row']], {'code': r[h['Formula Code']], 'note': r[h['Note']] or '', 'role': r[h['Variant']] or '', 'feeders': {}})
        f['feeders'][(r[h['Extruder']] or '', r[h['Feeder']])] = (r[h['Material (as issued)']] or '', r[h['Set']] or '')
    lines = []
    for lc in sorted({e['line'] for e in P['ext']}, key=lambda c: LN.get(c, 99)):
        m = meta.get(lc, {})
        col_of = {split_feeder(c): c for c in m.get('cols', [])}
        groups, grp_of = [], {}
        rows_l = [r for e in sorted(P['ext'], key=lambda e: e['scan_page']) if e['line'] == lc for r in e['rows']]
        for r in rows_l:
            f = forms.get((lc, r['order']))
            if not f:
                continue
            fl = [{'code': x['code'], 'note': x['note'], 'role': x['role'],
                   'feeders': [{'col': col_of.get(k, ' '.join(p for p in k if p)), 'mat': current(mat), 'set': st, 'id': MAP.get(current(mat))}
                               for k, (mat, st) in x['feeders'].items() if mat or st]}
                  for _, x in sorted(f.items())]
            sig = json.dumps(fl, sort_keys=True)
            gi = next((i for i, g in enumerate(groups) if g['_sig'] == sig), None)
            if gi is None:
                groups.append({'_sig': sig, 'orders': [], 'formulas': fl}); gi = len(groups) - 1
            groups[gi]['orders'].append({'order': r['order'], 'prod': r['prod_code'], 'colors': r.get('colors'), 'thk': r.get('thk')})
            grp_of[r['order']] = gi
        for g in groups:
            g.pop('_sig')
        used = [f['col'] for g in groups for x in g['formulas'] for f in x['feeders']]
        cols = m.get('cols') or sorted(set(used))
        sched = [{'order': r['order'], 'prod': r['prod_code'], 'colors': r.get('colors'), 'thk': r.get('thk'), 'gsm': r.get('gsm'),
                  'size': f"{r.get('order_width', '')} × {r.get('order_length', '')}", 'die': r.get('die'),
                  'plt': r.get('num_plt'), 'lbs': r.get('weight_lbs'), 'si': r.get('special_instructions') or '',
                  'hw': r.get('handwritten') or '', 'g': grp_of.get(r['order']), 'plt_real': real_pallets(r),
                  'exc': exc.get((lc, r['order']))} for r in rows_l]
        lines.append({'code': lc, 'no': LN.get(lc), 'dosing': DOS.get(lc, 'AUGER'), 'cols': cols, 'ac': m.get('ac', ''),
                      'groups': groups, 'notes': m.get('notes', []), 'sched': sched})
    return lines


def rules_any(day, meta):
    """The rules' answer by product for an order the daily run did not see (ui/rules_any.py, run for the built day):
    {'asof', 'prod': {product: key}, 'prod61': {product: key on SE61, where it differs}, 'lines': {line: {'groups':
    [formulas], 'key': {key: group index, or the Exception's reason}}}}. Empty when rules_any.py has not run for the day."""
    base = ROOT / 'work' / 'rules_any'
    src, kf, pf = base / f'FRM Draft {day}.xlsx', base / f'keys_{day}.json', base / 'packets' / f'packet_{day}.json'
    if not (src.exists() and kf.exists() and pf.exists()):
        return {}
    K = json.loads(kf.read_text(encoding='utf-8'))
    syn = json.loads(pf.read_text(encoding='utf-8'))
    lines = {}
    for ln in draft_lines(day, syn, meta, src):
        ordv = {o['order']: o['g'] if o['g'] is not None else (o['exc'] or '') for o in ln['sched']}
        lines[ln['code']] = {'groups': [g['formulas'] for g in ln['groups']],
                             'key': {k.split('|', 1)[1]: ordv.get(o, '') for k, o in K['order'].items() if k.split('|', 1)[0] == ln['code']}}
    common = next((K['keys'][ln] for ln in sorted(K['keys']) if ln != 'SE61'), {})
    p61 = {p: s for p, s in K['keys'].get('SE61', {}).items() if common.get(p) != s}
    return {'asof': datetime.datetime.fromtimestamp(src.stat().st_mtime).strftime('%Y-%m-%d %H:%M'), 'prod': common,
            'prod61': p61, 'lines': lines}


def rules_by_date(pk, day, meta):
    """The rules' answer for every other schedule on file, for the Generate box (ui/rules_days.py writes them to
    work/rules/): date -> {'asof': when the rules ran, 'lines': {line: {'groups': [formulas], 'ord': {order: group index,
    or the Exception's reason}}}}. The built day is the page's own lines."""
    out = {}
    for d in sorted(pk):
        src = ROOT / 'work' / 'rules' / f'FRM Draft {d}.xlsx'
        if d == day or not src.exists():
            continue
        ls = {}
        for ln in draft_lines(d, pk[d], meta, src):
            ls[ln['code']] = {'groups': [g['formulas'] for g in ln['groups']],
                              'ord': {o['order']: o['g'] if o['g'] is not None else (o['exc'] or '') for o in ln['sched']}}
        out[d] = {'asof': datetime.datetime.fromtimestamp(src.stat().st_mtime).strftime('%Y-%m-%d'), 'lines': ls}
    return out


def record_rows():
    """Rows in the three published records (read where they live)."""
    out = {}
    for key, (name, sheet) in {'frm': ('Formulation Report Record.xlsx', 'Issued'), 'ext': ('Extrusion Production Record.xlsx', 'Orders by Day'),
                               'cnv': ('Converting Production Record.xlsx', 'Orders by Day')}.items():
        p = config.published_path(name)
        if p and p.exists():
            config.record_read(p, 'record (interface copy)')
            out[key] = load_workbook(p, read_only=True)[sheet].max_row - 1
        else:
            out[key] = 0
    return out


ENG_URL = 'https://github.com/tesseract-ocr/tessdata_fast/raw/main/eng.traineddata'


def eng_data():
    """out/eng-data.js: Tesseract's English 'fast' LSTM model (tesseract-ocr/tessdata_fast, Apache-2.0; the same file the
    Tesseract 5.4 install on James's PC uses), gzipped and base64'd as a script published next to the page (Artifact
    `files`). The page cannot download it itself: an artifact page may load scripts from the CDN but not fetch data
    (tested 29 Sep 2026). 'fast' reads these schedules better than 'best_int' (61 vs 49 of 76 orders on 29 Sep)."""
    import base64
    import gzip
    import urllib.request
    out = config.OUTPUT_DIR / 'eng-data.js'
    tag = '/* tessdata_fast eng */'
    if out.exists() and out.read_text(encoding='ascii', errors='ignore')[:len(tag)] == tag:
        return out
    raw = config.WORK_DIR / 'eng_fast.traineddata'
    if not raw.exists():
        urllib.request.urlretrieve(ENG_URL, raw)
    gz = gzip.compress(raw.read_bytes(), 9)
    out.write_text(tag + 'window.ENG_TRAINEDDATA_GZ_B64="' + base64.b64encode(gz).decode() + '";', encoding='ascii')
    return out


GLYPH_Q = (500.0, 38.0)          # int8 scale of the 560 shape values, uint8 scale of the 3 size/position values


def quantize_bank(X):
    """The glyph bank as bytes: per sample 560 shape values as int8 (x500, they lie within +-0.25) then the 3 size and
    position values as uint8 (x38, they lie within 0..6.7). ocr_eval.bank() can read it back to check the page's reader."""
    import numpy as np
    shape = np.clip(np.round(X[:, :560] * GLYPH_Q[0]), -127, 127).astype(np.int8).view(np.uint8)
    size = np.clip(np.round(X[:, 560:] * GLYPH_Q[1]), 0, 255).astype(np.uint8)
    return np.concatenate([shape, size], axis=1)


def dequantize_bank(Q):
    import numpy as np
    return np.concatenate([Q[:, :560].view(np.int8).astype(np.float32) / GLYPH_Q[0], Q[:, 560:].astype(np.float32) / GLYPH_Q[1]], axis=1)


def glyph_bank_js():
    """out/glyph-bank.js: the plant's glyph bank (scan_reader/glyph_bank.npz, 10,844 printed characters from the 25 and
    28 Sep scans), quantized, gzipped and base64'd as a script published next to the page, for the page's second reader
    (the same nearest-neighbour reading as ui/ocr_eval.py decode_order / decode_product)."""
    import base64
    import gzip
    import numpy as np
    out = config.OUTPUT_DIR / 'glyph-bank.js'
    z = np.load(config.GLYPH_BANK)
    X, y = z['X'].astype(np.float32), [str(v) for v in z['y']]
    assert X.shape[1] == 563 and all(len(v) == 1 for v in y)
    gz = gzip.compress(quantize_bank(X).tobytes(), 9, mtime=0)   # same bytes every build
    out.write_text('/* glyph bank */window.GLYPH_BANK=' + json.dumps({'n': int(X.shape[0]), 'd': 563, 'q': GLYPH_Q, 'labels': ''.join(y),
                   'data': base64.b64encode(gz).decode()}) + ';', encoding='ascii')
    return out


def line_layout(pk):
    """Per line: number, dosing, and the layout of Tech's latest issued page (feeder columns, AC, footnotes, effective
    date), for the Word formulation the page generates (same layout as daily/render_frm.py)."""
    out = {c: {'no': n, 'dosing': DOS.get(c, 'AUGER'), 'cols': [], 'ac': '', 'notes': [], 'effective': ''} for c, n in LN.items()}
    for d in sorted(pk):
        for pg in pk[d]['frm']:
            if pg['line_code'] in out:
                out[pg['line_code']].update(cols=pg['feeder_columns'], ac=pg.get('header_note', ''), notes=pg.get('footnotes', []),
                                            effective=pg.get('effective_date', ''))
    return out


def product_codes():
    """Product codes in the published Product Master (to flag a code the page read that the plant does not know)."""
    p = config.published_path('Product Master.xlsx')
    if not (p and p.exists()):
        return []
    config.record_read(p, 'Product Master (interface copy)')
    ws = load_workbook(p, read_only=True)['Product Master']
    return sorted({r[0] for r in ws.iter_rows(min_row=2, max_col=1, values_only=True) if r[0]})


def history_pairs(days=730):
    """Order + product pairs on the extrusion schedule in the last two years, from the published Extrusion Production
    Record (the system's own schedules since 2020 + the daily scans): the page's reader takes a row that matches one
    exactly. H/SH orders are always under two years old on a schedule (order_number.py)."""
    import datetime
    p = config.published_path('Extrusion Production Record.xlsx')
    if not (p and p.exists()):
        return []
    config.record_read(p, 'record (interface copy: order history)')
    rows = [r for r in load_workbook(p, read_only=True)['Orders by Day'].iter_rows(min_row=2, max_col=4, values_only=True) if r[0]]
    as_d = lambda v: v.date() if hasattr(v, 'date') else datetime.date.fromisoformat(str(v)[:10])
    last = max(as_d(r[0]) for r in rows)
    return sorted({f'{r[2]}|{r[3]}' for r in rows if as_d(r[0]) >= last - datetime.timedelta(days=days) and r[2] and r[3]})


def draft_forms(day):
    """{(line, order): [[code, settings], ...] in Formula Row order} from the day's issued FRM Draft ({} when none)."""
    src = config.OUTPUT_DIR / f'FRM Draft {day}.xlsx'
    if not src.exists():
        return {}
    rows = list(load_workbook(src, read_only=True)['Draft'].iter_rows(values_only=True))
    h = {k: i for i, k in enumerate(rows[0])}
    by = {}
    for r in rows[1:]:
        f = by.setdefault((r[h['Line Code']], r[h['Order']]), {}).setdefault(r[h['Formula Row']], [r[h['Formula Code']], []])
        mat, st = r[h['Material (as issued)']], r[h['Set']]
        if mat or st:
            f[1].append(feeder_txt(r[h['Extruder']], r[h['Feeder']], current(mat or ''), st))
    return {k: [[c, ' | '.join(fs)] for _, (c, fs) in sorted(v.items())] for k, v in by.items()}


def feeder_txt(ext, feeder, mat, st):
    """'A V1 6502A 45' / 'H3 PP WB Reclaim 31': one feeder of a formulation, as short as the look-up table needs."""
    if isinstance(st, float):
        st = f'{st:g}'
    pos = ' '.join(x for x in (ext or '', re.sub(r'^Hopper\s*', 'H', str(feeder or ''))) if x)
    return f'{pos} {mat} {st if st not in (None, "") else ""}'.strip()


def sched_fields(r, hist=False):
    """The production schedule's important fields for one order on one day (James Kuo, 30 Oct 2026: "i do want to see
    all the important things from production schedule"): [spec, colours, thk, gsm, size, die, sheets, plts, lbs, due, si]."""
    if hist:
        return [f"{r['mat']} {r['mat_spec']}".strip(), r['colour'], r['thk'], r['gsm'], f"{r['width']} × {r['length']}", r['die'],
                r['total_sheets'], r['plts'], r['weight_lbs'], r['instr_date'], r['special']]
    sheets = ' + '.join(c.get('total_sheets') or '' for c in r.get('cut_rows') or [] if c.get('total_sheets'))
    si = r.get('special_instructions') or ''
    if r.get('handwritten'):
        si += (' / ' if si else '') + 'handwritten: ' + r['handwritten']
    return [r.get('mat_spec') or '', r.get('colors') or '', r.get('thk') or '', r.get('gsm') or '',
            f"{r.get('order_width', '')} × {r.get('order_length', '')}", r.get('die') or '', sheets, r.get('num_plt') or '',
            r.get('weight_lbs') or '', r.get('instr_date') or '', si]


ORDER_RX = re.compile(r'S?(?:RP|H)[0-9A-Z]{2}[0-9A-Z]\d{2,3}', re.I)


def calc_orders(text):
    """Order bases in a calc block's 'Order #' as Tech writes it: 'H66A138, 139, 140 & 148' (a bare 3-digit number
    replaces the last digits), 'H66A113-1 & 2' (a bare 1-2 digit number is another suffix of the same order)."""
    out, prev = [], None
    for tok in re.split(r'[,&;/]|and', text or ''):
        tok = tok.strip().upper()
        m = ORDER_RX.match(tok)
        if m:
            prev = m.group(0)
        elif prev and re.fullmatch(r'\d{3}(-\d+)?', tok):
            prev = prev[:-3] + tok[:3]
        elif not (prev and re.fullmatch(r'-?\d{1,2}', tok)):
            continue
        out.append(prev)
    return list(dict.fromkeys(out))


def calc_history(pk, day, days=365):
    """The look-up's record before the first packet (James Kuo, 30 Oct 2026: "can you add atleast 6 months of formulation
    record in", "i am looking for historic record"): each order on the system's own EXT schedules (history/prod_instr.py;
    the image-day scans where no PDF exists) with the formulation Tech calculated for it in the line's calc workbook
    (Formulation Calc Library 'Calc History', one block per order: code, run period, feeder settings). An order Tech's
    blocks do not name gets the block Tech last calculated for the product on that line before it ran, marked so; none
    = no record (SE42 has no workbook). -> [(date, line, order, prod, [[code, settings]], schedule fields)], one per day"""
    import csv
    lib = config.OUTPUT_DIR / 'Formulation Calc Library.xlsx'
    if not lib.exists():
        return []
    config.record_read(lib, "Tech's calc history (look-up)")
    ws = load_workbook(lib, read_only=True, data_only=True)['Calc History']
    it = ws.iter_rows(values_only=True)
    H = {k: i for i, k in enumerate(next(it))}
    blocks = {}
    for r in it:
        k = (r[H['Workbook']], r[H['Sheet']], r[H['Row']])
        if k not in blocks:
            run = r[H['Run Date (est.)']]
            blocks[k] = {'line': r[H['Line']], 'code': (r[H['Formula Code']] or '').strip(), 'orders': calc_orders(r[H['Orders']]),
                         'prods': re.findall(r'[A-Z]{3}[0-9A-Z]{2}[A-Z]{2}\d+', (r[H['Product(s)']] or '').upper()),
                         'run': run.date().isoformat() if hasattr(run, 'date') else '', 'feeders': []}
        if r[H['Setting']] not in (None, ''):
            blocks[k]['feeders'].append(feeder_txt(r[H['Extruder']], r[H['Feeder']], r[H['Material']], r[H['Setting']]))
    by_order, by_prod = {}, {}
    for b in blocks.values():
        if not b['code']:
            continue
        for o in b['orders']:
            by_order.setdefault((b['line'], o), []).append(b)
        for p in b['prods']:
            by_prod.setdefault((b['line'], p), []).append(b)
    first = min(pk)
    start = (datetime.date.fromisoformat(day) - datetime.timedelta(days=days)).isoformat()
    rows, pdf_days = [], set()
    for name in ('ext_history.csv', 'ext_history_scans.csv'):
        p = config.WORK_DIR / 'history' / name
        if not p.exists():
            continue
        got = [r for r in csv.DictReader(open(p, encoding='utf-8')) if start <= r['date'] < first]
        rows += [r for r in got if r['date'] not in pdf_days]
        pdf_days |= {r['date'] for r in got}
    res = []
    for r in rows:
        base = (r['order_base'] or r['order'].split('-')[0]).upper()
        bl = by_order.get((r['line'], base), [])
        if not bl:
            # the product's last calc on this line on or before the run (blocks with no date count as before)
            prior = sorted((b for b in by_prod.get((r['line'], r['prod_code']), []) if (b['run'] or '') <= r['date']), key=lambda b: b['run'] or '')
            bl = [b for b in prior if b['run'] == prior[-1]['run']] if prior else []
        forms = {}
        for b in sorted(bl, key=lambda b: b['run'] or ''):
            rec = ' | '.join(b['feeders'])
            forms[(b['code'], rec)] = [b['code'], rec]     # every distinct setting Tech wrote for the order
        res.append((r['date'], r['line'], r['order'], r['prod_code'], list(forms.values()), sched_fields(r, hist=True)))
    return res


def lookup_runs(pk, day):
    """The Look up tab: every order run on record for a year, one row per order + line + formulation over its run days,
    with the schedule's fields as of its last day. Packets: the formulation issued that day (Tech's page, else our FRM
    Draft); before the first packet: Tech's calc workbook (calc_history). -> {'runs': [[order, line, prod, first, last,
    days, [[code, settings index]], schedule fields]], 'recipes': [settings text]}"""
    days = calc_history(pk, day)
    for d, p in sorted(pk.items()):
        ext = {(e['line'], r['order']): r for e in p['ext'] for r in e['rows']}
        forms = {}
        for pg in p['frm']:
            for g in pg['groups']:
                fl = []
                for f in g['formulas']:
                    fs = []
                    for c, v in f['feeders'].items():
                        if v['material'] or v['set']:
                            ext_, feed = c.split(' ', 1) if re.match(r'^[A-D] ', c) else ('', c)
                            fs.append(feeder_txt(ext_, feed, current(v['material']), v['set']))
                    fl.append([f['formula_code'], ' | '.join(fs)])
                for o in g['orders']:
                    forms[(pg['line_code'], o)] = fl
        if not p['frm']:
            forms = draft_forms(d)
        for (lc, o), r in ext.items():
            days.append((d, lc, o, r.get('prod_code') or '', forms.get((lc, o), []), sched_fields(r)))
    ri, out = {}, {}
    for d, lc, o, prod, fl, sf in sorted(days, key=lambda x: x[0]):
        fl = [[c, ri.setdefault(rec, len(ri))] for c, rec in fl]
        k = (o, lc, prod, json.dumps(fl))
        run = out.setdefault(k, [o, lc, prod, d, d, 0, fl, sf])
        run[4], run[5], run[7] = d, run[5] + 1, sf
    return {'runs': list(out.values()), 'recipes': list(ri)}


def approved_decisions(sheet):
    """Engineer decisions (APPROVED Product to Formula rows with that line's settings), as daily/resolve.py uses them:
    {'LINE|PRODUCT': [formula, ...]} in run order (Primary first)."""
    text = {m['Material ID']: (m.get('FRM Text') or m.get('Name') or m['Material ID']) for m in sheet('Materials')}
    sets = {}
    for s in sheet('Line Settings'):
        sets.setdefault((s['Line Code'], s['Formula Code'], s['Variant']), []).append(s)
    out = {}
    for p in sheet('Product to Formula'):
        k = (p['Line Code'], p['Formula Code'], p['Variant'])
        if p['Status'] != 'Approved' or k not in sets:
            continue
        out.setdefault(f"{p['Line Code']}|{p['Product Code']}", []).append((p['Priority'] != 'Primary', {
            'code': p['Formula Code'], 'note': p['Variant'] if p['Variant'] != 'Primary' else '',
            'feeders': [{'col': (s['Extruder'] + ' ' + s['Feeder']).strip() if s['Extruder'] else s['Feeder'],
                         'mat': text.get(s['Material ID'], s['Material ID']), 'set': s['Set'], 'id': s['Material ID']} for s in sets[k]]}))
    return {k: [f for _, f in sorted(v, key=lambda x: x[0])] for k, v in out.items()}


def main():
    pk = {}
    for f in sorted(glob.glob(str(config.PACKETS_DIR / 'packet_*.json'))):
        d = json.load(open(f, encoding='utf-8'))
        pk[d['packet_date']] = d
    day = os.environ.get('PKT_DATE') or max(d for d in pk if pk[d]['frm'])
    P = pk[day]
    from db import preflight, schema
    ok, pf_lines, _ = preflight.check(schema.FM)          # hard rule: the master is read where it lives, pre-flighted
    if not ok:
        print(chr(10).join(pf_lines)); raise SystemExit('STOP: the Formulation Master has unaccepted changes; interface not built.')
    accepted = re.search(r'accepted version \(([^)]+)\)', pf_lines[0])
    fm = config.published_path(schema.FM)
    config.record_read(fm, 'master (interface copy)')
    wb = load_workbook(fm, read_only=True)

    def sheet(n):
        r = list(wb[n].iter_rows(values_only=True))
        return [dict(zip(r[0], x)) for x in r[1:]]

    mats = {m['Material ID']: {'code': m['Material Code'], 'name': m['Name'], 'st': m['IWPFT062 Status'], 'role': m['Role']}
            for m in sheet('Materials')}
    lines = [] if P['frm'] else draft_lines(day, P, line_layout({d: p for d, p in pk.items() if d <= day}))
    for pg in sorted(P['frm'], key=lambda x: LN[x['line_code']]):
        lc = pg['line_code']
        ext = {r['order']: r for e in P['ext'] if e['line'] == lc for r in e['rows']}
        groups = [{'orders': [{'order': o, 'prod': ext.get(o, {}).get('prod_code'), 'colors': ext.get(o, {}).get('colors'),
                               'thk': ext.get(o, {}).get('thk')} for o in g['orders']],
                   'formulas': [{'code': f['formula_code'], 'note': f.get('note', ''),
                                 'feeders': [{'col': c, 'mat': current(v['material']), 'set': v['set'], 'id': MAP.get(current(v['material']))}
                                             for c, v in f['feeders'].items() if v['material'] or v['set']]}
                                for f in g['formulas']]}
                  for g in pg['groups']]
        grp_of = {o: gi for gi, g in enumerate(pg['groups']) for o in g['orders']}
        sched = [{'order': r['order'], 'prod': r['prod_code'], 'colors': r.get('colors'), 'thk': r.get('thk'), 'gsm': r.get('gsm'),
                  'size': f"{r.get('order_width', '')} × {r.get('order_length', '')}", 'die': r.get('die'),
                  'plt': r.get('num_plt'), 'lbs': r.get('weight_lbs'), 'si': r.get('special_instructions') or '',
                  'hw': r.get('handwritten') or '', 'g': grp_of.get(r['order']), 'plt_real': real_pallets(r)}
                 for e in sorted(P['ext'], key=lambda e: e['scan_page']) if e['line'] == lc for r in e['rows']]
        lines.append({'code': lc, 'no': LN[lc], 'dosing': DOS.get(lc, 'AUGER'), 'cols': pg['feeder_columns'],
                      'ac': pg.get('header_note', ''), 'groups': groups, 'notes': pg.get('footnotes', []), 'sched': sched})
    stats = draft_vs_tech(day, P)
    stats['orders'] = sum(len(g['orders']) for ln in lines for g in ln['groups']) if P['frm'] else sum(len(ln['sched']) for ln in lines)
    stats['lines'] = len(lines)
    stats['source'] = 'tech' if P['frm'] else 'draft'     # 'draft': no Tech FRM pages yet, formulation as last issued
    data = {'day': day, 'dates': sorted(pk),
            'scans': {'packet': P['source_scan'], 'frm': (P.get('frm_source_scan') or P['source_scan']) if P['frm'] else ''},
            'lines': lines, 'materials': mats, 'look': lookup_runs(pk, day), 'stats': stats,
            'issues': [{'sev': i['Severity'], 'line': i['Line'], 'check': i['Check'], 'detail': i['Detail']} for i in sheet('Issues')],
            'master': {'formulas': len(sheet('Formulas')), 'settings': len(sheet('Line Settings')), 'materials': len(mats),
                       'accepted': accepted.group(1)[:16].replace('T', ' ') if accepted else '', 'changes': 0},
            'records': record_rows(), 'built': datetime.datetime.now().strftime('%Y-%m-%d %H:%M'),
            'open': json.loads((ROOT / 'ui' / 'open_items.json').read_text(encoding='utf-8')),
            'cust': json.loads((ROOT / 'ui' / 'customer_formulas.json').read_text(encoding='utf-8')),
            # James Kuo, 28 Oct 2026: "another tab for revision record" - "formulation change (logic or W%) and artifact
            # design change"; a row is added with each change
            'revs': json.loads((ROOT / 'ui' / 'revision_record.json').read_text(encoding='utf-8'))['rows'],
            'replaced': replaced(),
            # for "generate from a scan" on the Daily run tab (James Kuo, 29 Sep 2026: "let me put the same production
            # file in the artifact. see if it generate the right formulation"): every issued FRM by date, the
            # transcribed schedule by date (to check the page's reading), known product codes, approved decisions
            'issued': {d: {pg['line_code']: [{'orders': g['orders'],
                                              'formulas': [{'code': f['formula_code'], 'note': f.get('note', ''),
                                                            'feeders': [{'col': c, 'mat': current(v['material']), 'set': v['set'], 'id': MAP.get(current(v['material']))}
                                                                        for c, v in f['feeders'].items() if v['material'] or v['set']]}
                                                           for f in g['formulas']]} for g in pg['groups']]
                           for pg in p['frm']} for d, p in pk.items() if p['frm']},
            'sched_by_date': {d: [[e['line'], r['order'], r['prod_code']] for e in p['ext'] for r in e['rows']] for d, p in pk.items()},
            'products': product_codes(),
            'hist_pairs': history_pairs(),
            'approved': approved_decisions(sheet),
            'rules_by_date': rules_by_date(pk, day, line_layout({d: p for d, p in pk.items() if d <= day})),
            'rules_any': rules_any(day, line_layout({d: p for d, p in pk.items() if d <= day})),
            'linemeta': line_layout(pk),
            'code_rule': {'colours': sorted(product_code.COLOURS), 'families': sorted(product_code.FAMILIES)}}
    # every formulation shows its weight % (James, 27 Oct 2026: "make sure this weight is display on every one so i can
    # easily monitor them"): on an auger line each feeder's share = setting x the hopper's calibration slope over the total
    from render_frm import weight_pct, PRINT_AS
    for ln in data['lines']:
        for g in ln['groups']:
            for f in g['formulas']:
                for x in f['feeders']:
                    x['mat'] = PRINT_AS.get(x['mat'], x['mat'])      # as the Word file prints it (the black: KS NPC PE-B9000F)
                if ln['dosing'] == 'WEIGHT':
                    continue
                pct, why = weight_pct(ln['code'], {('', x['col']): (x['mat'], x['set'], None) for x in f['feeders']})
                f['wnote'] = why
                for x in f['feeders']:
                    x['w'] = (pct or {}).get(('', x['col']))
    t = (ROOT / 'ui' / 'page.template.html').read_text(encoding='utf-8')
    t = t.replace('/*__READER__*/', (ROOT / 'ui' / 'reader.js').read_text(encoding='utf-8'))   # the scan reader (tested on its own)
    blob = json.dumps(data, ensure_ascii=False, separators=(',', ':')).replace('</', '<\\/')
    page = t.replace('__DATA__', blob)
    eng_data()
    glyph_bank_js()
    # the artifact source is the page fragment (the Artifact tool adds the document skeleton); the dated copy kept in
    # SharePoint is opened as a file, so it carries its own skeleton and charset
    full = ('<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
            '</head><body>' + page + '</body></html>')
    for out, text in ((config.OUTPUT_DIR / f'Profile Formulation {day}.html', full), (config.OUTPUT_DIR / 'profile-formulation.html', page)):
        out.write_text(text, encoding='utf-8')
        print(out)


if __name__ == '__main__':
    main()
