"""The rules' answer for an order the daily run did not see, for the page's Generate box (James Kuo, 29 Oct 2026: the
revised SE23 page added H6AA271-1 after the run, and the page "has issues reading and finding the right formulation").

    python ui/rules_any.py <day>      -> work/rules_any/FRM Draft <day>.xlsx + work/rules_any/keys_<day>.json

Every product in the published Product Master is keyed the way daily/assign.py keys an order: derive() (prefix, grade,
colour, thickness; a two-colour premix by its middle colour) and the qualifiers its product alone carries (corn box,
SPA sign blank, 6.3 / 4.6 mm). One product stands for each key on each line, and the daily run's own resolve.py decides
them on the day's schedule with the rules and settings as they are now (a copy of the packets in work/rules_any, so
data/packets, out/ and the issued copies stay as they are). Instruction qualifiers (VOIDFORM, RUN WITH, roll, LTL, the
slitter) are not known from the product: the page says to check them on the paper.
"""
import datetime
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / 'daily'))
import config  # noqa: E402
import assign  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

OUT = ROOT / 'work' / 'rules_any'
ORDER = 'RULE{:05d}-1'          # the stand-in orders: never a real order number (no letter H / RP prefix)


def products():
    """Product Master -> [{prod_code, mat_spec, colors}] for the products whose material, grade, spec and colours are on file."""
    p = config.published_path('Product Master.xlsx')
    config.record_read(p, 'Product Master (rules for an order the daily run did not see)')
    ws = load_workbook(p, read_only=True)['Product Master']
    rows = ws.iter_rows(values_only=True)
    h = {k: i for i, k in enumerate(next(rows))}
    out = []
    for r in rows:
        code, mat, grade, spec, cols = (r[h[k]] for k in ('Product Code', 'Material', 'Grade', 'Spec', 'Colours (3 layers)'))
        if code and mat and grade and spec and cols:
            out.append({'prod_code': code, 'mat_spec': f'{mat} {grade} {spec}', 'colors': cols, 'special_instructions': ''})
    return out


def sig(line, row, day):
    """The key the rules decide by, as text, or None when the rules cannot key the product (assign.derive)."""
    key, _ = assign.derive(line, row)
    if key is None:
        return None
    want, _ = assign.wanted(line, row, day)
    return '|'.join(key) + '|' + ','.join(sorted(want))


def main(day):
    d = datetime.date.fromisoformat(day)
    src = config.packet_path(day)
    pk = json.loads(src.read_text(encoding='utf-8'))
    lines = sorted({pg['line'] for pg in pk['ext']})
    prods = products()
    rep, keys = {}, {}                  # (line, sig) -> stand-in order; per line: product -> sig
    rows = {ln: [] for ln in lines}
    for ln in lines:
        keys[ln] = {}
        for r in prods:
            s = sig(ln, r, d)
            if s is None:
                continue
            keys[ln][r['prod_code']] = s
            if (ln, s) not in rep:
                rep[(ln, s)] = ORDER.format(len(rep) + 1)
                rows[ln].append(dict(r, order=rep[(ln, s)], T='Y', die='', thk='', gsm='', cut_rows=[], handwritten='', unclear=''))
    # the day's packet with the stand-in rows as its schedule (its FRM pages kept: they are the day's issued settings)
    pdir = OUT / 'packets'
    if OUT.exists():
        shutil.rmtree(OUT)
    pdir.mkdir(parents=True)
    for p in config.PACKETS_DIR.glob('packet_*.json'):
        if p.name != src.name:
            shutil.copy2(p, pdir / p.name)
    syn = dict(pk, ext=[{'scan_page': i + 1, 'report_page': str(i + 1), 'line': ln, 'rows': rows[ln]} for i, ln in enumerate(lines)])
    (pdir / src.name).write_text(json.dumps(syn, ensure_ascii=False), encoding='utf-8')
    shutil.copy2(config.OUTPUT_DIR / 'Formulation Calc Library.xlsx', OUT / 'Formulation Calc Library.xlsx')
    env = dict(os.environ, PYTHONUTF8='1', OUTPUT_DIR=str(OUT), PACKETS_DIR=str(pdir), PKT_DATE=day)
    env.pop('HISTORY_BEFORE', None); env.pop('PKT_STAGE1', None)
    r = subprocess.run([sys.executable, str(ROOT / 'daily' / 'resolve.py')], env=env, capture_output=True, text=True, cwd=ROOT)
    if r.returncode:
        raise SystemExit('resolve.py failed:\n' + r.stdout[-2000:] + r.stderr[-2000:])
    (OUT / f'keys_{day}.json').write_text(json.dumps({'day': day, 'order': {f'{ln}|{s}': o for (ln, s), o in rep.items()},
                                                      'keys': keys}, ensure_ascii=False), encoding='utf-8')
    print(f'{day}: {len(rep)} keys on {len(lines)} lines from {len(prods)} products;', (r.stdout.strip().splitlines() or [''])[0])


if __name__ == '__main__':
    main(sys.argv[1])
