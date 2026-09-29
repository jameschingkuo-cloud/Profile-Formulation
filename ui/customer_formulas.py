"""Build ui/customer_formulas.json: the customer-specific formulas parked for James to name (29 Sep 2026).

James: "We should use this oppotunity to fix their mistake and create new fomulation code instead" and "put all
customer specific list on the side and let me think so we can name these". Nothing is renamed; this only lists them.
"""
import glob
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import config  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

APPS = [(r'void\s*form', 'VOIDFORM'), (r'sign\s*blank', 'Sign blank'), (r'corn\s*box', 'Corn box'), (r'roll', 'Roll material')]
PAT = re.compile(r'^([A-Z])([A-Z])([0-9A-Z])(\d{3})([A-Z]{2})([0-9A-Z])$')


def app_of(note):
    return next((a for rx, a in APPS if re.search(rx, note or '', re.I)), None)


def main():
    packets = [json.load(open(f, encoding='utf-8')) for f in sorted(glob.glob(str(config.PACKETS_DIR / 'packet_*.json')))]
    # every code known (Calc Library + FRM) -> sequence numbers taken per family
    codes = set()
    wb = load_workbook(config.published_path('Formulation Calc Library.xlsx'), read_only=True)
    rows = wb['Formula Library'].iter_rows(values_only=True)
    h = [str(x) for x in next(rows)]
    ci = next(i for i, x in enumerate(h) if 'Formula' in x and 'Code' in x)
    codes |= {str(r[ci]).strip() for r in rows if r[ci]}
    for p in packets:
        codes |= {fm['formula_code'] for pg in p['frm'] for g in pg['groups'] for fm in g['formulas']}
    fam = defaultdict(set)
    for c in codes:
        m = PAT.match(c)
        if m:
            fam[m[1] + m[2] + m[3] + '·' + m[5] + m[6]].add(m[4])

    items = {}
    for p in packets:
        prod = {r['order']: r['prod_code'] for pg in p['ext'] for r in pg['rows']}
        for pg in p['frm']:
            for g in pg['groups']:
                for fm in g['formulas']:
                    app = app_of(fm.get('note'))
                    if not app:
                        continue
                    k = (pg['line_code'], fm['formula_code'], app, fm.get('note', ''))
                    it = items.setdefault(k, {'line': pg['line_code'], 'code': fm['formula_code'], 'app': app,
                                              'note': fm.get('note', ''), 'orders': set(), 'products': set(), 'dates': set()})
                    it['orders'] |= set(g['orders'])
                    it['products'] |= {prod[o] for o in g['orders'] if o in prod}
                    it['dates'].add(p['packet_date'])
                    it['settings'] = [f"{c.replace('Extruder ', '')} {v['material']} {v['set']}" for c, v in fm['feeders'].items()
                                      if v['material'] or v['set']]
                    it['also_plain'] = any(f2['formula_code'] == fm['formula_code'] and not app_of(f2.get('note'))
                                           for pg2 in p['frm'] if pg2['line_code'] == pg['line_code']
                                           for g2 in pg2['groups'] for f2 in g2['formulas'])
    out = []
    for it in sorted(items.values(), key=lambda x: (x['app'], x['line'], x['code'])):
        m = PAT.match(it['code'])
        famkey = (m[1] + m[2] + m[3] + '·' + m[5] + m[6]) if m else None
        out.append({**it, 'orders': sorted(it['orders']), 'products': sorted(it['products']), 'dates': sorted(it['dates']),
                    'family': famkey, 'taken': sorted(fam.get(famkey, [])) if famkey else []})
    data = {'updated': '2026-09-29', 'items': out,
            'finding': {'codes_known': len(codes), 'fit_9': sum(1 for c in codes if PAT.match(c)),
                        'pattern': '[base][requirement][blend][sequence, 3 digits][colour, 2][thickness, 1]',
                        'not_fitting': sorted(c for c in codes if not PAT.match(c))}}
    (ROOT / 'ui' / 'customer_formulas.json').write_text(json.dumps(data, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')
    print(len(out), 'customer-specific formulas;', data['finding']['fit_9'], 'of', len(codes), 'codes fit the 9-character pattern')


if __name__ == '__main__':
    main()
