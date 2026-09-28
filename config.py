"""Paths for this machine, and the record of what each build read.

Every script imports this module. Paths come from, in order: an environment variable of the same name,
local_settings.json (next to this file, not in git), or the default below. The defaults keep everything
inside the repo, so nothing touches SharePoint until publish.py is run on purpose.
"""
import hashlib, json, os, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent
_LOCAL = ROOT / 'local_settings.json'
_S = json.loads(_LOCAL.read_text(encoding='utf-8')) if _LOCAL.exists() else {}


def _path(key, default):
    v = os.environ.get(key) or _S.get(key)
    return Path(v).expanduser() if v else (Path(default) if default else None)


# ---- where things live ----------------------------------------------------------------------------------
OUTPUT_DIR = _path('OUTPUT_DIR', ROOT / 'out')            # builds write here first (staging, not in git)
PUBLISH_DIR = _path('PUBLISH_DIR', None)                   # the SharePoint-synced "Production Formulation Automation" folder
CALC_DIR = _path('CALC_DIR', ROOT / 'inputs' / 'calc')     # Tech's "SExx Formulation.xls" workbooks (live location: handoff §10 Q14)
SCAN_DIR = _path('SCAN_DIR', ROOT / 'inputs' / 'scans')    # daily packet scans (PDF)
HANDOFF = _path('HANDOFF', None)                           # the handoff document (lives in PUBLISH_DIR)
WORK_DIR = _path('WORK_DIR', ROOT / 'work')                # intermediate files: parsed.pkl, issues.json, calc_products.json, reads.json
DATA_DIR = ROOT / 'data'
PACKETS_DIR = DATA_DIR / 'packets'                         # transcribed daily packets: packet_YYYY-MM-DD.json (kept in git)
TRUTH_CSV = DATA_DIR / 'ext_truth_2026-09-23.csv'          # 82 verified EXT rows (scan reader training labels)
GLYPH_BANK = ROOT / 'scan_reader' / 'glyph_bank.npz'

for _d in (OUTPUT_DIR, WORK_DIR):
    _d.mkdir(parents=True, exist_ok=True)


def packet_path(date):
    """data/packets/packet_YYYY-MM-DD.json for a packet date (str or date)."""
    return PACKETS_DIR / f'packet_{date}.json'


def latest_packet():
    ps = sorted(PACKETS_DIR.glob('packet_*.json'))
    return ps[-1] if ps else None


# ---- HARD RULE support: record what a build read, so publish.py can refuse to overwrite a file that changed since ----
READS = WORK_DIR / 'reads.json'


def sha256(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for chunk in iter(lambda: f.read(1 << 20), b''):
            h.update(chunk)
    return h.hexdigest()


def content_hash(path):
    """What a file SAYS, not its bytes. For .xlsx: SHA-256 over every sheet name and every row of cell values
    (formulas as text), so SharePoint's metadata rewrite on upload does not count as a change but any edited cell does.
    For other files: SHA-256 of the bytes. Same function as the handoff's checksum manifest (§11); first 16 hex shown there."""
    p = Path(path)
    if p.suffix.lower() != '.xlsx':
        return sha256(p)
    from openpyxl import load_workbook
    wb = load_workbook(p, read_only=True)
    h = hashlib.sha256()
    for sn in wb.sheetnames:
        h.update(sn.encode())
        for row in wb[sn].iter_rows(values_only=True):
            h.update(json.dumps(row, default=str).encode())
    wb.close()
    return h.hexdigest()


def record_read(path, role=''):
    """Call on every source file a build reads. Stores size, modified time, SHA-256 and content hash in work/reads.json."""
    p = Path(path)
    reads = json.loads(READS.read_text(encoding='utf-8')) if READS.exists() else {}
    st = p.stat()
    reads[str(p.resolve())] = {'role': role, 'bytes': st.st_size,
                               'modified': datetime.datetime.fromtimestamp(st.st_mtime).isoformat(timespec='seconds'),
                               'sha256': sha256(p), 'content': content_hash(p),
                               'read_at': datetime.datetime.now().isoformat(timespec='seconds')}
    READS.write_text(json.dumps(reads, indent=1), encoding='utf-8')
    return reads[str(p.resolve())]


def reads():
    return json.loads(READS.read_text(encoding='utf-8')) if READS.exists() else {}
